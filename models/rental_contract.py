# -*- coding: utf-8 -*-

# from datetime import datetime, timedelta, date
# from odoo.exceptions import UserError, ValidationError
# from dateutil.relativedelta import relativedelta
import base64
from werkzeug.urls import url_encode
import logging
import io
from openpyxl.styles import Font, PatternFill
from odoo import models, fields, api, _
from collections import defaultdict
import html
from odoo.exceptions import UserError, ValidationError
logger = logging.getLogger(__name__)
from ..services import rental_contract_services as rcs

from urllib.parse import quote
from openpyxl import load_workbook
from datetime import date, timedelta
import base64
import io

class RentalContract(models.Model):
    _name = 'rental.contract'
    _inherit = ['mail.thread', 'mail.activity.mixin', 'mc.group.mixin']
    _description = 'Rental Contract'
    _rec_name = "code"

    code = fields.Char(
        string='Mã hợp đồng',
        index=True,
        required=True,
        default=lambda self: self.env['ir.sequence'].next_by_code('rental.contract'),
    )
    name = fields.Char(string='Tên hợp đồng')
    contract_date = fields.Date(string='Contract date', default=date.today())
    contract_number = fields.Char(
        string='Số hợp đồng',
        help='Số hợp đồng ghi trên văn bản (xuất khi Tải HĐ / In HĐ).',
        tracking=True,
    )
    prices_include_tax = fields.Boolean(
        string='Đã bao gồm thuế',
        default=False,
        tracking=True,
        help='Đơn giá trên bảng báo giá đã gồm thuế hay chưa.',
    )
    rental_billing_mode = fields.Selection(
        selection=[
            ("day", "Thuê theo ngày"),
            ("month", "Thuê theo tháng"),
        ],
        string="Cách tính tiền thuê",
        default="month",
        required=True,
        tracking=True,
        help="Ảnh hưởng bảng thanh toán / hóa đơn: theo ngày dùng Giá thuê/ngày trên sản phẩm; "
             "theo tháng: (Giá thuê tháng / số ngày trong tháng kỳ lập hóa đơn) × số ngày × SL.",
    )
    monthly_day_basis = fields.Selection(
        selection=[
            ("calendar", "Theo số ngày thực của tháng"),
            ("fixed_30", "Cố định 30 ngày/tháng"),
        ],
        string="Cơ sở quy đổi giá ngày (thuê theo tháng)",
        default="calendar",
        required=True,
        tracking=True,
        help="Chỉ áp dụng khi 'Cách tính tiền thuê' = Thuê theo tháng.\n"
             "• Theo số ngày thực: giá ngày = giá tháng ÷ số ngày của tháng kỳ thanh toán.\n"
             "• Cố định 30 ngày: giá ngày = giá tháng ÷ 30 (không phụ thuộc vào tháng).",
    )
    include_start_day_bob = fields.Boolean(
        string="Dư đầu kỳ: tính cả ngày thuê",
        default=True,
        tracking=True,
        help="Khi bật, số ngày thuê của phần Dư đầu kỳ sẽ tính bao gồm cả Ngày thuê.",
    )
    include_start_day_current = fields.Boolean(
        string="Thuê kỳ này: tính cả ngày thuê",
        default=False,
        tracking=True,
        help="Khi bật, số ngày thuê của phần Thuê kỳ này sẽ tính bao gồm cả Ngày thuê.",
    )
    minimum_rental_billing_mode = fields.Selection(
        [
            ("upfront", "Tính đủ kỳ tối thiểu vào tháng trả hàng"),
            ("spread", "Rải kỳ tối thiểu theo từng tháng"),
        ],
        string="Cách tính kỳ tối thiểu",
        default="upfront",
        required=True,
        tracking=True,
        help="Cách xử lý phần trả sớm chưa đủ kỳ tối thiểu:\n"
             "• Tính đủ kỳ tối thiểu vào tháng trả hàng (mặc định): ngay tại bảng thanh toán "
             "của tháng KH trả hàng, phần SL trả sẽ được tính tiền cho tới hết kỳ tối thiểu "
             "(ví dụ thuê 10/06, trả 20/06, kỳ 2 tháng → tính luôn 10/06→10/08 trong bảng tháng 6).\n"
             "• Rải kỳ tối thiểu theo từng tháng: phần trả sớm vẫn được coi là đang thuê cho tới "
             "hết kỳ tối thiểu và tính tiền dàn trải qua từng tháng tương ứng.",
    )
    minimum_rental_months = fields.Integer(
        string="Kỳ thuê tối thiểu (tháng)",
        default=2,
        tracking=True,
        help="Quy định thuê tối thiểu. Nếu trả sản phẩm trước khi đủ số tháng này (tính từ "
             "ngày giao của lô tương ứng) thì vẫn tính tiền đủ kỳ tối thiểu. Đặt 0 để tắt.",
    )
    minimum_penalty_current_period_only = fields.Boolean(
        string="Chỉ phạt sản phẩm thuê từ tháng này",
        default=False,
        tracking=True,
        help="Khi bật: chỉ áp dụng phạt kỳ tối thiểu cho số lượng giao trong kỳ thanh toán "
             "đang lập. Phần trả vượt quá SL thuê trong kỳ (khớp vào dư đầu kỳ / lô trước kỳ) "
             "tính và hiển thị như không phạt. Khi tắt: giữ hành vi cũ (phạt theo từng lô).",
    )
    transport_fee_share_min_months = fields.Integer(
        string="Ngưỡng chia sẻ phí VC (tháng)",
        default=0,
        tracking=True,
        help="Thuê luôn có 2 chiều vận chuyển (đi và về). Nếu thời gian thuê ≥ giá trị này "
             "thì bên cho thuê và bên thuê cùng chịu phí vận chuyển; nếu ngắn hơn thì bên thuê "
             "chịu hết. Đặt 0 = luôn bên thuê chịu hết (không đạt ngưỡng chia sẻ). "
             "Field dùng cho điều khoản / xuất báo giá; chưa đổi logic chia phí trên hóa đơn.",
    )
    price_change_ids = fields.One2many(
        'rental.contract.line.price',
        'contract_id',
        string="Điều chỉnh giá theo thời gian",
        help="Các mốc thay đổi đơn giá thuê có hiệu lực từ một ngày nhất định.",
    )

    company_partner_id = fields.Many2one(
        'res.partner',
        string='Company Partner (helper)',
        related='company_id.partner_id',
        store=False
    )
    currency_id = fields.Many2one(
        'res.currency', string="Currency", required=True,
        default=lambda self: self.env.company.currency_id.id
    )
    staff_id = fields.Many2one(
        'res.users',
        string="Staff",
        default=lambda self: self.env.user,  # current logged-in user
        index=True,
        tracking=True,
        help="User responsible for this contract."
    )
    status = fields.Selection([
        ('new', 'Mới'),
        ('need_fix', 'Cần sửa'),
        ('need_approve', 'Cần xác nhận'),
        ('leader_approved', 'Leader đã xác nhận'),
        ('customer_confirmed', 'Khách hàng đã xác nhận'),
        ('active', 'Đang hiệu lực'),
        ('finish', 'Kết thúc'),
        ('break', 'Tạm dừng'),
    ], default='new', string='Trạng thái', tracking=True)
    edit_unlocked = fields.Boolean(
        string="Mở khóa để sửa",
        default=False,
        tracking=True,
        help="Khi bật, hợp đồng đang cho phép sửa ở các trạng thái ngoài 'new'/'need_fix'.",
    )
    can_edit = fields.Boolean(
        string="Có thể sửa",
        compute="_compute_can_edit",
    )

    a_company_party = fields.Many2one(
        'res.partner',
        string='A company',
        domain="[('customer_type', '=', 'renter'), ('is_company', '=', True), ('company_id', 'in', allowed_company_ids)]",
        required=True,
        tracking=True,
    )
    a_party = fields.Many2one(
        'res.partner',
        string='A representative user',
        domain="[('parent_id', '=', a_company_party)]",
        readonly="[('a_company_party', '=', False)]",
        required=True,
        tracking=True,
    )
    a_name = fields.Char(string='A Name', compute='_compute_a_party', store=True)
    a_address = fields.Char(string='A Address', compute='_compute_a_party', store=True)
    a_function = fields.Char(string='A Function', compute='_compute_a_party', store=True)
    a_phone = fields.Char(string='A Phone', compute='_compute_a_party', store=True)
    a_bank_account = fields.Char(string='A Bank account', compute='_compute_a_party', store=True)
    a_vat = fields.Char(string='A Vat', compute='_compute_a_party', store=True)

    a_executor_id = fields.Many2one(
        'res.partner',
        string='Người thực hiện bên A',
        domain="[('parent_id', '=', a_company_party)]",
        tracking=True,
        help="Người thực hiện bên A.",
    )
    a_executor_name = fields.Char(string='Tên người thực hiện bên A', compute='_compute_a_party', store=True)
    a_executor_function = fields.Char(string='Chức vụ người thực hiện bên A', compute='_compute_a_party', store=True)

    b_company_party = fields.Many2one(
        'res.partner',
        string='B company',
        domain="[('parent_id', '=', company_partner_id)]",
        default=lambda self: self.env.company.partner_id,
        required=True,
        tracking=True,
    )
    b_party = fields.Many2one(
        'res.partner',
        string='B representative user',
        domain="[('parent_id', '=', b_company_party)]",
        readonly="[('b_company_party', '=', False)]",
        required=True,
        tracking=True,
    )
    b_name = fields.Char(string='B Name', compute='_compute_b_party', store=True)
    b_address = fields.Char(string='B Address', compute='_compute_b_party', store=True)
    b_function = fields.Char(string='B Function', compute='_compute_b_party', store=True)
    b_phone = fields.Char(string='B Phone', compute='_compute_b_party', store=True)
    b_bank_account = fields.Char(string='B Bank account', compute='_compute_b_party', store=True)
    b_vat = fields.Char(string='B Vat', compute='_compute_b_party', store=True)

    b_executor_id = fields.Many2one(
        'res.partner',
        string='Người thực hiện bên B',
        domain="[('parent_id', '=', b_company_party)]",
        tracking=True,
        help="Người thực hiện bên B.",
    )
    b_executor_name = fields.Char(string='Tên người thực hiện bên B', compute='_compute_b_party', store=True)
    b_executor_function = fields.Char(string='Chức vụ người thực hiện bên B', compute='_compute_b_party', store=True)

    rental_contract_line_ids = fields.One2many('rental.contract.line', 'contract_id', string="Contract Lines", copy=True, tracking=True)

    rr_transport_ids = fields.One2many('rr.transport', 'rental_contract_id', string='Transports', tracking=True)
    rr_transport_line_ids = fields.One2many('rr.transport.line', 'rental_contract_id', string='Rental product', tracking=True)
    rr_transport_matrix_json = fields.Json(
        string='Transport matrix json',
        compute='_compute_rr_transport_matrix_html',
        help='This store matrix data of truck that will export excel and show: Confirmation table for rental volume'
    )
    # rr_transport_matrix_html = fields.Html(
    #     string='Confirmation table for rental volume',
    #     compute='_compute_rr_transport_matrix_html',
    #     sanitize=False,  # Allow our table classes/styles
    # )

    rental_invoice_ids = fields.One2many('rental.invoice', 'rental_contract_id', string='Rental invoice', tracking=True)
    amount_invoiced = fields.Monetary(string="Invoiced (Total)",
                                      currency_field="currency_id",
                                      compute="_compute_accounting_totals", store=False)
    amount_paid = fields.Monetary(string="Paid",
                                  currency_field="currency_id",
                                  compute="_compute_accounting_totals", store=False)
    amount_due = fields.Monetary(string="Amount Due",
                                 currency_field="currency_id",
                                 compute="_compute_accounting_totals", store=False)
    # invoice_start_date = fields.Date(string="Invoice start date")
    # invoice_end_date = fields.Date(string="Invoice end date")

    construction_work_id = fields.Many2one('construction.work', string="Gói Thầu", tracking=True, help="Each contract will be for on Construction work")
    construction_work_project_id = fields.Many2one(related='construction_work_id.project_id', string="Dự án", store=True, readonly=True)
    construction_work_address_ids = fields.Many2many(
        related='construction_work_id.address_ids',
        string='Địa chỉ công trình',
        readonly=True,
    )
    construction_work_address = fields.Text(
        string='Địa chỉ công trình (text)',
        compute='_compute_construction_work_address',
        store=True,
    )

    account_move_ids = fields.One2many(
        'account.move',
        'rental_contract_id',
        string='Invoices',
        tracking=True
    )
    invoice_count = fields.Integer(string='Invoice count', compute='_compute_invoice_count')

    transport_matrix_ids = fields.One2many(
        'rental.transport.matrix',
        'rental_contract_id',
        string='Confirmation tables',
        read_only=True,
    )
    transport_matrix_count = fields.Integer(
        string='Confirmation tables',
        compute='_compute_transport_matrix_count',
    )

    holiday_ids = fields.One2many(
        'rental.holiday',
        'rental_contract_id',
        string='Ngày nghỉ riêng',
        help="Các kỳ nghỉ áp dụng riêng cho hợp đồng này (cộng thêm vào ngày nghỉ "
             "toàn hệ thống khi tính số ngày thuê).",
    )

    deposit = fields.Float(string='Deposit', tracking=True)
    product_list_template_id = fields.Many2one(
        "rental.product.template.set",
        string="Mẫu sản phẩm",
        domain="[('company_id', '=', company_id)]",
        tracking=True,
        help="Chọn mẫu để tự động nạp sẵn danh sách product.template vào bảng báo giá.",
    )

    note = fields.Html(string='Note')

    @api.depends("status", "edit_unlocked")
    def _compute_can_edit(self):
        for rec in self:
            rec.can_edit = rec.status in ("new", "need_fix") or bool(rec.edit_unlocked)

    def _check_can_edit(self):
        for rec in self:
            if not rec.can_edit:
                raise UserError(_("Hợp đồng hiện đang bị khóa, không cho phép sửa. Vui lòng yêu cầu Leader mở khóa."))  # noqa: E501

    _sql_constraints = [
        ('unique_rental_contract_code', 'UNIQUE(code)',
         'The rental contract with this code had been created'),
    ]

    @api.depends('account_move_ids')
    def _compute_invoice_count(self):
        for rec in self:
            rec.invoice_count = len(rec.account_move_ids)

    @api.depends('transport_matrix_ids')
    def _compute_transport_matrix_count(self):
        for rec in self:
            rec.transport_matrix_count = len(rec.transport_matrix_ids)

    def action_open_transport_matrices(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'name': _('Confirmation tables'),
            'res_model': 'rental.transport.matrix',
            'view_mode': 'tree,form',
            'domain': [('rental_contract_id', '=', self.id)],
            'context': {'default_rental_contract_id': self.id},
        }

    def action_open_transport_import_wizard(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'name': _('Import xuất nhập kho từ Excel'),
            'res_model': 'rental.transport.import.wizard',
            'view_mode': 'form',
            'target': 'new',
            'context': {
                'default_rental_contract_id': self.id,
            },
        }

    def action_download_transport_import_template(self):
        self.ensure_one()
        url = f'/rental/rental-contract/transport-import-template/{self.id}/download'
        return {
            'type': 'ir.actions.act_url',
            'url': url,
            'target': 'self',
        }

    @api.depends('b_party', 'company_id')
    def _compute_accounting_totals(self):
        # Batch all contracts in one grouped query
        if not self:
            return
        Move = self.env['account.move'].with_context(active_test=False)

        domain = [
            ('state', '=', 'posted'),
            ('move_type', 'in', ('out_invoice', 'out_refund')),
            ('rental_contract_id', 'in', self.ids),
            ('company_id', 'in', self.mapped('company_id').ids),
        ]
        # amount_total and amount_residual are stored on account.move => fast to aggregate
        groups = Move.read_group(
            domain,
            ['amount_total:sum', 'amount_residual:sum', 'rental_contract_id'],
            ['rental_contract_id'],
        )
        by_contract = {g['rental_contract_id'][0]: g for g in groups}

        for contract in self:
            g = by_contract.get(contract.id)
            total = g['amount_total'] if g else 0.0
            residual = g['amount_residual'] if g else 0.0
            contract.amount_invoiced = total
            contract.amount_paid = total - residual
            contract.amount_due = residual

    @api.depends(
        'construction_work_id',
        'construction_work_id.address_ids',
        'construction_work_id.address_ids.name',
    )
    def _compute_construction_work_address(self):
        for rec in self:
            if rec.construction_work_id and rec.construction_work_id.address_ids:
                rec.construction_work_address = "\n".join(rec.construction_work_id.address_ids.mapped("name"))
            else:
                rec.construction_work_address = ""

    @api.depends(
        'a_party',
        'a_party.function',
        'a_company_party',
        'a_executor_id',
        'a_executor_id.name',
        'a_executor_id.function',
    )
    def _compute_a_party(self):
        for rec in self:
            a_executor = rec.a_executor_id
            rec.a_name = rec.a_party.name
            rec.a_address = rec.a_company_party.street
            rec.a_function = rec.a_party.function
            rec.a_phone = rec.a_company_party.phone
            rec.a_bank_account = rec.a_company_party.bank_account
            rec.a_vat = rec.a_company_party.vat
            rec.a_executor_name = a_executor.name if a_executor else False
            rec.a_executor_function = a_executor.function if a_executor else False

    @api.depends(
        'b_party',
        'b_party.function',
        'b_company_party',
        'b_executor_id',
        'b_executor_id.name',
        'b_executor_id.function',
    )
    def _compute_b_party(self):
        for rec in self:
            b_executor = rec.b_executor_id
            rec.b_name = rec.b_party.name
            rec.b_address = rec.b_company_party.street
            rec.b_function = rec.b_party.function
            rec.b_phone = rec.b_company_party.phone
            rec.b_bank_account = rec.b_company_party.bank_account
            rec.b_vat = rec.b_company_party.vat
            rec.b_executor_name = b_executor.name if b_executor else False
            rec.b_executor_function = b_executor.function if b_executor else False

    @api.model
    def create(self, vals):
        records = super().create(vals)
        for rec in records:
            staff = rec.staff_id
            leader = staff and staff.rental_leader_id or False
            if leader and leader.partner_id:
                rec.message_subscribe(partner_ids=[leader.partner_id.id])
        return records

    def write(self, vals):
        # Lock enforcement.
        # When contract is locked, we block writes to the editable fields only.
        # This avoids breaking side-effects coming from other models (e.g. stock picking save).
        if not self.env.context.get("rental_contract_allow_locked_write"):
            # Fields that correspond to "Base info" tab edits.
            blocked_fields = {
                "a_company_party",
                "a_party",
                "a_function",
                "a_executor_id",
                "a_address",
                "a_phone",
                "a_bank_account",
                "a_vat",
                "b_company_party",
                "b_party",
                "b_function",
                "b_executor_id",
                "b_address",
                "b_phone",
                "b_bank_account",
                "b_vat",
                "construction_work_id",
                "note",
                "rental_contract_line_ids",
                "price_change_ids",
                "product_list_template_id",
                # These are editable toggles; handled separately via context.
                # "status",
                # "edit_unlocked",
            }
            vals_keys = set(vals.keys())
            for rec in self:
                if not rec.can_edit and vals_keys & blocked_fields:
                    raise UserError(_("Hợp đồng hiện đang bị khóa, không cho phép sửa."))
        return super().write(vals)

    def _prepare_contract_line_vals_from_template(self, template):
        self.ensure_one()
        line_vals = []
        template_lines = template.line_ids.sorted(key=lambda l: (l.sequence, l.id))
        for sequence, template_line in enumerate(template_lines, start=1):
            product_tmpl = template_line.product_tmpl_id
            line_vals.append((0, 0, {
                "sequence": template_line.sequence or (sequence * 10),
                "product_tmpl_id": product_tmpl.id,
                "name": template_line.name or getattr(
                    product_tmpl,
                    "get_product_multiline_description_sale",
                    lambda: product_tmpl.display_name,
                )(),
                "product_uom_qty": template_line.product_uom_qty or 1,
                "price_unit": template_line.price_unit,
                "standard_price": template_line.standard_price or 0.0,
                "compensation_price": template_line.compensation_price or 0.0,
                "uom_id": template_line.uom_id.id or product_tmpl.uom_id.id or False,
            }))
        return line_vals

    @api.onchange("product_list_template_id")
    def _onchange_product_list_template_id(self):
        # No wipe-on-change: use action_apply_product_list_template (confirm button).
        return

    def action_apply_product_list_template(self):
        """Replace bảng báo giá lines from the selected product template set."""
        self.ensure_one()
        self._check_can_edit()
        if not self.product_list_template_id:
            raise UserError(_("Vui lòng chọn mẫu sản phẩm trước khi nạp."))
        self.write({
            "rental_contract_line_ids": (
                [(5, 0, 0)]
                + self._prepare_contract_line_vals_from_template(self.product_list_template_id)
            ),
        })
        return True

    def _rental_invoice_xlsx_load_workbook(self):
        self.ensure_one()
        data, start_row = self.env["rental.template"].get_rental_invoice_xlsx(
            self.company_id,
        )
        return load_workbook(io.BytesIO(data)), start_row

    def _rental_invoice_xlsx_apply_placeholders(self, ws, start_date, end_date):
        replacements = {
            "{{today_is}}": date.today().strftime("ngày %d tháng %m năm %Y"),
            "{{invoice_number}}": "",
            "{{contract_number}}": self.contract_number or "",
            "{{contract_date}}": self._quotation_contract_date_display(),
            "{{construction_work_project}}": self.construction_work_project_id.name or "",
            "{{construction_work_name}}": self.construction_work_id.name or "",
            "{{construction_work_address}}": self.construction_work_address or "",
            "{{b_company}}": self.b_party.parent_id.name or "",
            "{{b_address}}": self.b_address or "",
            "{{b_representative}}": self.b_name or "",
            "{{b_function}}": self.b_function or "",
            "{{a_company}}": self.a_party.parent_id.name or "",
            "{{a_representative}}": self.a_name or "",
            "{{a_function}}": self.a_function or "",
            "{{a_confirmer_function}}": self.a_function or "",
            "{{b_confirmer_function}}": self.b_function or "",
            "{{start_date}}": start_date.strftime("%d/%m/%Y"),
            "{{end_date}}": end_date.strftime("%d/%m/%Y"),
            "{{end_date_month_year}}": end_date.strftime("%m/%Y"),
        }
        from ..helper.xlsx_template_utils import replace_placeholders_in_sheet

        replace_placeholders_in_sheet(ws, replacements)

    def _transport_fee_line_dicts(self, transports):
        """Normalize rr.transport records into fee line dicts for HSTT / invoice."""
        return [
            {
                "date": t.start_rental_or_return_date,
                "code": t.code or "",
                "type": t.type,
                "amount": t.fee or 0.0,
                "transport_id": t.id,
            }
            for t in transports
        ]

    def _rental_unbilled_transport_fee_lines(self, transport_fee_until_date):
        """Phí VC chưa tính: fee > 0, chưa đánh dấu, ngày tính ≤ cutoff.

        Cutoff trống → không lấy phí nào (hoãn toàn bộ kỳ này).
        Không lọc theo kỳ thuê start/end — để gom phí bị hoãn từ tháng trước.
        """
        self.ensure_one()
        if not transport_fee_until_date:
            return []
        transports = self.rr_transport_ids.filtered(
            lambda t: t.state == "done"
            and t.fee
            and t.start_rental_or_return_date
            and not t.fee_billed_date
            and not t.fee_invoice_id
            and t.start_rental_or_return_date <= transport_fee_until_date
        ).sorted("start_rental_or_return_date")
        return self._transport_fee_line_dicts(transports)

    def _rental_invoice_transport_fee_lines(self, move):
        """Phí VC đã gắn vào hóa đơn (dùng khi regenerate KLCT+HSTT)."""
        self.ensure_one()
        if not move:
            return []
        transports = self.rr_transport_ids.filtered(
            lambda t: t.state == "done"
            and t.fee
            and t.fee_invoice_id
            and t.fee_invoice_id.id == move.id
        ).sorted("start_rental_or_return_date")
        return self._transport_fee_line_dicts(transports)

    def _rental_period_transport_fee_lines(
        self, start_date, end_date, transport_fee_until_date=None, fee_invoice=None
    ):
        """Resolve fee lines for export.

        - fee_invoice: regenerate — phí đã gắn hóa đơn đó.
        - otherwise: unbilled with cutoff (mặc định end_date nếu không truyền).
        """
        self.ensure_one()
        if fee_invoice:
            return self._rental_invoice_transport_fee_lines(fee_invoice)
        cutoff = transport_fee_until_date
        if cutoff is None:
            # Backward-compat for callers that only pass period dates.
            cutoff = end_date
        return self._rental_unbilled_transport_fee_lines(cutoff)

    def _mark_transport_fees_billed(self, fee_lines, move, billed_date):
        """Auto-mark transports included in HSTT as billed on the given invoice."""
        self.ensure_one()
        if not fee_lines or not move or not billed_date:
            return
        transport_ids = [fl.get("transport_id") for fl in fee_lines if fl.get("transport_id")]
        if not transport_ids:
            return
        transports = self.env["rr.transport"].browse(transport_ids).exists()
        transports.write({
            "fee_billed_date": billed_date,
            "fee_invoice_id": move.id,
        })

    def _rental_period_compensation_lines(self, start_date, end_date):
        """Các dòng đền bù (tiền phạt mất/hỏng) phát sinh trong kỳ.

        Lấy theo phiếu Đền bù (rr.transport.type == 'compensation', state=done) có
        ngày tính nằm trong [start_date, end_date] và có tiền phạt > 0.
        """
        self.ensure_one()
        lines = self.env["rr.transport.line"].search(
            [
                ("transport_id", "in", self.rr_transport_ids.ids),
                ("transport_id.type", "=", "compensation"),
                ("transport_id.state", "=", "done"),
                ("start_rental_or_return_date", ">=", start_date),
                ("start_rental_or_return_date", "<=", end_date),
            ],
            order="start_rental_or_return_date ASC, id ASC",
        )
        result = []
        for line in lines:
            if not line.fine_amount:
                continue
            result.append({
                "date": line.start_rental_or_return_date,
                "product_id": line.product_id.id,
                "product_name": line.product_id.display_name,
                "uom_name": line.product_id._get_staff_display_uom().name,
                "qty": line.qty or 0,
                "damage_ratio": line._damage_ratio_value(),
                "amount": line.fine_amount or 0.0,
            })
        return result

    def _payment_table_subtotal(self, bob_map, map_map, fee_lines=None, compensation_lines=None):
        """Sum line amounts written to column I of the payment table."""
        total = 0.0
        for bucket in (bob_map, map_map):
            if not bucket:
                continue
            for _tmpl_id in bucket:
                for line in bucket[_tmpl_id].get("lines", []):
                    total += line.get("total_amount") or 0.0
        for fee_line in fee_lines or []:
            total += fee_line.get("amount") or 0.0
        for comp_line in compensation_lines or []:
            total += comp_line.get("amount") or 0.0
        return total

    def _rental_invoice_xlsx_apply_payment_totals(self, ws, bob_map, map_map, fee_lines=None, compensation_lines=None):
        """Fill summary placeholders and amount-in-words anywhere in the sheet."""
        from ..helper.xlsx_template_utils import replace_placeholders_in_sheet

        subtotal = self._payment_table_subtotal(bob_map, map_map, fee_lines, compensation_lines)
        vat = round(subtotal * 0.08)
        total_after_tax = int(round(subtotal + vat))
        amount_words = self.env["amount_to_text.vi"].vn_amount_to_text(total_after_tax)
        replace_placeholders_in_sheet(
            ws,
            {
                "{{total_after_tax_string}}": amount_words,
                "{{subtotal_before_tax}}": int(round(subtotal)),
                "{{vat_amount}}": int(round(vat)),
                "{{total_after_tax}}": total_after_tax,
            },
        )
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if not isinstance(val, str):
                    continue
                if "{{total_after_tax_string}}" in val:
                    cell.value = val.replace("{{total_after_tax_string}}", amount_words)
                elif "Bằng chữ" in val or "bằng chữ" in val:
                    cell.value = f"(Bằng chữ: {amount_words}/.)"

    # Chỉ tô vàng dòng "Cộng" cho dễ nhìn; các tiêu đề phụ chỉ in đậm.
    _FILL_YELLOW = PatternFill("solid", fgColor="FFFF00")

    @staticmethod
    def _font_like(cell, color):
        base = cell.font
        return Font(name=base.name, size=base.size, bold=base.bold, italic=base.italic, color=color)

    def _red_font_like(self, cell):
        """Font đỏ (dòng trả hàng)."""
        return self._font_like(cell, "FFFF0000")

    def _blue_font_like(self, cell):
        """Font xanh nhạt (dòng chuyển thừa trả lại — không tính tiền)."""
        return self._font_like(cell, "FF5B9BD5")

    @staticmethod
    def _fill_row(ws, r, col_from, col_to, fill):
        for c in range(col_from, col_to + 1):
            ws.cell(r, c).fill = fill

    def _write_subheader(self, ws, r, text):
        cell = ws.cell(r, 4)
        cell.value = text
        cell.font = Font(bold=True)

    def _holiday_days_for_billing_span(self, start_date, end_date, include_start_day):
        """Holiday days on the same effective range as rental_days_between_with_holiday."""
        self.ensure_one()
        if include_start_day:
            effective_start = start_date
        else:
            effective_start = start_date + timedelta(days=1)
        if effective_start > end_date:
            return 0
        return rcs._holiday_days_between(self.env, self, effective_start, end_date)

    @staticmethod
    def _rental_days_excel_formula(row, include_start_day, holiday_days=0):
        """Excel formula for rental days: (C-B[+1]) minus holidays, floored at 0."""
        base = f"C{row}-B{row}+1" if include_start_day else f"C{row}-B{row}"
        holiday_days = int(holiday_days or 0)
        if holiday_days <= 0:
            return f"={base}"
        return f"=MAX(0,{base}-{holiday_days})"

    def _write_billing_line(
        self,
        ws,
        start_row,
        count,
        line,
        content,
        use_excel_formulas=False,
        klct_meta=None,
        month_day_dim=0,
    ):
        """Một dòng tính tiền: ngày bắt đầu → kết thúc, SL, số ngày, đơn giá, thành tiền."""
        r = start_row + count
        if use_excel_formulas:
            self._write_excel_date(ws.cell(r, 2), line["start_date"])
            self._write_excel_date(ws.cell(r, 3), line["end_date"])
        else:
            ws.cell(r, 2).value = line["start_date"].strftime("%d/%m/%Y")
            ws.cell(r, 3).value = line["end_date"].strftime("%d/%m/%Y")
        ws.cell(r, 4).value = content
        ws.cell(r, 5).value = line["uom_name"]
        qty_formula = None
        if use_excel_formulas and klct_meta:
            qty_formula = self._klct_qty_formula_for_line(klct_meta, line)
        if qty_formula:
            ws.cell(r, 6).value = qty_formula
        else:
            ws.cell(r, 6).value = line["qty"]
        # Month mode + Excel formulas: H must be day_price × period dim so I=F*G*H/dim
        # matches Python (penalty lines may have end_date in another month → wrong display).
        if use_excel_formulas:
            include = line.get("include_start_day")
            if include is None:
                include = bool(line.get("is_bob"))
            holiday_days = self._holiday_days_for_billing_span(
                line["start_date"], line["end_date"], include
            )
            ws.cell(r, 7).value = self._rental_days_excel_formula(r, include, holiday_days)
            if self.rental_billing_mode == "month" and month_day_dim:
                ws.cell(r, 8).value = float(line.get("unit_price") or 0) * int(month_day_dim)
                ws.cell(r, 9).value = f"=F{r}*G{r}*H{r}/{int(month_day_dim)}"
            else:
                ws.cell(r, 8).value = line.get("display_unit_price", line["unit_price"])
                ws.cell(r, 9).value = f"=F{r}*G{r}*H{r}"
        else:
            ws.cell(r, 7).value = line["rental_days"]
            ws.cell(r, 8).value = line.get("display_unit_price", line["unit_price"])
            ws.cell(r, 9).value = line["total_amount"]
        return count + 1

    def _write_aggregated_return_row(
        self,
        ws,
        start_row,
        count,
        agg,
        block,
        content,
        use_excel_formulas=False,
        month_day_dim=0,
    ):
        """Dòng gộp phần trả (phạt/đã trả): không ngày bắt đầu, chỉ SL + số ngày + tiền."""
        r = start_row + count
        ws.cell(r, 4).value = content
        ws.cell(r, 5).value = block["uom_name"]
        ws.cell(r, 6).value = agg["qty"]
        if use_excel_formulas and agg.get("start_date") and agg.get("end_date"):
            self._write_excel_date(ws.cell(r, 2), agg["start_date"])
            self._write_excel_date(ws.cell(r, 3), agg["end_date"])
            include = agg.get("include_start_day")
            if include is None:
                include = bool(agg.get("is_bob")) or (
                    agg.get("deliver_date") and agg.get("start_date")
                    and agg["deliver_date"] <= agg["start_date"]
                )
            holiday_days = self._holiday_days_for_billing_span(
                agg["start_date"], agg["end_date"], include
            )
            ws.cell(r, 7).value = self._rental_days_excel_formula(r, include, holiday_days)
            if self.rental_billing_mode == "month" and month_day_dim:
                ws.cell(r, 8).value = float(agg.get("unit_price") or 0) * int(month_day_dim)
                ws.cell(r, 9).value = f"=F{r}*G{r}*H{r}/{int(month_day_dim)}"
            else:
                ws.cell(r, 8).value = agg.get("display_unit_price", agg["unit_price"])
                ws.cell(r, 9).value = f"=F{r}*G{r}*H{r}"
        else:
            ws.cell(r, 7).value = agg["rental_days"]
            ws.cell(r, 8).value = agg.get("display_unit_price", agg["unit_price"])
            ws.cell(r, 9).value = agg["total_amount"]
        return count + 1

    @staticmethod
    def _write_excel_date(cell, value):
        """Write a Python date as an Excel date with DD/MM/YYYY display format."""
        cell.value = value
        cell.number_format = "DD/MM/YYYY"

    @staticmethod
    def _excel_sheet_ref(sheet_title):
        """Quote sheet name for formula references when needed."""
        title = sheet_title or ""
        if any(ch in title for ch in (" ", "-", "'")):
            safe = title.replace("'", "''")
            return f"'{safe}'"
        return title

    @classmethod
    def _klct_qty_formula_for_line(cls, klct_meta, line):
        """Build F-column formula pointing at KLCT qty (or Tổng MD) for this billing line.

        Only link when the KLCT cell(s) match billing ``line["qty"]``. Otherwise return
        None so the writer stores a literal (avoids double-counting split returns /
        penalties — DEC-16, extended to in-period rows).
        """
        if not klct_meta:
            return None
        tmpl_id = line.get("tmpl_id")
        if not tmpl_id:
            return None
        col = (klct_meta.get("qty_col_by_tmpl_id") or {}).get(tmpl_id)
        if not col:
            return None
        from openpyxl.utils import get_column_letter

        letter = get_column_letter(col)
        sheet_ref = cls._excel_sheet_ref(klct_meta.get("sheet_title") or "")
        line_qty = float(line.get("qty") or 0)
        qty_by_row = klct_meta.get("qty_by_row_and_tmpl_id") or {}

        def _rows_qty(rows):
            return sum(float(qty_by_row.get((r, tmpl_id), 0) or 0) for r in rows)

        if line.get("is_bob"):
            opening_row = klct_meta.get("opening_row")
            if not opening_row:
                return None
            opening_qty_by_tmpl = klct_meta.get("opening_qty_by_tmpl_id")
            if opening_qty_by_tmpl is not None:
                opening_qty = opening_qty_by_tmpl.get(tmpl_id, 0) or 0
                if abs(float(opening_qty) - line_qty) > 1e-6:
                    return None
            elif qty_by_row:
                if abs(_rows_qty([opening_row]) - line_qty) > 1e-6:
                    return None
            return f"={sheet_ref}!{letter}{opening_row}"

        deliver_date = line.get("deliver_date")
        rows_by_date_and_direction = klct_meta.get("data_rows_by_date_and_direction")
        if rows_by_date_and_direction is not None and line_qty < 0:
            rows = (
                (rows_by_date_and_direction.get(deliver_date) or {}).get("return")
                or []
            )
        else:
            # Positive merged billing lines may net deliveries and orphan returns from
            # the same day, so keep all movement rows. This is also the fallback for
            # callers providing the old metadata shape.
            rows = (klct_meta.get("data_rows_by_date") or {}).get(deliver_date) or []
        if not rows:
            return None
        if qty_by_row and abs(_rows_qty(rows) - line_qty) > 1e-6:
            return None
        if len(rows) == 1:
            return f"={sheet_ref}!{letter}{rows[0]}"
        refs = ",".join(f"{sheet_ref}!{letter}{r}" for r in rows)
        return f"=SUM({refs})"

    @staticmethod
    def _copy_xlsx_sheet(src_ws, dest_wb, title):
        """Copy a worksheet (values, styles, merges, dimensions) into another workbook."""
        import copy

        dest_ws = dest_wb.create_sheet(title=title)
        for row in src_ws.iter_rows():
            for cell in row:
                new_cell = dest_ws.cell(row=cell.row, column=cell.column, value=cell.value)
                if cell.has_style:
                    new_cell.font = copy.copy(cell.font)
                    new_cell.border = copy.copy(cell.border)
                    new_cell.fill = copy.copy(cell.fill)
                    new_cell.number_format = cell.number_format
                    new_cell.protection = copy.copy(cell.protection)
                    new_cell.alignment = copy.copy(cell.alignment)
        for merged in list(src_ws.merged_cells.ranges):
            dest_ws.merge_cells(str(merged))
        for col_letter, dim in src_ws.column_dimensions.items():
            dest_ws.column_dimensions[col_letter].width = dim.width
            dest_ws.column_dimensions[col_letter].hidden = dim.hidden
        for idx, dim in src_ws.row_dimensions.items():
            dest_ws.row_dimensions[idx].height = dim.height
            dest_ws.row_dimensions[idx].hidden = dim.hidden
        dest_ws.sheet_view.showGridLines = src_ws.sheet_view.showGridLines
        if src_ws.freeze_panes:
            dest_ws.freeze_panes = src_ws.freeze_panes
        return dest_ws

    @staticmethod
    def _return_row_content(block, agg, penalty=False):
        """Text dòng đối ứng trả: '{SP} trả trong tháng {tháng trả} - thuê từ {ngày giao}
        [- phạt {N} tháng]'."""
        rd = agg.get("return_date")
        deliver = agg.get("deliver_date")
        text = _("%(name)s trả trong tháng %(m)s - thuê từ %(d)s") % {
            "name": block["product_name"],
            "m": rd.strftime("%m/%Y") if rd else "",
            "d": deliver.strftime("%d/%m/%Y") if deliver else "",
        }
        if penalty and agg.get("min_months"):
            text += _(" - phạt %s tháng") % agg["min_months"]
        return text

    def _write_product_payment_block(
        self,
        ws,
        block,
        start_row,
        count,
        use_excel_formulas=False,
        klct_meta=None,
        month_day_dim=0,
    ):
        """Một sản phẩm = một khối: (1) đơn thuê bình thường; (2) trả không phạt
        (credit âm từ ngày trả → cuối kỳ); (3) đối ứng + phạt kỳ tối thiểu; (4) chuyển
        thừa; (5) dòng Cộng SL đang thuê."""
        # Phần 1: đơn thuê bình thường (kể cả SL đã gộp từ trả không phạt).
        for line in block["normal_lines"]:
            content = line["product_name"]
            if line.get("is_bob"):
                content = _("%s (dư đầu kỳ)") % content
            count = self._write_billing_line(
                ws,
                start_row,
                count,
                line,
                content,
                use_excel_formulas=use_excel_formulas,
                klct_meta=klct_meta,
                month_day_dim=month_day_dim,
            )

        rc = block["return_calc"]
        if rc:
            mode = rc.get("display_mode") or "legacy"

            # Trả không phạt: dòng đỏ có tiền âm (ảnh 2) — không cần sub-header đối ứng.
            for cred in rc.get("credit_returns") or []:
                count = self._write_credit_return_row(
                    ws,
                    start_row,
                    count,
                    cred,
                    block,
                    use_excel_formulas=use_excel_formulas,
                    klct_meta=klct_meta,
                    month_day_dim=month_day_dim,
                )

            # Chuyển thừa trả lại khi credit-only (không vào khối đối ứng phạt).
            if mode == "credit" and rc.get("excess_return_qty"):
                r = start_row + count
                c2 = ws.cell(r, 2)
                c2.value = _("Dư đầu kỳ")
                c4 = ws.cell(r, 4)
                c4.value = _("%s (Chuyển thừa, trả lại)") % block["product_name"]
                c5 = ws.cell(r, 5)
                c5.value = block["uom_name"]
                c6 = ws.cell(r, 6)
                c6.value = rc["excess_return_qty"]
                c8 = ws.cell(r, 8)
                c8.value = block["display_unit_price"]
                for c in (c2, c4, c5, c6, c8):
                    c.font = self._blue_font_like(c)
                count += 1

            # Phần phạt / legacy đối ứng.
            if mode in ("legacy", "mixed") and (
                rc.get("offset_deliveries")
                or rc.get("returns")
                or rc.get("penalty_rows")
                or rc.get("leftover_present")
                or (mode == "legacy" and rc.get("excess_return_qty"))
            ):
                self._write_subheader(
                    ws,
                    start_row + count,
                    _("Số lượng SP đã thuê để đối ứng với phần trả hàng"),
                )
                count += 1
                for od in rc["offset_deliveries"]:
                    r = start_row + count
                    if od.get("is_bob"):
                        ws.cell(r, 2).value = _("Dư đầu kỳ")
                    else:
                        ws.cell(r, 2).value = od["date"].strftime("%d/%m/%Y")
                    ws.cell(r, 4).value = block["product_name"]
                    ws.cell(r, 5).value = block["uom_name"]
                    ws.cell(r, 6).value = od["qty"]
                    ws.cell(r, 8).value = block["display_unit_price"]
                    count += 1
                if rc.get("excess_return_qty") and mode != "credit":
                    r = start_row + count
                    c2 = ws.cell(r, 2)
                    c2.value = _("Dư đầu kỳ")
                    c4 = ws.cell(r, 4)
                    c4.value = _("%s (Chuyển thừa, trả lại)") % block["product_name"]
                    c5 = ws.cell(r, 5)
                    c5.value = block["uom_name"]
                    c6 = ws.cell(r, 6)
                    c6.value = rc["excess_return_qty"]
                    c8 = ws.cell(r, 8)
                    c8.value = block["display_unit_price"]
                    for c in (c2, c4, c5, c6, c8):
                        c.font = self._blue_font_like(c)
                    count += 1
                for ret in rc["returns"]:
                    r = start_row + count
                    c_date = ws.cell(r, 2)
                    c_date.value = ret["date"].strftime("%d/%m/%Y")
                    c_name = ws.cell(r, 4)
                    c_name.value = _("%s (trả hàng)") % block["product_name"]
                    c_uom = ws.cell(r, 5)
                    c_uom.value = block["uom_name"]
                    c_qty = ws.cell(r, 6)
                    c_qty.value = -ret["qty"]
                    c_price = ws.cell(r, 8)
                    c_price.value = block["display_unit_price"]
                    for c in (c_date, c_name, c_uom, c_qty, c_price):
                        c.font = self._red_font_like(c)
                    count += 1

                if rc.get("leftover_present") or rc.get("penalty_rows") or rc.get("returned_rows"):
                    self._write_subheader(ws, start_row + count, _("Đối ứng sản phẩm trả"))
                    count += 1
                    for line in rc.get("leftover_present") or []:
                        content = _("%(name)s (dư từ lô %(d)s)") % {
                            "name": line["product_name"],
                            "d": line["deliver_date"].strftime("%d/%m/%Y"),
                        }
                        count = self._write_billing_line(
                            ws,
                            start_row,
                            count,
                            line,
                            content,
                            use_excel_formulas=use_excel_formulas,
                            klct_meta=None,
                            month_day_dim=month_day_dim,
                        )
                    for agg in rc.get("returned_rows") or []:
                        count = self._write_aggregated_return_row(
                            ws,
                            start_row,
                            count,
                            agg,
                            block,
                            self._return_row_content(block, agg),
                            use_excel_formulas=use_excel_formulas,
                            month_day_dim=month_day_dim,
                        )
                    for agg in rc.get("penalty_rows") or []:
                        count = self._write_aggregated_return_row(
                            ws,
                            start_row,
                            count,
                            agg,
                            block,
                            self._return_row_content(block, agg, penalty=True),
                            use_excel_formulas=use_excel_formulas,
                            month_day_dim=month_day_dim,
                        )

        # Chuyển thừa (chuyển dư) — KHÔNG tính tiền, chỉ hiển thị để dễ quản lý.
        if block.get("excess_qty"):
            r = start_row + count
            ws.cell(r, 4).value = _("Chuyển thừa đang giữ (không tính tiền)")
            ws.cell(r, 4).font = Font(italic=True)
            ws.cell(r, 5).value = block["uom_name"]
            ws.cell(r, 6).value = block["excess_qty"]
            count += 1

        # Dòng Cộng: tổng SL đang thuê cuối kỳ (KHÔNG ghi cột Thành tiền để không lẫn vào tổng).
        r = start_row + count
        ws.cell(r, 4).value = _("Cộng (đang thuê cuối kỳ)")
        ws.cell(r, 4).font = Font(bold=True)
        ws.cell(r, 6).value = block["present_total_qty"]
        ws.cell(r, 6).font = Font(bold=True)
        self._fill_row(ws, r, 2, 9, self._FILL_YELLOW)
        count += 1
        return count

    def _write_credit_return_row(
        self,
        ws,
        start_row,
        count,
        cred,
        block,
        use_excel_formulas=False,
        klct_meta=None,
        month_day_dim=0,
    ):
        """Dòng trả không phạt: trừ tiền từ ngày trả → cuối kỳ (qty & thành tiền âm)."""
        r = start_row + count
        content = _("%s (trả hàng)") % block["product_name"]
        if use_excel_formulas and cred.get("start_date") and cred.get("end_date"):
            self._write_excel_date(ws.cell(r, 2), cred["start_date"])
            self._write_excel_date(ws.cell(r, 3), cred["end_date"])
        else:
            if cred.get("start_date"):
                ws.cell(r, 2).value = cred["start_date"].strftime("%d/%m/%Y")
            if cred.get("end_date"):
                ws.cell(r, 3).value = cred["end_date"].strftime("%d/%m/%Y")
        c_name = ws.cell(r, 4)
        c_name.value = content
        c_uom = ws.cell(r, 5)
        c_uom.value = block["uom_name"]
        c_qty = ws.cell(r, 6)
        qty_formula = None
        signed_qty = -float(cred.get("qty") or 0)
        if use_excel_formulas and klct_meta:
            qty_formula = self._klct_qty_formula_for_line(klct_meta, {
                "tmpl_id": block.get("tmpl_id"),
                "qty": signed_qty,
                "deliver_date": cred.get("return_date") or cred.get("date"),
            })
        c_qty.value = qty_formula if qty_formula else signed_qty
        c_days = ws.cell(r, 7)
        c_price = ws.cell(r, 8)
        c_amt = ws.cell(r, 9)
        if use_excel_formulas and cred.get("start_date") and cred.get("end_date"):
            include = bool(cred.get("include_start_day"))
            holiday_days = self._holiday_days_for_billing_span(
                cred["start_date"], cred["end_date"], include
            )
            c_days.value = self._rental_days_excel_formula(r, include, holiday_days)
            if self.rental_billing_mode == "month" and month_day_dim:
                c_price.value = float(cred.get("unit_price") or 0) * int(month_day_dim)
                c_amt.value = f"=F{r}*G{r}*H{r}/{int(month_day_dim)}"
            else:
                c_price.value = cred.get("display_unit_price", cred.get("unit_price"))
                c_amt.value = f"=F{r}*G{r}*H{r}"
        else:
            c_days.value = cred["rental_days"]
            c_price.value = cred.get("display_unit_price", cred.get("unit_price"))
            c_amt.value = cred["total_amount"]
        for c in (ws.cell(r, 2), ws.cell(r, 3), c_name, c_uom, c_qty, c_days, c_price, c_amt):
            c.font = self._red_font_like(c)
        return count + 1
    def _rental_invoice_xlsx_write_table(
        self,
        ws,
        blocks,
        fee_lines=None,
        use_excel_formulas=False,
        klct_meta=None,
        month_day_dim=0,
        start_row=None,
    ):
        """Mỗi sản phẩm là một khối: đơn thuê bình thường + tính toán trả hàng tối thiểu."""
        if not start_row or start_row < 2:
            start_row = self.env["rental.template"]._DATA_START_ROW_DEFAULT
        max_row = 200
        count = 0
        if self.rental_billing_mode == "month":
            ws.cell(start_row - 1, 8).value = "Đơn giá thuê/\n1 tháng (chưa VAT)"
        else:
            ws.cell(start_row - 1, 8).value = "Đơn giá thuê/\n1 ngày (chưa VAT)"

        for block in blocks or []:
            count = self._write_product_payment_block(
                ws,
                block,
                start_row,
                count,
                use_excel_formulas=use_excel_formulas,
                klct_meta=klct_meta,
                month_day_dim=month_day_dim,
            )

        if fee_lines:
            ws.cell(start_row + count, 4).value = "Phí vận chuyển"
            ws.cell(start_row + count, 4).font = Font(bold=True)
            count += 1
            r = start_row + count
            total_fee = sum((fl.get("amount") or 0.0) for fl in fee_lines)
            n_trips = len(fee_lines)
            ws.cell(r, 4).value = _("Phí vận chuyển (%(n)s chuyến)") % {"n": n_trips}
            ws.cell(r, 9).value = total_fee
            count += 1

        for r in range(start_row + count, start_row + max_row):
            ws.row_dimensions[r].hidden = True

    def _rental_invoice_xlsx_write_fee_detail_sheet(self, wb, fee_lines):
        """Sheet 'Chi tiết phí VC' — ngày / mã phiếu / loại / số tiền (không phá layout HSTT)."""
        if not fee_lines:
            return None
        title = "Chi tiết phí VC"
        # Drop existing sheet with same title if regenerating into a copied workbook.
        existing = self._rental_invoice_get_sheet(wb, title)
        if existing is not None:
            wb.remove(existing)
        ws = wb.create_sheet(title)
        headers = [_("Ngày"), _("Mã phiếu"), _("Loại"), _("Số tiền")]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(1, col)
            cell.value = header
            cell.font = Font(bold=True)
        row = 2
        for fee_line in fee_lines:
            type_label = _("Nhập") if fee_line.get("type") in ("return", "compensation") else _("Xuất")
            fee_date = fee_line.get("date")
            ws.cell(row, 1).value = fee_date.strftime("%d/%m/%Y") if fee_date else ""
            ws.cell(row, 2).value = fee_line.get("code") or ""
            ws.cell(row, 3).value = type_label
            ws.cell(row, 4).value = fee_line.get("amount") or 0.0
            row += 1
        ws.cell(row, 3).value = _("Tổng")
        ws.cell(row, 3).font = Font(bold=True)
        ws.cell(row, 4).value = sum((fl.get("amount") or 0.0) for fl in fee_lines)
        ws.cell(row, 4).font = Font(bold=True)
        return ws

    @staticmethod
    def _rental_invoice_get_sheet(wb, name):
        """Tìm sheet theo tên (không phân biệt hoa/thường, bỏ khoảng trắng đầu/cuối)."""
        target = (name or "").strip().lower()
        for sheet in wb.worksheets:
            if (sheet.title or "").strip().lower() == target:
                return sheet
        return None

    def _rental_invoice_xlsx_write_compensation_sheet(self, ws, compensation_lines):
        """Ghi bảng tiền đền bù (mất/hỏng) sang sheet 'GT đền bù' + tổng cộng riêng."""
        from ..helper.xlsx_template_utils import replace_placeholders_in_sheet

        start_row = 13
        max_row = 200
        count = 0
        for comp_line in compensation_lines:
            r = start_row + count
            ws.cell(r, 2).value = comp_line["date"].strftime("%d/%m/%Y")
            ws.cell(r, 4).value = _("%(name)s - đền bù %(ratio)s%% (mất/hỏng)") % {
                "name": comp_line.get("product_name") or "",
                "ratio": comp_line.get("damage_ratio") or 0,
            }
            ws.cell(r, 5).value = comp_line.get("uom_name") or ""
            qty = comp_line.get("qty") or 0
            amount = comp_line.get("amount") or 0.0
            ws.cell(r, 6).value = qty
            ws.cell(r, 8).value = (amount / qty) if qty else amount
            ws.cell(r, 9).value = amount
            count += 1

        for r in range(start_row + count, start_row + max_row):
            ws.row_dimensions[r].hidden = True

        subtotal = sum((c.get("amount") or 0.0) for c in compensation_lines)
        vat = round(subtotal * 0.08)
        total_after_tax = int(round(subtotal + vat))
        amount_words = self.env["amount_to_text.vi"].vn_amount_to_text(total_after_tax)
        replace_placeholders_in_sheet(
            ws,
            {
                "{{total_after_tax_string}}": amount_words,
                "{{subtotal_before_tax}}": int(round(subtotal)),
                "{{vat_amount}}": int(round(vat)),
                "{{total_after_tax}}": total_after_tax,
            },
        )
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if not isinstance(val, str):
                    continue
                if "Bằng chữ" in val or "bằng chữ" in val:
                    cell.value = f"(Bằng chữ: {amount_words}/.)"

    def _build_rental_payment_xlsx_buffer(
        self, start_date, end_date, transport_fee_until_date=None, fee_invoice=None
    ):
        """Bảng thanh toán gộp theo mẫu SP (cùng file gắn Business XLSX trên hóa đơn)."""
        self.ensure_one()
        bob_map, map_map = rcs.calc_rental_payment_table_grouped_by_template(
            self.env, self, start_date, end_date
        )
        blocks = rcs.calc_rental_payment_blocks_by_template(
            self.env, self, start_date, end_date
        )
        fee_lines = self._rental_period_transport_fee_lines(
            start_date,
            end_date,
            transport_fee_until_date=transport_fee_until_date,
            fee_invoice=fee_invoice,
        )
        compensation_lines = self._rental_period_compensation_lines(start_date, end_date)
        wb, data_start_row = self._rental_invoice_xlsx_load_workbook()
        # Thay placeholder chung (công ty, đại diện, ngày, chức vụ người xác nhận...) cho
        # mọi sheet, gồm 'GT thuê' và 'GT đền bù'.
        for sheet in wb.worksheets:
            self._rental_invoice_xlsx_apply_placeholders(sheet, start_date, end_date)

        # Sheet 'GT thuê' (active): tiền thuê + phí vận chuyển (không gồm đền bù).
        ws = wb.active
        self._rental_invoice_xlsx_write_table(
            ws, blocks, fee_lines, start_row=data_start_row
        )
        self._rental_invoice_xlsx_apply_payment_totals(ws, bob_map, map_map, fee_lines)

        # Sheet 'GT đền bù': chỉ ghi khi có đền bù, ngược lại bỏ sheet cho gọn.
        comp_ws = self._rental_invoice_get_sheet(wb, "GT đền bù")
        if comp_ws is not None:
            if compensation_lines:
                self._rental_invoice_xlsx_write_compensation_sheet(comp_ws, compensation_lines)
            elif len(wb.worksheets) > 1:
                wb.remove(comp_ws)

        self._rental_invoice_xlsx_write_fee_detail_sheet(wb, fee_lines)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    def action_export_invoice_excel(self, start_date, end_date, transport_fee_until_date=None):
        self.ensure_one()
        buffer = self._build_rental_payment_xlsx_buffer(
            start_date, end_date, transport_fee_until_date=transport_fee_until_date
        )
        filename = f"HSTT {end_date.strftime('%m-%Y')} - {self.code}"
        filename_ascii = quote(filename)
        att_id = self.action_create_rental_invoice(
            start_date,
            end_date,
            excel_buffer=buffer,
            file_name=filename_ascii,
            transport_fee_until_date=transport_fee_until_date,
        )
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{att_id}?download=1",
            "target": "self",
        }

    def _build_klct_hstt_xlsx_buffer(
        self, start_date, end_date, transport_fee_until_date=None, fee_invoice=None
    ):
        """One workbook: KLCT + HSTT (+ optional sheets) + ĐCCN with Excel formulas.

        Returns (buffer, hstt_subtotal) where hstt_subtotal matches the HSTT sheet
        footer (rent + transport fees, before VAT; excludes compensation sheet).
        """
        self.ensure_one()
        from ..helper.transport_matrix_export import build_transport_matrix_into_workbook

        period_label = end_date.strftime("%m-%Y")
        klct_title = f"KLCT {period_label}"
        hstt_title = f"HSTT {period_label}"

        wb, klct_meta = build_transport_matrix_into_workbook(
            self.env, self, start_date, end_date
        )
        wb.active.title = klct_title
        klct_meta["sheet_title"] = klct_title
        # Drop any extra sheets from the volume template (keep KLCT only).
        for sheet in list(wb.worksheets):
            if sheet.title != klct_title:
                wb.remove(sheet)

        bob_map, map_map = rcs.calc_rental_payment_table_grouped_by_template(
            self.env, self, start_date, end_date
        )
        blocks = rcs.calc_rental_payment_blocks_by_template(
            self.env, self, start_date, end_date
        )
        fee_lines = self._rental_period_transport_fee_lines(
            start_date,
            end_date,
            transport_fee_until_date=transport_fee_until_date,
            fee_invoice=fee_invoice,
        )
        compensation_lines = self._rental_period_compensation_lines(start_date, end_date)
        month_day_dim = rcs.month_day_basis(self, end_date) if self.rental_billing_mode == "month" else 0
        hstt_subtotal = self._payment_table_subtotal(bob_map, map_map, fee_lines)

        inv_wb, data_start_row = self._rental_invoice_xlsx_load_workbook()
        for sheet in inv_wb.worksheets:
            self._rental_invoice_xlsx_apply_placeholders(sheet, start_date, end_date)

        hstt_src = inv_wb.active
        self._rental_invoice_xlsx_write_table(
            hstt_src,
            blocks,
            fee_lines,
            use_excel_formulas=True,
            klct_meta=klct_meta,
            month_day_dim=month_day_dim,
            start_row=data_start_row,
        )
        self._rental_invoice_xlsx_apply_payment_totals(hstt_src, bob_map, map_map, fee_lines)
        self._copy_xlsx_sheet(hstt_src, wb, hstt_title)

        # Optional compensation sheet as 3rd sheet (same rule as standalone HSTT export).
        comp_ws = self._rental_invoice_get_sheet(inv_wb, "GT đền bù")
        if comp_ws is not None and compensation_lines:
            self._rental_invoice_xlsx_write_compensation_sheet(comp_ws, compensation_lines)
            self._copy_xlsx_sheet(comp_ws, wb, "GT đền bù")

        self._rental_invoice_xlsx_write_fee_detail_sheet(wb, fee_lines)

        dccn_title = f"ĐCCN {period_label}"
        self._build_debt_confirmation_into_workbook(
            wb, start_date, end_date, sheet_title=dccn_title
        )

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer, hstt_subtotal

    def _get_or_create_hstt_sale_tax(self):
        """Return the company's 8% sale tax used by HSTT total invoices."""
        self.ensure_one()
        tax = self.env["account.tax"].search([
            ("company_id", "=", self.company_id.id),
            ("type_tax_use", "=", "sale"),
            ("amount", "=", 8.0),
            ("amount_type", "=", "percent"),
            ("active", "=", True),
        ], limit=1)
        if not tax:
            tax = self.env["account.tax"].create({
                "name": "8%",
                "amount": 8.0,
                "amount_type": "percent",
                "type_tax_use": "sale",
                "company_id": self.company_id.id,
            })
        return tax

    def _get_or_create_hstt_total_product(self, end_date):
        """Service product 'Tổng thanh toán: mm-YYYY' with fixed 8% sale tax."""
        self.ensure_one()
        period = end_date.strftime("%m-%Y")
        name = _("Tổng thanh toán: %s") % period
        tax = self._get_or_create_hstt_sale_tax()
        Product = self.env["product.product"]
        product = Product.search([
            ("name", "=", name),
            ("company_id", "in", [False, self.company_id.id]),
        ], limit=1)
        if product:
            if product.taxes_id != tax:
                product.taxes_id = [(6, 0, tax.ids)]
            return product

        # Prefer income account from a contract line product; else journal default.
        income_account = False
        for cl in self.rental_contract_line_ids:
            variant = cl.product_tmpl_id.product_variant_id
            if not variant:
                continue
            income_account = variant._get_product_accounts().get("income")
            if income_account:
                break
        if not income_account:
            journal = self.env["account.journal"].search([
                ("type", "=", "sale"),
                ("company_id", "=", self.company_id.id),
            ], limit=1)
            income_account = journal.default_account_id if journal else False

        vals = {
            "name": name,
            "type": "service",
            "list_price": 0.0,
            "sale_ok": True,
            "purchase_ok": False,
            "company_id": self.company_id.id,
            "taxes_id": [(6, 0, tax.ids)],
        }
        if income_account:
            vals["property_account_income_id"] = income_account.id
        return Product.create(vals)

    def _rental_posted_invoice_same_period(self, start_date, end_date):
        """Return posted rental invoices for this contract covering the same period."""
        self.ensure_one()
        return self.env["account.move"].search([
            ("rental_contract_id", "=", self.id),
            ("state", "=", "posted"),
            ("move_type", "in", ("out_invoice", "out_refund")),
            ("rental_start_date", "=", start_date),
            ("rental_end_date", "=", end_date),
        ])

    def _attach_business_xlsx_to_move(self, move, excel_buffer, file_name):
        """Create/replace business XLSX attachment on a rental invoice."""
        self.ensure_one()
        name = f"{file_name}.xlsx" if not str(file_name).endswith(".xlsx") else file_name
        datas = base64.b64encode(excel_buffer.getvalue())
        mimetype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if move.business_xlsx_attachment_id:
            move.business_xlsx_attachment_id.write({
                "name": name,
                "datas": datas,
                "mimetype": mimetype,
            })
            return move.business_xlsx_attachment_id.id
        attachment = self.env["ir.attachment"].create({
            "name": name,
            "res_model": "account.move",
            "res_id": move.id,
            "type": "binary",
            "datas": datas,
            "mimetype": mimetype,
        })
        move.business_xlsx_attachment_id = attachment.id
        return attachment.id

    def _create_rental_invoice_from_hstt_total(
        self,
        start_date,
        end_date,
        subtotal,
        transport_fee_until_date=None,
        fee_lines=None,
        post=True,
    ):
        """Create out_invoice with a single line matching the HSTT Excel subtotal.

        When ``post`` is True (default), posts immediately so AR / ĐCCN see the period.
        """
        self.ensure_one()
        partner = self.a_party
        product = self._get_or_create_hstt_total_product(end_date)

        journal = self.env["account.journal"].search([
            ("type", "=", "sale"),
            ("company_id", "=", self.company_id.id),
        ], limit=1)
        if not journal:
            raise UserError(_("No Sale journal found for %s") % self.company_id.display_name)

        fpos = partner.property_account_position_id
        accounts = product._get_product_accounts()
        income_account = accounts.get("income") or journal.default_account_id
        if not income_account:
            raise UserError(_("No income account set for %s") % product.display_name)
        if fpos:
            income_account = fpos.map_account(income_account)
        # HSTT total invoices are explicitly configured at 8%; do not let a stale
        # product tax or partner fiscal-position mapping replace it with 10%.
        taxes = self._get_or_create_hstt_sale_tax()

        # Cutoff trống trên wizard = không tính phí; lưu False. None từ caller cũ → end_date.
        fee_until = transport_fee_until_date
        if fee_until is None and fee_lines:
            fee_until = end_date

        move = self.env["account.move"].create({
            "rental_start_date": start_date,
            "rental_end_date": end_date,
            "transport_fee_until_date": fee_until or False,
            "move_type": "out_invoice",
            "partner_id": partner.id,
            "rental_contract_id": self.id,
            "invoice_date": fields.Date.context_today(self),
            "journal_id": journal.id,
            "invoice_line_ids": [(0, 0, {
                "product_id": product.id,
                "name": product.display_name,
                "quantity": 1,
                "price_unit": subtotal,
                "account_id": income_account.id,
                "tax_ids": [(6, 0, taxes.ids)],
            })],
        })

        if fee_lines and fee_until:
            self._mark_transport_fees_billed(fee_lines, move, fee_until)

        if post:
            move.action_post()
        return move

    def action_export_klct_hstt_excel(self, start_date, end_date, transport_fee_until_date=None):
        """Create volume matrix + posted invoice (1 line = HSTT total) + download combined XLSX."""
        self.ensure_one()
        Matrix = self.env["rental.transport.matrix"]
        overlap_domain = [
            ("rental_contract_id", "=", self.id),
            ("start_date", "<=", end_date),
            ("end_date", ">=", start_date),
        ]
        if Matrix.search_count(overlap_domain):
            wiz = self.env["rental.transport.matrix.overlap.wizard"].create({
                "rental_contract_id": self.id,
            })
            return {
                "type": "ir.actions.act_window",
                "name": _("Trùng khoảng thời gian"),
                "res_model": "rental.transport.matrix.overlap.wizard",
                "view_mode": "form",
                "target": "new",
                "res_id": wiz.id,
            }

        existing = self._rental_posted_invoice_same_period(start_date, end_date)
        if existing:
            raise UserError(
                _(
                    "Đã có hóa đơn đã đăng sổ cho kỳ %(start)s → %(end)s "
                    "(%(moves)s). Hãy hủy hóa đơn cũ rồi xuất lại."
                )
                % {
                    "start": start_date.strftime("%d/%m/%Y"),
                    "end": end_date.strftime("%d/%m/%Y"),
                    "moves": ", ".join(existing.mapped("name")),
                }
            )

        Matrix.create({
            "rental_contract_id": self.id,
            "start_date": start_date,
            "end_date": end_date,
            "name": f"Confirmation table {self.code}: {start_date.strftime('%d/%m/%Y')} → {end_date.strftime('%d/%m/%Y')}",
        })

        fee_lines = self._rental_period_transport_fee_lines(
            start_date, end_date, transport_fee_until_date=transport_fee_until_date
        )
        bob_map, map_map = rcs.calc_rental_payment_table_grouped_by_template(
            self.env, self, start_date, end_date
        )
        hstt_subtotal = self._payment_table_subtotal(bob_map, map_map, fee_lines)
        # Match Excel HSTT footer {{subtotal_before_tax}} = int(round(subtotal)).
        move = self._create_rental_invoice_from_hstt_total(
            start_date,
            end_date,
            subtotal=float(int(round(hstt_subtotal))),
            transport_fee_until_date=transport_fee_until_date,
            fee_lines=fee_lines,
            post=True,
        )
        # Build after post so ĐCCN sheet includes this period's residual.
        buffer, _subtotal = self._build_klct_hstt_xlsx_buffer(
            start_date,
            end_date,
            transport_fee_until_date=transport_fee_until_date,
            fee_invoice=move,
        )
        filename = f"KLCT-HSTT {end_date.strftime('%m-%Y')} - {self.code}"
        att_id = self._attach_business_xlsx_to_move(move, buffer, filename)
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{att_id}?download=1",
            "target": "self",
        }

    def action_open_account_moves(self):
        self.ensure_one()
        # Build action in code only — do not ref() account actions (restricted to Admin/Settings).
        action = {
            'type': 'ir.actions.act_window',
            'name': _('Invoices / Journal Entries'),
            'res_model': 'account.move',
            'view_mode': 'tree,form',
            'target': 'current',
            'domain': [('id', 'in', self.account_move_ids.ids)],
            # Optional: start directly on the first move’s form
            # 'res_id': self.account_move_ids[:1].id if self.account_move_ids else False,
            # 'views': [(False, 'tree'), (self.env.ref('account.view_move_form').id, 'form')],
        }
        return action

    def action_create_rental_invoice(
        self, start_date, end_date, excel_buffer=False, file_name='', transport_fee_until_date=None
    ):
        self.ensure_one()
        partner = self.a_party
        bob_map_product_id_2_line, map_product_id_2_line = rcs.calc_rental_contract_invoice(self.env, self,
                                                                                            start_date, end_date)
        product_items = []
        bob_map_product_id_2_line |= map_product_id_2_line
        for product_id in bob_map_product_id_2_line:
            product_info = bob_map_product_id_2_line[product_id]
            product_items.append({
                'product_id': product_id,
                'qty': product_info['invoice_qty'],
                'price_unit': product_info['unit_price'],
            })

        journal = self.env['account.journal'].search([
            ('type', '=', 'sale'),
            ('company_id', '=', self.company_id.id)
        ], limit=1)
        if not journal:
            raise UserError(_("No Sale journal found for %s") % self.company_id.display_name)

        fpos = partner.property_account_position_id
        invoice_lines = []
        Product = self.env['product.product']
        for item in product_items:
            product = Product.browse(item['product_id'])
            if not product.exists():
                raise UserError(_("Product not found (ID %s)") % item['product_id'])

            qty = item.get('qty', 0.0)
            # if qty <= 0:
            #     raise UserError(_("Quantity must be positive for %s") % product.display_name)
            price_unit = item.get('price_unit', product.lst_price)
            discount = item.get('discount', 0.0)
            # Income account
            accounts = product._get_product_accounts()
            income_account = accounts.get('income')  # Pull code from OCA git and update account_payment?
            if not income_account:
                raise UserError(_("No income account set for %s") % product.display_name)
            if fpos:
                income_account = fpos.map_account(income_account)

            # Taxes
            taxes = product.taxes_id.filtered(lambda t: t.company_id == self.company_id)
            if fpos:
                taxes = fpos.map_tax(taxes)

            invoice_lines.append((0, 0, {
                'product_id': product.id,
                'name': product.display_name,
                'quantity': qty,
                'price_unit': price_unit,
                'discount': discount,
                'account_id': income_account.id,
                'tax_ids': [(6, 0, taxes.ids)],
            }))

        # Dòng tiền đền bù (mất/hỏng) — tách riêng cho từng phiếu đền bù trong kỳ.
        for comp_line in self._rental_period_compensation_lines(start_date, end_date):
            product = Product.browse(comp_line['product_id'])
            if not product.exists():
                continue
            accounts = product._get_product_accounts()
            income_account = accounts.get('income')
            if not income_account:
                raise UserError(_("No income account set for %s") % product.display_name)
            if fpos:
                income_account = fpos.map_account(income_account)
            taxes = product.taxes_id.filtered(lambda t: t.company_id == self.company_id)
            if fpos:
                taxes = fpos.map_tax(taxes)
            comp_qty = comp_line.get('qty') or 0
            comp_amount = comp_line.get('amount') or 0.0
            comp_price_unit = (comp_amount / comp_qty) if comp_qty else comp_amount
            invoice_lines.append((0, 0, {
                'product_id': product.id,
                'name': _("%(name)s - đền bù %(ratio)s%% (mất/hỏng)") % {
                    "name": product.display_name,
                    "ratio": comp_line.get('damage_ratio') or 0,
                },
                'quantity': comp_qty or 1,
                'price_unit': comp_price_unit,
                'account_id': income_account.id,
                'tax_ids': [(6, 0, taxes.ids)],
            }))

        # Phí vận chuyển — đưa vào hóa đơn để khớp với bảng thanh toán xuất ra (Excel cộng
        # phí vận chuyển vào tổng tiền). Trước đây hóa đơn thiếu phần này nên bị lệch.
        fee_lines = self._rental_period_transport_fee_lines(
            start_date, end_date, transport_fee_until_date=transport_fee_until_date
        )
        if fee_lines:
            # Dùng chung tài khoản/thuế với dòng sản phẩm (thường VAT 8%) để tổng khớp Excel.
            fee_account = False
            fee_taxes = self.env['account.tax']
            if invoice_lines:
                first_vals = invoice_lines[0][2]
                fee_account = self.env['account.account'].browse(first_vals['account_id'])
                fee_tax_ids = first_vals.get('tax_ids') or []
                if fee_tax_ids and fee_tax_ids[0][2]:
                    fee_taxes = self.env['account.tax'].browse(fee_tax_ids[0][2])
            if not fee_account:
                fee_account = journal.default_account_id
            total_fee = sum((fl.get("amount") or 0.0) for fl in fee_lines)
            invoice_lines.append((0, 0, {
                'name': _("Phí vận chuyển (%(n)s chuyến)") % {"n": len(fee_lines)},
                'quantity': 1,
                'price_unit': total_fee,
                'account_id': fee_account.id,
                'tax_ids': [(6, 0, fee_taxes.ids)],
            }))

        fee_until = transport_fee_until_date
        if fee_until is None and fee_lines:
            fee_until = end_date

        move = self.env['account.move'].create({
            'rental_start_date': start_date,
            'rental_end_date': end_date,
            'transport_fee_until_date': fee_until or False,
            'move_type': 'out_invoice',
            'partner_id': partner.id,
            'rental_contract_id': self.id,
            'invoice_date': fields.Date.context_today(self),
            'journal_id': journal.id,
            'invoice_line_ids': invoice_lines,
        })

        if fee_lines and fee_until:
            self._mark_transport_fees_billed(fee_lines, move, fee_until)

        attachment = self.env["ir.attachment"].create({
            "name": f"{file_name}.xlsx",
            "res_model": "account.move",
            "res_id": move.id,
            "type": "binary",
            "datas": base64.b64encode(excel_buffer.getvalue()),
            "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        })
        move.business_xlsx_attachment_id = attachment.id
        return attachment.id


    def action_create_invoice(self, start_date, end_date):  # This is manually invoice model, i want to use default invoice of Odoo
        for rec in self:
            rental_invoice = self.env['rental.invoice'].create({
                'name': f"{rec.a_party.name}-{start_date.strftime('%d/%m/%Y')}->{end_date.strftime('%d/%m/%Y')}",
                'rental_contract_id': rec.id,
                'start_date': start_date,
                'end_date': end_date,
            })

            map_product_and_date_to_line = {}
            ratios = rcs.contract_line_ratios_by_template(rec)
            # The key will be product_id and the start_date.
            # If more lines have the same product_id and start_date ==> sum them
            for line in rec.rr_transport_line_ids:
                if line.transport_id.state != "done":
                    continue
                if not line.billable_qty:
                    continue
                invoice_date = line.start_rental_or_return_date
                if invoice_date < start_date:
                    # We start invoice from start_date
                    invoice_date = start_date
                if invoice_date <= end_date:
                    key = f"{line.product_id.id}_{invoice_date.strftime('%Y%m%d')}"
                    if key not in map_product_and_date_to_line:
                        tmpl_id = line.product_id.product_tmpl_id.id
                        unit_price = rcs.unit_price_for_transport_line(
                            rec,
                            line.product_id,
                            ratios.get(tmpl_id, 1.0),
                            end_date,
                        )
                        map_product_and_date_to_line[key] = {
                            'start_date': invoice_date,
                            'product_id': line.product_id.id,
                            'qty': line.billable_qty,
                            'unit_price': unit_price,
                        }
                    else:
                        map_product_and_date_to_line[key]['qty'] += line.billable_qty

            if map_product_and_date_to_line:
                for key in map_product_and_date_to_line:
                    line = map_product_and_date_to_line[key]
                    self.env['rental.invoice.line'].create({
                        'rental_invoice_id': rental_invoice.id,
                        'start_date': line['start_date'],
                        'end_date': end_date,
                        'product_id': line['product_id'],
                        'qty': line['qty'],
                        'unit_price': line['unit_price'],
                    })

    def _collect_debt_confirmation_data(self, start_date, end_date):
        """Aggregate posted rental invoices into beginning / in-period debt buckets.

        Payment-in-period uses amount_total - amount_residual (lifetime paid snapshot
        on invoices in the period), not payment journal dates.
        """
        beginning_debit = 0.0
        total_amount_in_period = 0.0
        total_paid_in_period = 0.0
        total_remain = 0.0
        beginning_moves = self.env["account.move"]
        in_period_moves = self.env["account.move"]
        customer_id = False

        for rec in self:
            if not customer_id:
                customer_id = rec.a_party.id
            elif customer_id != rec.a_party.id:
                raise UserError(_("Only allow calculate debt for one customer each time"))
            for inv in rec.account_move_ids:
                if inv.state in ("draft", "cancel") or not inv.rental_end_date:
                    continue
                if inv.move_type not in ("out_invoice", "out_refund"):
                    continue
                if inv.rental_end_date < start_date:
                    if inv.amount_residual > 0:
                        beginning_moves |= inv
                        beginning_debit += inv.amount_residual
                        total_remain += inv.amount_residual
                elif inv.rental_end_date <= end_date:
                    in_period_moves |= inv
                    total_amount_in_period += inv.amount_total
                    total_paid_in_period += inv.amount_total - inv.amount_residual
                    total_remain += inv.amount_residual

        return {
            "customer_id": customer_id,
            "beginning_debit": beginning_debit,
            "total_amount_in_period": total_amount_in_period,
            "total_paid_in_period": total_paid_in_period,
            "total_remain": total_remain,
            "beginning_moves": beginning_moves.sorted(
                key=lambda m: (m.rental_end_date or m.invoice_date or date.min, m.id)
            ),
            "in_period_moves": in_period_moves.sorted(
                key=lambda m: (m.rental_end_date or m.invoice_date or date.min, m.id)
            ),
        }

    def _debt_confirmation_amount_to_text(self, amount):
        try:
            AmountVi = self.env["amount_to_text.vi"]
        except KeyError:
            return ""
        if hasattr(AmountVi, "vn_amount_to_text"):
            return AmountVi.vn_amount_to_text(amount)
        return ""

    def _write_debt_confirmation_sheet(self, ws, start_date, end_date, data):
        """Fill ĐCCN template: header placeholders + summary totals + invoice list."""
        self.ensure_one()
        total_remain = data["total_remain"]
        replacements = {
            "{{a_company}}": self.a_party.parent_id.name or "",
            "{{a_address}}": self.a_address or "",
            "{{a_representative}}": self.a_name or "",
            "{{a_function}}": self.a_function or "",
            "{{b_company}}": self.b_party.parent_id.name or "",
            "{{b_address}}": self.b_address or "",
            "{{b_representative}}": self.b_name or "",
            "{{b_function}}": self.b_function or "",
            "{{construction_work_project}}": self.construction_work_project_id.name or "",
            "{{construction_work_name}}": self.construction_work_id.name or "",
            "{{construction_work_address}}": self.construction_work_address or "",
            "{{start_date}}": start_date.strftime("%d/%m/%Y"),
            "{{end_date}}": end_date.strftime("%d/%m/%Y"),
            "{{total_remain}}": f"{total_remain:,.0f}" if total_remain else "0",
            "{{total_remain_string}}": self._debt_confirmation_amount_to_text(total_remain),
        }
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    for key, val in replacements.items():
                        if key in cell.value:
                            cell.value = cell.value.replace(key, val)

        # Summary totals on template section rows.
        ws.cell(16, 8).value = data["beginning_debit"]
        ws.cell(17, 8).value = data["total_amount_in_period"]
        ws.cell(17, 6).value = None  # clear stale template SUM cache
        ws.cell(115, 8).value = data["total_paid_in_period"]

        detail_start = 19  # under "Tiền thuê hàng tháng" (row 18)
        payment_row = 115
        max_detail = payment_row - 1
        row = detail_start

        def _write_inv_row(r, inv, label, amount):
            inv_date = inv.invoice_date or inv.rental_end_date
            ws.cell(r, 2).value = inv.name or ""
            ws.cell(r, 3).value = inv_date.strftime("%d/%m/%Y") if inv_date else ""
            period = ""
            if inv.rental_start_date and inv.rental_end_date:
                period = (
                    f"{inv.rental_start_date.strftime('%d/%m/%Y')}"
                    f" → {inv.rental_end_date.strftime('%d/%m/%Y')}"
                )
            ws.cell(r, 4).value = f"{label} {period}".strip()
            ws.cell(r, 8).value = amount

        for inv in data["beginning_moves"]:
            if row >= max_detail:
                break
            _write_inv_row(row, inv, _("Nợ đầu kỳ"), inv.amount_residual)
            row += 1
        for inv in data["in_period_moves"]:
            if row >= max_detail:
                break
            _write_inv_row(row, inv, _("Phát sinh"), inv.amount_total)
            row += 1

        # Hide unused detail rows between last written line and payment section.
        for r in range(row, payment_row):
            ws.row_dimensions[r].hidden = True

    def _load_debt_confirmation_workbook(self):
        """Load company ĐCCN template into a workbook (active sheet filled later)."""
        company = self.company_id if len(self) == 1 else self[:1].company_id
        data, _source = self.env["rental.template"].get_template_bytes(
            company,
            "debt_confirmation_xlsx",
        )
        return load_workbook(io.BytesIO(data))

    def _build_debt_confirmation_into_workbook(
        self, dest_wb, start_date, end_date, sheet_title=None
    ):
        """Fill ĐCCN template and copy as a sheet into ``dest_wb``."""
        self.ensure_one()
        data = self._collect_debt_confirmation_data(start_date, end_date)
        src_wb = self._load_debt_confirmation_workbook()
        ws = src_wb.active
        self._write_debt_confirmation_sheet(ws, start_date, end_date, data)
        title = sheet_title or f"ĐCCN {end_date.strftime('%m-%Y')}"
        return self._copy_xlsx_sheet(ws, dest_wb, title)

    def export_debt_confirmation_comparison_table(self, start_date, end_date):
        """Standalone ĐCCN XLSX download (same builder as combined KLCT+HSTT sheet)."""
        data = self._collect_debt_confirmation_data(start_date, end_date)
        if not data["customer_id"]:
            raise UserError(_("No customer on selected contracts."))
        contract = self[:1]
        src_wb = contract._load_debt_confirmation_workbook()
        ws = src_wb.active
        contract._write_debt_confirmation_sheet(ws, start_date, end_date, data)

        buffer = io.BytesIO()
        src_wb.save(buffer)
        buffer.seek(0)

        filename = f"ĐCCN {end_date.strftime('%m-%Y')} - {contract.code}"
        attachment = self.env["ir.attachment"].create({
            "name": f"{filename}.xlsx",
            "res_model": "res.partner",
            "res_id": data["customer_id"],
            "type": "binary",
            "datas": base64.b64encode(buffer.getvalue()),
            "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        })
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{attachment.id}?download=1",
            "target": "self",
        }

    def action_print_quotation_price_table(self):
        """Return the PDF report action."""
        self.ensure_one()
        return self.env.ref('rental.action_report_quotation_price').report_action(self)

    def action_preview_quotation_price_table(self):
        self.ensure_one()
        # Reuse the same report but as HTML; add debug=assets for easier CSS dev.
        return (
            self.env.ref('rental.action_report_quotation_price_html')
            .with_context(preview_mode=True, debug='assets')
            .report_action(self)
        )

    def action_export_quotation_excel(self):
        self.ensure_one()
        url = f'/rental/rental-contract/quotation/{self.id}/download'
        return {
            'type': 'ir.actions.act_url',
            'url': url,
            'target': 'self',  # or 'new' to open in new tab
        }

    @staticmethod
    def _format_months_padded(months):
        """Zero-pad month count for Excel (2 -> '02', 11 -> '11')."""
        return f"{int(months or 0):02d}"

    @staticmethod
    def _months_to_days(months):
        return int(months or 0) * 30

    def _quotation_contract_date_display(self):
        """Vietnamese date: ngày DD tháng MM năm YYYY (zero-pad day/month)."""
        self.ensure_one()
        d = self.contract_date
        if not d:
            return ''
        return f"ngày {d.day:02d} tháng {d.month:02d} năm {d.year}"

    def _quotation_xlsx_placeholder_replacements(self):
        """Placeholders for contract_quotation_xlsx template."""
        self.ensure_one()
        min_months = self.minimum_rental_months or 0
        share_months = self.transport_fee_share_min_months or 0
        return {
            '{{b_company}}': self.b_party.parent_id.name or '',
            '{{b_address}}': self.b_address or '',
            '{{b_representative}}': self.b_name or '',
            '{{b_phone}}': self.b_phone or '',
            '{{b_email}}': self.b_party.email or '',
            '{{today_is}}': date.today().strftime('ngày %d tháng %m năm %Y'),

            '{{a_representative}}': self.a_name or '',
            '{{a_company}}': self.a_party.parent_id.name or '',

            '{{contract_date}}': self._quotation_contract_date_display(),
            '{{construction_work}}': self.construction_work_id.name or '',
            '{{construction_work_project}}': self.construction_work_project_id.name or '',
            '{{construction_work_name}}': self.construction_work_id.name or '',
            '{{construction_work_address}}': self.construction_work_address or '',

            '{{minimum_rental_months}}': self._format_months_padded(min_months),
            '{{minimum_rental_months_to_day}}': str(self._months_to_days(min_months)),
            '{{transport_fee_share_min_months}}': self._format_months_padded(share_months),
            '{{transport_fee_share_min_months_to_day}}': str(self._months_to_days(share_months)),
            '{{prices_include_tax}}': _('đã') if self.prices_include_tax else _('chưa'),
        }

    def action_export_transport_matrix_excel(self, start_date, end_date):
        self.ensure_one()
        params = {
            'start_date': fields.Date.to_string(start_date),
            'end_date': fields.Date.to_string(end_date),
        }
        url = f"/rental/rental-contract/transport-matrix/{self.id}/download?{url_encode(params)}"
        return {
            'type': 'ir.actions.act_url',
            'url': url,
            'target': 'self',  # or 'new' to open in new tab
        }

    def action_create_transport_matrix_record(self, start_date, end_date):
        """
        New flow: persist a matrix record then open it.
        (Replaces direct Excel download.)
        """
        self.ensure_one()
        Matrix = self.env["rental.transport.matrix"]
        overlap_domain = [
            ("rental_contract_id", "=", self.id),
            ("start_date", "<=", end_date),
            ("end_date", ">=", start_date),
        ]
        if Matrix.search_count(overlap_domain):
            wiz = self.env["rental.transport.matrix.overlap.wizard"].create({
                "rental_contract_id": self.id,
            })
            return {
                "type": "ir.actions.act_window",
                "name": _("Trùng khoảng thời gian"),
                "res_model": "rental.transport.matrix.overlap.wizard",
                "view_mode": "form",
                "target": "new",
                "res_id": wiz.id,
            }
        matrix = Matrix.create({
            "rental_contract_id": self.id,
            "start_date": start_date,
            "end_date": end_date,
            "name": f"Confirmation table {self.code}: {start_date.strftime('%d/%m/%Y')} → {end_date.strftime('%d/%m/%Y')}",
        })
        return {
            "type": "ir.actions.act_window",
            "name": _("Confirmation table for rental volume"),
            "res_model": "rental.transport.matrix",
            "view_mode": "form",
            "target": "current",
            "res_id": matrix.id,
        }

    def action_download_rental_contract(self):
        self.ensure_one()
        url = f'/rental/rental-contract/contract/{self.id}/download'
        return {
            'type': 'ir.actions.act_url',
            'url': url,
            'target': 'self',  # or 'new' to open in new tab
        }

    def action_print_rental_contract(self):
        self.ensure_one()
        url = f'/rental/rental-contract/contract/{self.id}/print'
        return {
            'type': 'ir.actions.act_url',
            'url': url,
            'target': 'new',  # 'self' for in this tab or 'new' to open in new tab
        }

    # @api.depends(
    #     'rr_transport_ids.date',
    #     'rr_transport_ids.truck_id',
    #     'rr_transport_ids.transport_line_ids.product_id',
    #     'rr_transport_ids.transport_line_ids.qty',
    # )
    # def _compute_rr_transport_matrix_html(self):
    #     """
    #     Build an HTML table:
    #       - Rows: (date, truck)
    #       - Columns: equipment (product)
    #       - Cell: sum(qty)
    #     """
    #     for rec in self:
    #         # ---- Collect all lines for this contract ----
    #         # Expected structure:
    #         # rr_transport_ids(date, truck_id) -> transport_line_ids(product_id, qty)
    #         rows = []  # (date_str, truck_id) -> dict
    #         products_order = []  # keep stable order as we discover
    #         prod_seen = set()
    #         col_totals = defaultdict(float)
    #
    #         for transport in rec.rr_transport_ids:
    #             start_rental_or_return_date = transport.start_rental_or_return_date and transport.start_rental_or_return_date.strftime('%-m/%-d/%Y')
    #             plate = transport.plate
    #
    #             row = {
    #                 'date': start_rental_or_return_date,
    #                 'truck_name': plate,
    #                 'cells': defaultdict(float),
    #                 'row_total': 0.0,
    #             }
    #             rows.append(row)
    #
    #             for line in transport.transport_line_ids:
    #                 product_id = line.product_id.id
    #                 if product_id not in prod_seen:
    #                     prod_seen.add(product_id)
    #                     products_order.append((product_id, line.product_id.display_name))
    #                 qty = float(line.qty or 0.0)
    #                 row['cells'][product_id] += qty
    #                 row['row_total'] += qty
    #                 col_totals[product_id] += qty
    #
    #         # ---- Sort rows by date then truck name ----
    #         rows = list(rows_keyed.values())
    #         rows.sort(key=lambda r: (r['date'], r['truck_name']))
    #
    #         # ---- Grand total ----
    #         grand_total = sum(col_totals.values())
    #
    #         # ---- Render HTML (bootstrap-like table) ----
    #         # You can tweak styles to match your Excel screenshot.
    #         def esc(x):
    #             return html.escape(str(x) if x is not None else '')
    #
    #         css = """
    # <style>
    # .tbl-matrix { border-collapse: collapse; width: 100%; table-layout: auto; }
    # .tbl-matrix th, .tbl-matrix td { border: 1px solid #ddd; padding: 4px 6px; }
    # .tbl-matrix thead th { background: #f6f6f6; text-align: center; font-weight: 600; }
    # .tbl-matrix tfoot th { background: #f0f0f0; }
    # .text-end { text-align: right; }
    # .text-center { text-align: center; }
    # .sticky-head thead th { position: sticky; top: 0; z-index: 2; }
    # .sticky-first th[rowspan], .sticky-first td:first-child { position: sticky; left: 0; z-index: 1; background: #fff; }
    # </style>
    # """
    #
    #         # Header cells
    #         head_cols_html = "".join(f"<th>{esc(name)}</th>" for pid, name in products_order)
    #
    #         # Body rows
    #         body_rows_html = []
    #         for idx, r in enumerate(rows, start=1):
    #             cells_html = "".join(
    #                 f"<td class='text-end'>{esc(r['cells'].get(pid, 0.0))}</td>"
    #                 for pid, _ in products_order
    #             )
    #             body_rows_html.append(
    #                 f"<tr>"
    #                 f"<td class='text-center'>{idx}</td>"
    #                 f"<td>{esc(r['date'])}</td>"
    #                 f"<td>{esc(r['truck_name'])}</td>"
    #                 f"{cells_html}"
    #                 f"<td class='text-end'><b>{esc(r['row_total'])}</b></td>"
    #                 f"</tr>"
    #             )
    #
    #         # Footer totals
    #         foot_cols_html = "".join(
    #             f"<th class='text-end'>{esc(col_totals.get(pid, 0.0))}</th>"
    #             for pid, _ in products_order
    #         )
    #
    #         html_table = f"""
    # {css}
    # <div style="max-height: calc(100vh - 260px); overflow:auto;">
    #   <table class="tbl-matrix sticky-head sticky-first">
    #     <thead>
    #       <tr>
    #         <th rowspan="2" style="min-width:50px;">STT</th>
    #         <th rowspan="2" style="min-width:120px;">Ngày tháng</th>
    #         <th rowspan="2" style="min-width:140px;">BKS Xe</th>
    #         <th colspan="{len(products_order)}">Chủng loại / Khối lượng (cây/cái/chân/cặp)</th>
    #         <th rowspan="2" style="min-width:110px;">Tổng hàng</th>
    #       </tr>
    #       <tr>
    #         {head_cols_html}
    #       </tr>
    #     </thead>
    #     <tbody>
    #       {''.join(body_rows_html) if body_rows_html else "<tr><td colspan='999'>No data</td></tr>"}
    #     </tbody>
    #     <tfoot>
    #       <tr>
    #         <th colspan="3" class="text-end">Tổng cột</th>
    #         {foot_cols_html}
    #         <th class="text-end">{esc(grand_total)}</th>
    #       </tr>
    #     </tfoot>
    #   </table>
    # </div>
    # """
    #         rec.rr_transport_matrix_html = html_table

    def request_confirm_from_leader(self):
        for rec in self:
            if rec.status in ['new', 'need_fix'] or rec.edit_unlocked:
                rec.with_context(rental_contract_allow_locked_write=True).write({
                    'status': 'need_approve',
                    'edit_unlocked': False,
                })
                rec.message_post(
                    body=_("Yêu cầu xác nhận hợp đồng thuê."),
                    message_type="comment",
                    subtype_xmlid="mail.mt_comment",
                )

    def leader_confirm_rental_contract(self):
        for rec in self:
            if rec.status in ['need_approve']:
                rec.with_context(rental_contract_allow_locked_write=True).write({
                    'status': 'leader_approved',
                    'edit_unlocked': False,
                })

    def customer_confirm_rental_contract(self):
        for rec in self:
            if rec.status in ['new', 'leader_approved']:
                rec.with_context(rental_contract_allow_locked_write=True).write({
                    'status': 'customer_confirmed',
                    'edit_unlocked': False,
                })

    def active_rental_contract(self):
        for rec in self:
            if rec.status in ['new', 'leader_approved', 'customer_confirmed']:
                rec.with_context(rental_contract_allow_locked_write=True).write({
                    'status': 'active',
                    'edit_unlocked': False,
                })

    def finish_rental_contract(self):
        for rec in self:
            if rec.status in ['active']:
                rec.with_context(rental_contract_allow_locked_write=True).write({
                    'status': 'finish',
                    'edit_unlocked': False,
                })

    def reactivate_rental_contract(self):
        for rec in self:
            if rec.status in ['finish']:
                rec.with_context(rental_contract_allow_locked_write=True).write({
                    'status': 'active',
                    'edit_unlocked': False,
                })

    def action_open_reject_confirm_wizard(self):
        self.ensure_one()
        if self.status != "need_approve":
            raise UserError(_("Hợp đồng không ở trạng thái cần xác nhận."))  # noqa: E501
        return {
            "type": "ir.actions.act_window",
            "res_model": "rental.reject.confirm.wizard",
            "view_mode": "form",
            "target": "new",
            "context": {
                "active_model": "rental.contract",
                "active_ids": self.ids,
            },
        }

    def action_unlock_edit(self):
        self.ensure_one()
        if not self.env.user.has_group("rental.group_rental_leader"):
            raise UserError(_("Chỉ Leader mới có quyền mở khóa để sửa."))  # noqa: E501
        if self.status in ("new", "need_fix"):
            # In those states, it is already editable.
            return False
        if self.edit_unlocked:
            return False
        self.with_context(rental_contract_allow_locked_write=True).write({"edit_unlocked": True})
        self.message_post(
            body=_("Leader đã mở khóa hợp đồng để chỉnh sửa."),
            message_type="comment",
            subtype_xmlid="mail.mt_comment",
        )

    def action_lock_edit(self):
        self.ensure_one()
        if not self.edit_unlocked:
            return False
        self.with_context(rental_contract_allow_locked_write=True).write({"edit_unlocked": False})
        self.message_post(
            body=_("Hợp đồng đã được khóa lại."),
            message_type="comment",
            subtype_xmlid="mail.mt_comment",
        )


class RentalContractLine(models.Model):
    _name = "rental.contract.line"
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _description = "Rental Contract Line"
    _order = "sequence, id"

    sequence = fields.Integer(default=10)
    contract_id = fields.Many2one('rental.contract', string="Contract", required=True, ondelete='cascade', index=True)
    company_id = fields.Many2one(related='contract_id.company_id', store=True, readonly=True, index=True)
    company_group_id = fields.Many2one(related='contract_id.company_group_id', store=True, index=True, readonly=True)
    currency_id = fields.Many2one(related='contract_id.currency_id', store=True, readonly=True)
    partner_id = fields.Many2one(related='contract_id.a_party', store=True, readonly=True)

    product_tmpl_id = fields.Many2one(
        'product.template', string="Product", required=True,
        domain=[('sale_ok', '=', True), ('active', '=', True)]
    )
    name = fields.Text(string="Description")
    product_uom_qty = fields.Integer(string="Quantity", default=1, readonly=True)
    price_unit = fields.Float(
        string="Đơn giá thuê / tháng",
        required=True,
        default=0,
        help="Đơn giá thuê theo tháng trên bảng báo giá (cùng đơn vị với list_price / giá tháng catalog).",
    )
    price_unit_day = fields.Float(
        string="Đơn giá thuê / ngày",
        compute="_compute_price_unit_day",
        readonly=True,
        help="Quy đổi từ đơn giá tháng ÷ 30 (cùng quy ước xuất báo giá).",
    )
    standard_price = fields.Float(string="Standard Price")
    compensation_price = fields.Float(string="Compensation Price")
    uom_id = fields.Many2one('uom.uom', 'Unit of Measure')
    minimum_rental_months = fields.Integer(
        string="Kỳ tối thiểu (tháng)",
        default=0,
        help="Ghi đè kỳ thuê tối thiểu riêng cho sản phẩm này. 0 = dùng theo hợp đồng.",
    )
    price_history_ids = fields.One2many(
        'rental.contract.line.price',
        'contract_line_id',
        string="Lịch sử điều chỉnh giá",
    )

    @api.depends('price_unit')
    def _compute_price_unit_day(self):
        for line in self:
            line.price_unit_day = (line.price_unit or 0.0) / 30.0

    @api.constrains('contract_id', 'product_tmpl_id')
    def _check_unique_product_tmpl_per_contract(self):
        for line in self:
            if not line.contract_id or not line.product_tmpl_id:
                continue
            duplicates = self.search_count([
                ('contract_id', '=', line.contract_id.id),
                ('product_tmpl_id', '=', line.product_tmpl_id.id),
                ('id', '!=', line.id),
            ])
            if duplicates:
                raise ValidationError(_(
                    "Sản phẩm «%s» đã có trên bảng báo giá của hợp đồng này."
                ) % (line.product_tmpl_id.display_name,))

    def _check_contract_can_edit_lines(self):
        if self.env.context.get("rental_contract_allow_locked_write"):
            return
        for line in self:
            contract = line.contract_id
            if contract and not contract.can_edit:
                raise UserError(_(
                    "Hợp đồng hiện đang bị khóa, không cho phép sửa bảng báo giá. "
                    "Vui lòng yêu cầu Leader mở khóa."
                ))

    @api.model_create_multi
    def create(self, vals_list):
        lines = super().create(vals_list)
        lines._check_contract_can_edit_lines()
        return lines

    def write(self, vals):
        self._check_contract_can_edit_lines()
        return super().write(vals)

    def unlink(self):
        self._check_contract_can_edit_lines()
        return super().unlink()

    @api.onchange('product_tmpl_id')
    def _onchange_product_tmpl_id(self):
        for line in self:
            pt = line.product_tmpl_id
            if not pt:
                continue
            line.name = getattr(pt, 'get_product_multiline_description_sale', lambda: pt.display_name)()
            line.price_unit = pt.list_price or 0
            line.standard_price = pt.standard_price or 0
            line.compensation_price = pt.compensation_price or 0
            line.uom_id = pt.uom_id or 0

    def _effective_price_unit(self, as_of_date=None):
        """Đơn giá thuê / tháng hiệu lực tại as_of_date (mốc gần nhất có date_from <= as_of_date).

        Nếu chưa có mốc nào áp dụng (hoặc không truyền ngày) thì dùng price_unit gốc
        (giá ban đầu của hợp đồng).
        """
        self.ensure_one()
        if as_of_date and self.price_history_ids:
            applicable = self.price_history_ids.filtered(
                lambda h: h.date_from and h.date_from <= as_of_date
            ).sorted('date_from')
            if applicable:
                return applicable[-1].price_unit
        return self.price_unit


class RentalContractLinePrice(models.Model):
    _name = "rental.contract.line.price"
    _description = "Điều chỉnh đơn giá thuê theo thời gian"
    _order = "date_from, id"

    contract_line_id = fields.Many2one(
        'rental.contract.line',
        string="Dòng báo giá",
        required=True,
        ondelete='cascade',
        index=True,
    )
    contract_id = fields.Many2one(
        related='contract_line_id.contract_id',
        store=True,
        index=True,
        readonly=True,
    )
    company_id = fields.Many2one(related='contract_line_id.company_id', store=True, index=True, readonly=True)
    company_group_id = fields.Many2one(
        related='contract_line_id.company_group_id', store=True, index=True, readonly=True
    )
    currency_id = fields.Many2one(related='contract_line_id.currency_id', readonly=True)
    product_tmpl_id = fields.Many2one(
        related='contract_line_id.product_tmpl_id',
        string="Sản phẩm",
        store=True,
        readonly=True,
    )
    date_from = fields.Date(string="Hiệu lực từ ngày", required=True)
    price_unit = fields.Float(
        string="Đơn giá thuê / tháng mới",
        required=True,
        help="Đơn giá thuê theo tháng sau mốc điều chỉnh (cùng đơn vị bảng báo giá).",
    )

    def _check_contract_can_edit_price_changes(self):
        if self.env.context.get("rental_contract_allow_locked_write"):
            return
        for rec in self:
            contract = rec.contract_id or rec.contract_line_id.contract_id
            if contract and not contract.can_edit:
                raise UserError(_(
                    "Hợp đồng hiện đang bị khóa, không cho phép sửa điều chỉnh giá. "
                    "Vui lòng yêu cầu Leader mở khóa."
                ))

    @api.model_create_multi
    def create(self, vals_list):
        records = super().create(vals_list)
        records._check_contract_can_edit_price_changes()
        return records

    def write(self, vals):
        self._check_contract_can_edit_price_changes()
        return super().write(vals)

    def unlink(self):
        self._check_contract_can_edit_price_changes()
        return super().unlink()

    @api.constrains('price_unit')
    def _check_price_unit(self):
        for rec in self:
            if rec.price_unit < 0:
                raise ValidationError(_("Đơn giá thuê không được âm."))


class ConstructionWork(models.Model):
    _name = 'construction.work'
    _inherit = ['mail.thread', 'mail.activity.mixin', 'mc.group.mixin']

    name = fields.Char(string='Tên gói thầu', tracking=True, required=True)
    project_id = fields.Many2one('construction.project', string='Dự án', tracking=True, required=True)
    address_ids = fields.Many2many('construction.address', string='Địa chỉ', tracking=True, required=True)
    # Legacy field kept for compatibility with mail tracking and old data
    address = fields.Char(string='Address', tracking=True)
