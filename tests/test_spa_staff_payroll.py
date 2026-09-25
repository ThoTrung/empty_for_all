# -*- coding: utf-8 -*-

import inspect
from datetime import date, datetime, timedelta
from unittest import SkipTest
from unittest.mock import patch

from odoo import fields
from odoo.exceptions import UserError, ValidationError
from odoo.tests.common import TransactionCase, tagged


@tagged("post_install", "-at_install")
class TestSpaStaffPayroll(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        if "hr.employee" not in cls.env.registry:
            raise SkipTest("hr is not installed in this database")
        cls.company = cls.env.company
        cls.partner = cls.env["res.partner"].create({"name": "Payroll Test Partner"})
        cls.user = cls.env["res.users"].create({
            "name": "Therapist Payroll",
            "login": "therapist_payroll_test",
            "email": "tp@test.local",
            "groups_id": [(6, 0, [cls.env.ref("base.group_user").id])],
        })
        cls.level_spec_b = cls.env["spa.staff.level"].search([("level_group", "=", "spec_b")], limit=1)
        if not cls.level_spec_b:
            cls.level_spec_b = cls.env["spa.staff.level"].create({
                "name": "CV B (test)",
                "level_group": "spec_b",
                "rank": 1,
            })
        cls.user.sudo().write({"spa_staff_level_id": cls.level_spec_b.id})
        cls.employee = cls.env["hr.employee"].create({
            "name": "Therapist Payroll",
            "user_id": cls.user.id,
            "company_id": cls.company.id,
        })
        cls.contract = cls.env["hr.contract"].create({
            "name": "Test Contract",
            "employee_id": cls.employee.id,
            "date_start": date(2020, 1, 1),
            "wage": 10000000,
            "spa_overtime_hourly_rate": 100000,
            "state": "open",
        })
        cls.profile = cls.env["spa.product.payroll.profile"].create({
            "name": "Test payroll profile",
            "company_id": cls.company.id,
            "payout_line_ids": [
                (0, 0, {
                    "staff_level_id": cls.level_spec_b.id,
                    "amount_fixed": 123000,
                    "percent": 0.0,
                }),
            ],
        })
        cls.product_tmpl = cls.env["product.template"].create({
            "name": "Payroll Service Product",
            "detailed_type": "service",
            "spa_sessions_per_unit": 5,
            "service_employee_salary": 500000,
            "spa_payroll_profile_id": cls.profile.id,
            "list_price": 1000000,
            "spa_duration_minutes": 60,
            "spa_staff_capacity_percent": 100,
        })
        cls.product = cls.product_tmpl.product_variant_id

        cls.env["spa.payroll.long.shift.tier"].create({
            "company_id": cls.company.id,
            "name": ">=1 long shift",
            "min_count": 1,
            "max_count": False,
            "amount_per_shift": 5000,
        })
        cls.env["spa.payroll.requested.shift.tier"].create({
            "company_id": cls.company.id,
            "name": ">=1 requested long shift",
            "min_count": 1,
            "max_count": False,
            "amount_per_shift": 7000,
        })
        cls.card = cls.env["spa.treatment.card"].create({
            "partner_id": cls.partner.id,
            "product_id": cls.product.id,
            "total_sessions": 10,
        })

        cls.product_tmpl_b = cls.env["product.template"].create({
            "name": "Payroll Service Product B",
            "detailed_type": "service",
            "spa_sessions_per_unit": 5,
            "service_employee_salary": 111000,
            "spa_payroll_profile_id": cls.profile.id,
            "list_price": 900000,
        })
        cls.product_b = cls.product_tmpl_b.product_variant_id
        cls.bed = cls.env["spa.bed"].create({"name": "Payroll Test Bed"})

        # PAY-DEC-2026-09-16-01: fixtures cho điểm/voucher — reward product không thuế,
        # không gán category (nhận % HH mặc định của "All" = 0 trừ khi test tự đổi).
        cls.reward_product = cls.env["product.product"].create({
            "name": "Loyalty Reward (Payroll test)",
            "detailed_type": "service",
            "list_price": 1,
            "taxes_id": [(6, 0, [])],
        })
        cls.env["ir.config_parameter"].sudo().set_param(
            "spa_loyalty.reward_product_id", str(cls.reward_product.id)
        )

    def _grant_points(self, order, points):
        """Cấp điểm để test đổi điểm — ledger link ``order`` (đã confirm) để point_balance tính vào."""
        self.env["spa.loyalty.ledger"].create({
            "partner_id": order.partner_id.id,
            "point": points,
            "sale_order_ids": [(4, order.id)],
        })
        order.partner_id.invalidate_recordset(["point_balance"])

    def _redeem_points(self, so, points):
        self._grant_points(so, points)
        wiz = self.env["spa.redeem.points.wizard"].create({
            "sale_order_id": so.id,
            "point_to_redeem": points,
        })
        wiz.action_apply()
        so.invalidate_recordset()

    def _apply_amount_off_voucher(self, so, amount):
        template = self.env["spa.voucher"].create({
            "name": "Amount off (payroll test)",
            "voucher_kind": "order",
            "benefit_type": "amount_off",
            "value": amount,
            "validity_mode": "days_after_issue",
            "validity_days": 30,
        })
        voucher = self.env["spa.partner.voucher"].create({
            "partner_id": so.partner_id.id,
            "template_id": template.id,
            "source_type": "birthday",
            "source_year": 2025,
            "source_month": 8,
        })
        wiz = self.env["spa.apply.voucher.wizard"].create({
            "sale_order_id": so.id,
            "voucher_id": voucher.id,
        })
        wiz.action_apply()
        so.invalidate_recordset()
        return voucher

    def _ensure_shift_lines(self, start_dt, line_specs):
        """line_specs: [(user_ids, start_h, dur_h), ...]"""
        local = fields.Datetime.context_timestamp(self.env.user, start_dt)
        if getattr(local, "tzinfo", None):
            local = local.replace(tzinfo=None)
        day = local.date()
        cfg = self.env["booking.shift.config"].search([("shift_date", "=", day)], limit=1)
        if cfg:
            cfg.line_ids.unlink()
        else:
            cfg = self.env["booking.shift.config"].create(
                {"name": "Payroll test shift", "shift_date": day}
            )
        Line = self.env["booking.shift.config.line"]
        for uids, st, dur in line_specs:
            Line.create({
                "config_id": cfg.id,
                "name": "Ca",
                "shift_start_time_hours": float(st),
                "shift_duration_hours": float(dur),
                "user_ids": [(6, 0, list(uids))],
            })
        return cfg

    def _create_booking(self, **extra):
        start = datetime.now() + timedelta(days=1)
        start = start.replace(hour=10, minute=0, second=0, microsecond=0)
        self._ensure_shift_lines(start, [([self.user.id], 8.0, 10.0)])
        vals = {
            "partner_id": self.partner.id,
            "card_id": self.card.id,
            "product_id": self.product.id,
            "start_datetime": start,
            "duration": 60,
            "staff_ids": [(6, 0, [self.user.id])],
            "bed_id": self.bed.id,
        }
        vals.update(extra)
        return self.env["spa.service.booking"].create(vals)

    def test_session_split_creates_service_line(self):
        Session = self.env["spa.treatment.session"]
        session = Session.create({
            "card_id": self.card.id,
            "date": fields.Datetime.to_datetime(date(2025, 6, 15)),
            "therapist_id": self.user.id,
            "therapist_ids": [(6, 0, [self.user.id])],
            "duration_minutes": 60,
            "state": "done",
        })
        self.assertTrue(session.product_id)

        payroll = self.env["spa.staff.payroll"].create({
            "employee_id": self.employee.id,
            "company_id": self.company.id,
            "date_from": date(2025, 6, 1),
            "date_to": date(2025, 6, 30),
            "wage_fixed": self.contract.wage,
            "overtime_hourly_rate": self.contract.spa_overtime_hourly_rate,
            "currency_id": self.company.currency_id.id,
        })
        payroll.action_recompute_service_lines()
        self.assertEqual(len(payroll.service_line_ids), 1)
        self.assertEqual(payroll.service_line_ids.amount_share, 500000.0)

    def test_overtime_approved_in_payroll(self):
        ot = self.env["spa.staff.overtime.request"].create({
            "employee_id": self.employee.id,
            "overtime_date": date(2025, 6, 10),
            "hours": 2.0,
            "state": "submitted",
        })
        ot.sudo().action_approve()
        payroll = self.env["spa.staff.payroll"].create({
            "employee_id": self.employee.id,
            "company_id": self.company.id,
            "date_from": date(2025, 6, 1),
            "date_to": date(2025, 6, 30),
            "wage_fixed": self.contract.wage,
            "overtime_hourly_rate": self.contract.spa_overtime_hourly_rate,
            "currency_id": self.company.currency_id.id,
        })
        payroll.action_recompute_overtime_lines()
        self.assertEqual(len(payroll.overtime_line_ids), 1)
        self.assertEqual(payroll.overtime_line_ids.amount_subtotal, 200000.0)

    def test_session_payout_and_shift_bonuses(self):
        Session = self.env["spa.treatment.session"]
        start = fields.Datetime.to_datetime(date(2025, 6, 15))
        s1 = Session.create({
            "card_id": self.card.id,
            "date": start,
            "therapist_id": self.user.id,
            "therapist_ids": [(6, 0, [self.user.id])],
            "duration_minutes": 60,
            "state": "done",
        })
        s1.write({
            "spa_payroll_shift_kind": "long",
            "spa_payroll_customer_requested": True,
        })
        s2 = Session.create({
            "card_id": self.card.id,
            "date": start,
            "therapist_id": self.user.id,
            "therapist_ids": [(6, 0, [self.user.id])],
            "duration_minutes": 60,
            "state": "done",
        })
        s2.write({
            "spa_payroll_shift_kind": "short",
            "spa_payroll_customer_requested": False,
        })
        self.assertTrue(s1.product_id)
        self.assertTrue(s2.product_id)

        payroll = self.env["spa.staff.payroll"].create({
            "employee_id": self.employee.id,
            "company_id": self.company.id,
            "date_from": date(2025, 6, 1),
            "date_to": date(2025, 6, 30),
            "wage_fixed": self.contract.wage,
            "overtime_hourly_rate": self.contract.spa_overtime_hourly_rate,
            "currency_id": self.company.currency_id.id,
        })
        payroll.action_recompute_payroll_extras()

        svc_lines = payroll.line_ids.filtered(lambda l: l.category == "service_payout_session")
        self.assertEqual(len(svc_lines), 1)
        # Fixed payout line on profile is per session; ledger is aggregated.
        self.assertEqual(svc_lines.amount, 123000 * 2)
        self.assertEqual(len(payroll.service_line_ids), 2)
        self.assertAlmostEqual(
            sum(payroll.service_line_ids.mapped("amount_share")),
            svc_lines.amount,
        )

        long_bonus = payroll.line_ids.filtered(lambda l: l.category == "long_shift_bonus")
        req_bonus = payroll.line_ids.filtered(lambda l: l.category == "requested_bonus")
        self.assertEqual(len(long_bonus), 1)
        self.assertEqual(len(req_bonus), 1)
        self.assertEqual(long_bonus.amount, 5000.0)
        self.assertEqual(req_bonus.amount, 7000.0)

    def test_customer_requested_requires_staff_on_booking(self):
        booking = self._create_booking(staff_ids=[(5, 0, 0)])
        with self.assertRaises(ValidationError):
            booking.write({"spa_payroll_customer_requested": True})

    def test_customer_requested_ok_with_staff(self):
        booking = self._create_booking(spa_payroll_customer_requested=True)
        self.assertTrue(booking.spa_payroll_customer_requested)
        self.assertTrue(booking.staff_ids)

    def test_customer_requested_line_requires_staff(self):
        booking = self._create_booking(staff_ids=[(5, 0, 0)], bed_id=False)
        line = self.env["spa.service.booking.line"].create({
            "booking_id": booking.id,
            "product_id": self.product.id,
            "start_datetime": booking.start_datetime,
            "duration_minutes": 60,
        })
        with self.assertRaises(ValidationError):
            line.write({"spa_payroll_customer_requested": True})

    def test_customer_requested_calendar_color_draft_and_confirmed(self):
        ICP = self.env["ir.config_parameter"].sudo()
        ICP.set_param("spa.booking_calendar_hex_color_customer_requested", "#FF8C00")
        ICP.set_param("spa.booking_calendar_hex_text_color_customer_requested", "#000000")
        ICP.set_param("spa.booking_calendar_hex_color_draft", "#5e5ce6")
        ICP.set_param("spa.booking_calendar_hex_color_confirmed", "#2a9d8f")
        ICP.set_param("spa.booking_calendar_hex_color_doing", "#f4a261")

        booking = self._create_booking(spa_payroll_customer_requested=True)
        booking._compute_state_calendar_hex_color()
        booking._compute_state_calendar_hex_text_color()
        booking._compute_draft_special_colors()
        self.assertEqual(booking.state_calendar_hex_color, "#FF8C00")
        self.assertEqual(booking.state_calendar_hex_text_color, "#000000")
        self.assertEqual(booking.draft_special_hex_color, "#FF8C00")

        booking.write({"state": "confirmed"})
        booking._compute_state_calendar_hex_color()
        booking._compute_draft_special_colors()
        self.assertEqual(booking.state_calendar_hex_color, "#FF8C00")
        self.assertFalse(booking.draft_special_hex_color)

        booking.write({"state": "doing"})
        booking._compute_state_calendar_hex_color()
        self.assertEqual(booking.state_calendar_hex_color, "#f4a261")

    def test_customer_requested_overrides_draft_special(self):
        ICP = self.env["ir.config_parameter"].sudo()
        ICP.set_param("spa.booking_calendar_hex_color_customer_requested", "#FF8C00")
        ICP.set_param("spa.booking_calendar_hex_color_draft_weekly", "#3E51BA")

        booking = self._create_booking(
            spa_payroll_customer_requested=True,
            recurring_is_active=True,
        )
        booking._compute_draft_special_colors()
        self.assertEqual(booking.draft_special_hex_color, "#FF8C00")

    def test_customer_requested_line_colors_parent_booking(self):
        ICP = self.env["ir.config_parameter"].sudo()
        ICP.set_param("spa.booking_calendar_hex_color_customer_requested", "#FF8C00")
        ICP.set_param("spa.booking_calendar_hex_color_draft", "#5e5ce6")

        booking = self._create_booking(staff_ids=[(5, 0, 0)], bed_id=False)
        self.env["spa.service.booking.line"].create({
            "booking_id": booking.id,
            "product_id": self.product.id,
            "start_datetime": booking.start_datetime,
            "duration_minutes": 60,
            "staff_id": self.user.id,
            "spa_payroll_customer_requested": True,
        })
        booking.invalidate_recordset()
        booking._compute_state_calendar_hex_color()
        booking._compute_draft_special_colors()
        self.assertTrue(booking._spa_payroll_is_customer_requested())
        self.assertEqual(booking.state_calendar_hex_color, "#FF8C00")
        self.assertEqual(booking.draft_special_hex_color, "#FF8C00")

    def test_no_double_count_service_and_ledger(self):
        Session = self.env["spa.treatment.session"]
        Session.create({
            "card_id": self.card.id,
            "date": fields.Datetime.to_datetime(date(2025, 6, 15)),
            "therapist_id": self.user.id,
            "therapist_ids": [(6, 0, [self.user.id])],
            "duration_minutes": 60,
            "state": "done",
        })
        payroll = self.env["spa.staff.payroll"].create({
            "employee_id": self.employee.id,
            "company_id": self.company.id,
            "date_from": date(2025, 6, 1),
            "date_to": date(2025, 6, 30),
            "wage_fixed": 1_000_000,
            "overtime_hourly_rate": 0,
            "currency_id": self.company.currency_id.id,
        })
        payroll.action_recompute_service_lines()
        payroll.action_recompute_payroll_extras()
        svc_legacy = payroll.amount_service
        svc_ledger = sum(
            payroll.line_ids.filtered(lambda l: l.category == "service_payout_session").mapped("amount")
        )
        self.assertGreater(svc_legacy, 0)
        self.assertGreater(svc_ledger, 0)
        # Total must NOT include legacy amount_service
        expected = payroll.wage_fixed + payroll.amount_overtime + payroll.amount_payroll_lines
        self.assertAlmostEqual(payroll.amount_total_payable, expected)

    def test_multi_therapist_splits_payout(self):
        user2 = self.env["res.users"].create({
            "name": "Therapist 2",
            "login": "therapist_payroll_test_2",
            "email": "tp2@test.local",
            "groups_id": [(6, 0, [self.env.ref("base.group_user").id])],
        })
        user2.sudo().write({"spa_staff_level_id": self.level_spec_b.id})
        emp2 = self.env["hr.employee"].create({
            "name": "Therapist 2",
            "user_id": user2.id,
            "company_id": self.company.id,
        })
        Session = self.env["spa.treatment.session"]
        Session.create({
            "card_id": self.card.id,
            "date": fields.Datetime.to_datetime(date(2025, 7, 10)),
            "therapist_ids": [(6, 0, [self.user.id, user2.id])],
            "duration_minutes": 60,
            "state": "done",
        })
        p1 = self.env["spa.staff.payroll"].create({
            "employee_id": self.employee.id,
            "company_id": self.company.id,
            "date_from": date(2025, 7, 1),
            "date_to": date(2025, 7, 31),
            "wage_fixed": 0,
            "currency_id": self.company.currency_id.id,
        })
        p2 = self.env["spa.staff.payroll"].create({
            "employee_id": emp2.id,
            "company_id": self.company.id,
            "date_from": date(2025, 7, 1),
            "date_to": date(2025, 7, 31),
            "wage_fixed": 0,
            "currency_id": self.company.currency_id.id,
        })
        p1.action_recompute_payroll_extras()
        p2.action_recompute_payroll_extras()
        a1 = sum(p1.line_ids.filtered(lambda l: l.category == "service_payout_session").mapped("amount"))
        a2 = sum(p2.line_ids.filtered(lambda l: l.category == "service_payout_session").mapped("amount"))
        self.assertAlmostEqual(a1, 123000 / 2)
        self.assertAlmostEqual(a2, 123000 / 2)

    def test_booking_flags_sync_to_session(self):
        booking = self._create_booking(
            spa_payroll_shift_kind="long",
            spa_payroll_customer_requested=True,
        )
        session = self.env["spa.treatment.session"].create({
            "card_id": self.card.id,
            "booking_id": booking.id,
            "date": fields.Datetime.to_datetime(date(2025, 8, 5)),
            "therapist_ids": [(6, 0, [self.user.id])],
            "duration_minutes": 60,
            "state": "done",
        })
        self.assertEqual(session.spa_payroll_shift_kind, "long")
        self.assertTrue(session.spa_payroll_customer_requested)
        booking.write({"spa_payroll_shift_kind": "short", "spa_payroll_customer_requested": False})
        session.invalidate_recordset()
        self.assertEqual(session.spa_payroll_shift_kind, "short")
        self.assertFalse(session.spa_payroll_customer_requested)

    def test_seniority_excludes_long_leave(self):
        self.employee.spa_work_start_date = date(2020, 1, 1)
        self.env["spa.staff.dayoff"].create({
            "name": "Long leave",
            "employee_id": self.employee.id,
            "date_start": date(2021, 1, 1),
            "date_end": date(2021, 3, 1),
            "dayoff_type": "unpaid",
        })
        self.employee.invalidate_recordset()
        self.assertTrue(self.employee.spa_seniority_display)
        self.assertGreaterEqual(self.employee.spa_seniority_years, 4)

    def test_load_default_tiers_idempotent(self):
        """Global seed (company_id trống): 11 KPI + 2 long + 2 requested; ensure idempotent."""
        Config = self.env["spa.payroll.config"]
        Config._spa_ensure_default_tiers(self.company)
        Config._spa_ensure_default_tiers(self.company)

        Long = self.env["spa.payroll.long.shift.tier"]
        Req = self.env["spa.payroll.requested.shift.tier"]
        Kpi = self.env["spa.payroll.kpi.revenue.tier"]
        Prize = self.env["spa.payroll.ranking.prize"]

        longs = Long.search([("company_id", "=", False)])
        reqs = Req.search([("company_id", "=", False)])
        kpis = Kpi.search([("company_id", "=", False)])
        self.assertGreaterEqual(len(longs), 2)
        self.assertGreaterEqual(len(reqs), 2)
        self.assertGreaterEqual(len(kpis), 11)
        self.assertEqual(
            len(Long.search([("company_id", "=", False), ("min_count", "=", 38)])),
            1,
        )
        prizes = Prize.search([
            ("company_id", "=", False),
            ("period_type", "=", "month"),
            ("metric", "=", "sales_revenue"),
            ("rank", "=", 1),
        ])
        self.assertEqual(len(prizes), 1)
        self.assertEqual(prizes.amount, 1_000_000)

    def test_submit_approve_workflow(self):
        payroll = self.env["spa.staff.payroll"].create({
            "employee_id": self.employee.id,
            "company_id": self.company.id,
            "date_from": date(2025, 9, 1),
            "date_to": date(2025, 9, 30),
            "wage_fixed": 1000,
            "currency_id": self.company.currency_id.id,
        })
        payroll.action_submit_approve()
        self.assertEqual(payroll.state, "to_approve")
        payroll.action_confirm()
        self.assertEqual(payroll.state, "done")

    def test_profile_without_company_usable_on_product(self):
        """Profile company_id trống = dùng chung mọi công ty."""
        shared = self.env["spa.product.payroll.profile"].create({
            "name": "Shared profile no company",
            "company_id": False,
            "payout_line_ids": [
                (0, 0, {
                    "staff_level_id": self.level_spec_b.id,
                    "amount_fixed": 88000,
                    "percent": 0.0,
                }),
            ],
        })
        line = shared.payout_line_ids[:1]
        self.assertTrue(line.currency_id)
        self.assertEqual(line.currency_id, self.company.currency_id)

        tmpl = self.env["product.template"].create({
            "name": "Shared profile product",
            "detailed_type": "service",
            "spa_sessions_per_unit": 1,
            "service_employee_salary": 0,
            "spa_payroll_profile_id": shared.id,
            "list_price": 100000,
            "company_id": self.company.id,
        })
        # Domain on Many2one would allow False company; assert search domain used on product.
        domain = ["|", ("company_id", "=", False), ("company_id", "=", self.company.id)]
        self.assertIn(shared, self.env["spa.product.payroll.profile"].search(domain))

        card = self.env["spa.treatment.card"].create({
            "partner_id": self.partner.id,
            "product_id": tmpl.product_variant_id.id,
            "total_sessions": 5,
        })
        self.env["spa.treatment.session"].create({
            "card_id": card.id,
            "date": fields.Datetime.to_datetime(date(2025, 10, 5)),
            "therapist_ids": [(6, 0, [self.user.id])],
            "duration_minutes": 60,
            "state": "done",
        })
        payroll = self.env["spa.staff.payroll"].create({
            "employee_id": self.employee.id,
            "company_id": self.company.id,
            "date_from": date(2025, 10, 1),
            "date_to": date(2025, 10, 31),
            "wage_fixed": 0,
            "currency_id": self.company.currency_id.id,
        })
        payroll.action_recompute_payroll_extras()
        svc = payroll.line_ids.filtered(lambda l: l.category == "service_payout_session")
        self.assertEqual(len(svc), 1)
        self.assertAlmostEqual(svc.amount, 88000.0)

    def test_sales_commission_percent_from_category_and_parent(self):
        parent = self.env["product.category"].create({
            "name": "Parent Comm",
            "spa_sales_commission_percent": 3.0,
        })
        child = self.env["product.category"].create({
            "name": "Child Comm",
            "parent_id": parent.id,
            "spa_sales_commission_percent": 0.0,
        })
        tmpl = self.env["product.template"].create({
            "name": "Comm walk product",
            "detailed_type": "service",
            "categ_id": child.id,
            "list_price": 1000,
        })
        self.assertAlmostEqual(tmpl._spa_get_sales_commission_percent(), 3.0)
        child.write({"spa_sales_commission_percent": 5.0})
        self.assertAlmostEqual(tmpl._spa_get_sales_commission_percent(), 5.0)
        # Legacy SP field ignored
        tmpl.write({"spa_sales_commission_percent": 99.0})
        self.assertAlmostEqual(tmpl._spa_get_sales_commission_percent(), 5.0)

    def test_sales_commission_from_category_on_payroll(self):
        categ = self.env["product.category"].create({
            "name": "Comm 5pct",
            "spa_sales_commission_percent": 5.0,
        })
        tmpl = self.env["product.template"].create({
            "name": "Invoice Comm Product",
            "detailed_type": "service",
            "categ_id": categ.id,
            "list_price": 200000,
            "spa_sales_commission_percent": 50.0,  # ignored
        })
        product = tmpl.product_variant_id
        Move = self.env["account.move"].with_context(
            spa_allow_invoice_without_so=True
        )
        inv = Move.create({
            "move_type": "out_invoice",
            "partner_id": self.partner.id,
            "invoice_date": date(2025, 11, 10),
            "invoice_user_id": self.user.id,
            "company_id": self.company.id,
            "invoice_line_ids": [(0, 0, {
                "product_id": product.id,
                "quantity": 1,
                "price_unit": 200000,
                "name": product.display_name,
            })],
        })
        inv.action_post()
        self.env.cr.execute(
            "UPDATE account_move SET payment_state = 'paid' WHERE id = %s",
            (inv.id,),
        )
        inv.invalidate_recordset(["payment_state"])
        self.assertEqual(inv.payment_state, "paid")

        payroll = self.env["spa.staff.payroll"].create({
            "employee_id": self.employee.id,
            "company_id": self.company.id,
            "date_from": date(2025, 11, 1),
            "date_to": date(2025, 11, 30),
            "wage_fixed": 0,
            "currency_id": self.company.currency_id.id,
        })
        payroll.action_recompute_payroll_extras()
        comm = payroll.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertEqual(len(comm), 1)
        # 200000 * 5% = 10000 (taxes may affect subtotal; use line price_subtotal)
        expected = sum(
            inv.invoice_line_ids.filtered(lambda l: l.display_type == "product").mapped("price_subtotal")
        ) * 0.05
        self.assertAlmostEqual(comm.amount, expected)
        self.assertEqual(len(payroll.commission_line_ids), 1)
        self.assertAlmostEqual(
            sum(payroll.commission_line_ids.mapped("commission_amount")),
            expected,
        )
        detail = payroll.commission_line_ids
        self.assertEqual(detail.product_id, product)
        self.assertEqual(detail.move_id, inv)
        self.assertAlmostEqual(detail.commission_percent, 5.0)

    def _spa_pay_invoice(self, invoice, pay_date, amount=None):
        """Post inbound payment and reconcile (full residual or partial ``amount``)."""
        journal = self.env["account.journal"].search([
            ("company_id", "=", self.company.id),
            ("type", "in", ("cash", "bank")),
        ], limit=1)
        if not journal:
            raise SkipTest("No cash/bank journal for payment test")
        pay_amt = float(amount if amount is not None else invoice.amount_residual)
        payment = self.env["account.payment"].create({
            "payment_type": "inbound",
            "partner_type": "customer",
            "partner_id": invoice.partner_id.id,
            "amount": pay_amt,
            "date": pay_date,
            "journal_id": journal.id,
            "currency_id": invoice.currency_id.id,
        })
        payment.action_post()
        lines = (invoice.line_ids + payment.move_id.line_ids).filtered(
            lambda l: l.account_id.account_type == "asset_receivable" and not l.reconciled
        )
        lines.reconcile()
        invoice.invalidate_recordset()
        return payment

    def _make_comm_product(self, name, pct, price=200000):
        categ = self.env["product.category"].create({
            "name": "Categ %s" % name,
            "spa_sales_commission_percent": pct,
        })
        tmpl = self.env["product.template"].create({
            "name": name,
            "detailed_type": "service",
            "invoice_policy": "order",
            "categ_id": categ.id,
            "list_price": price,
        })
        return tmpl.product_variant_id

    def _make_comm_card_product(self, name, pct, price=200000, sessions=10):
        """Comm product usable as a treatment-card package (spa_sessions_per_unit set)."""
        product = self._make_comm_product(name, pct, price)
        product.product_tmpl_id.write({
            "spa_sessions_per_unit": sessions,
            "taxes_id": [(6, 0, [])],
        })
        return product

    def _make_so_with_invoice(
        self,
        products_prices,
        user=None,
        invoice_date=None,
        discount=0.0,
        create_invoice=True,
        post_invoice=True,
    ):
        """Create confirmed SO; optionally invoice+post.

        ``products_prices``: list of (product, price_unit) or product alone (uses list_price).
        """
        user = user or self.user
        invoice_date = invoice_date or date(2025, 5, 10)
        lines = []
        for item in products_prices:
            if isinstance(item, tuple):
                product, price = item
            else:
                product = item
                price = product.list_price
            lines.append((0, 0, {
                "product_id": product.id,
                "product_uom_qty": 1,
                "price_unit": price,
                "discount": discount,
            }))
        so = self.env["sale.order"].create({
            "partner_id": self.partner.id,
            "user_id": user.id,
            "company_id": self.company.id,
            "order_line": lines,
        })
        so.action_confirm()
        inv = self.env["account.move"]
        if create_invoice:
            inv = so._create_invoices()
            inv.write({
                "invoice_date": invoice_date,
                "invoice_user_id": user.id,
            })
            if post_invoice:
                inv.action_post()
        return so, inv

    def _payroll_for_month(self, year, month, employee=None):
        employee = employee or self.employee
        last = 31 if month != 2 else 28
        if month in (4, 6, 9, 11):
            last = 30
        return self.env["spa.staff.payroll"].create({
            "employee_id": employee.id,
            "company_id": self.company.id,
            "date_from": date(year, month, 1),
            "date_to": date(year, month, last),
            "wage_fixed": 0,
            "currency_id": self.company.currency_id.id,
        })

    def _expected_so_commission(self, so):
        """Công thức tính lại độc lập (không gọi code production)."""
        base_lines = so.order_line.filtered(
            lambda l: not l._spa_is_order_level_discount()
            and not l.display_type
            and not l.is_downpayment
        )
        base_total = sum(base_lines.mapped("price_subtotal"))
        discount_untaxed = sum(
            so.order_line.filtered(lambda l: l._spa_is_order_level_discount()).mapped("price_subtotal")
        )
        total = 0.0
        for line in so.order_line:
            if line.display_type or not line.product_id:
                continue
            if line.is_downpayment or line._spa_is_order_level_discount():
                continue
            pct = float(line.product_id.product_tmpl_id._spa_get_sales_commission_percent() or 0.0)
            if not pct:
                continue
            share = discount_untaxed * (line.price_subtotal / base_total) if base_total else 0.0
            total += (float(line.price_subtotal or 0.0) + share) * (pct / 100.0)
        return total

    def test_sales_commission_so_settled_on_full_payment_month(self):
        """SO invoice May, pay July → HH on July payroll only; base after line discount."""
        product = self._make_comm_product("SO Comm Product", 5.0, 200000)
        so, inv = self._make_so_with_invoice(
            [(product, 200000)],
            invoice_date=date(2025, 5, 10),
            discount=10.0,
        )
        self.assertFalse(self.company.currency_id.is_zero(inv.amount_residual))

        pay_may = self._payroll_for_month(2025, 5)
        pay_may.action_recompute_payroll_extras()
        self.assertFalse(
            pay_may.line_ids.filtered(lambda l: l.category == "sales_commission"),
            "Deposit / unpaid SO must not pay commission in May",
        )

        self._spa_pay_invoice(inv, date(2025, 7, 15))
        self.assertTrue(self.company.currency_id.is_zero(inv.amount_residual))
        self.assertEqual(inv.payment_state, "paid")

        pay_jul = self._payroll_for_month(2025, 7)
        pay_jul.action_recompute_payroll_extras()
        comm = pay_jul.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertEqual(len(comm), 1)
        self.assertAlmostEqual(comm.amount, self._expected_so_commission(so))

        pay_may.action_recompute_payroll_extras()
        self.assertFalse(
            pay_may.line_ids.filtered(lambda l: l.category == "sales_commission"),
        )

    def test_sales_commission_so_partial_payment_no_hh(self):
        """T1: partial pay leaves residual → no commission / no SO KPI."""
        product = self._make_comm_product("SO Partial", 5.0, 200000)
        so, inv = self._make_so_with_invoice([(product, 200000)])
        self._spa_pay_invoice(inv, date(2025, 7, 10), amount=50000)
        self.assertFalse(self.company.currency_id.is_zero(inv.amount_residual))

        pay_jul = self._payroll_for_month(2025, 7)
        pay_jul.action_recompute_payroll_extras()
        self.assertFalse(pay_jul.line_ids.filtered(lambda l: l.category == "sales_commission"))
        self.assertAlmostEqual(pay_jul.kpi_revenue_base, 0.0)

    def test_sales_commission_so_two_products_different_pct(self):
        """T2: two lines with different category % → sum after settle."""
        p5 = self._make_comm_product("SO Multi 5", 5.0, 100000)
        p10 = self._make_comm_product("SO Multi 10", 10.0, 200000)
        so, inv = self._make_so_with_invoice([(p5, 100000), (p10, 200000)])
        self._spa_pay_invoice(inv, date(2025, 7, 20))

        pay_jul = self._payroll_for_month(2025, 7)
        pay_jul.action_recompute_payroll_extras()
        comm = pay_jul.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertEqual(len(comm), 1)
        self.assertAlmostEqual(comm.amount, self._expected_so_commission(so))
        # 100000*5% + 200000*10% = 5000 + 20000 = 25000
        self.assertAlmostEqual(comm.amount, 25000.0)

    def test_kpi_so_settled_month_july_not_may(self):
        """T3: KPI base/line only in settle month (July)."""
        self.env["spa.payroll.kpi.revenue.tier"].create({
            "company_id": False,
            "name": "KPI test from 1",
            "min_revenue": 1.0,
            "percent": 1.0,
            "currency_id": self.company.currency_id.id,
        })
        product = self._make_comm_product("SO KPI", 0.0, 500000)
        so, inv = self._make_so_with_invoice([(product, 500000)])

        pay_may = self._payroll_for_month(2025, 5)
        pay_may.action_recompute_payroll_extras()
        self.assertAlmostEqual(pay_may.kpi_revenue_base, 0.0)
        self.assertFalse(pay_may.line_ids.filtered(lambda l: l.category == "kpi_revenue"))

        self._spa_pay_invoice(inv, date(2025, 7, 15))
        pay_jul = self._payroll_for_month(2025, 7)
        pay_jul.action_recompute_payroll_extras()
        self.assertAlmostEqual(pay_jul.kpi_revenue_base, so._spa_recognized_amount_untaxed())
        kpi = pay_jul.line_ids.filtered(lambda l: l.category == "kpi_revenue")
        self.assertEqual(len(kpi), 1)
        self.assertAlmostEqual(kpi.amount, so._spa_recognized_amount_untaxed() * 0.01)

    def test_sales_commission_standalone_invoice_paid_date_not_invoice_date(self):
        """T4: standalone HĐ dated May, paid July → HH in July only."""
        product = self._make_comm_product("Standalone MayJul", 5.0, 200000)
        inv = self.env["account.move"].with_context(
            spa_allow_invoice_without_so=True
        ).create({
            "move_type": "out_invoice",
            "partner_id": self.partner.id,
            "invoice_date": date(2025, 5, 12),
            "invoice_user_id": self.user.id,
            "company_id": self.company.id,
            "invoice_line_ids": [(0, 0, {
                "product_id": product.id,
                "quantity": 1,
                "price_unit": 200000,
                "name": product.display_name,
            })],
        })
        inv.action_post()

        pay_may = self._payroll_for_month(2025, 5)
        pay_may.action_recompute_payroll_extras()
        self.assertFalse(pay_may.line_ids.filtered(lambda l: l.category == "sales_commission"))

        self._spa_pay_invoice(inv, date(2025, 7, 18))
        pay_jul = self._payroll_for_month(2025, 7)
        pay_jul.action_recompute_payroll_extras()
        comm = pay_jul.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertEqual(len(comm), 1)
        expected = sum(
            inv.invoice_line_ids.filtered(lambda l: l.display_type == "product").mapped("price_subtotal")
        ) * 0.05
        self.assertAlmostEqual(comm.amount, expected)

    def test_sales_commission_wrong_salesperson_zero(self):
        """T5: SO salesperson ≠ payroll employee user → no HH on that payslip."""
        other_user = self.env["res.users"].create({
            "name": "Other Sales",
            "login": "other_sales_comm_test",
            "email": "other_sales@test.local",
            "groups_id": [(6, 0, [self.env.ref("base.group_user").id])],
        })
        product = self._make_comm_product("Wrong User SO", 5.0, 150000)
        _so, inv = self._make_so_with_invoice(
            [(product, 150000)],
            user=other_user,
        )
        self._spa_pay_invoice(inv, date(2025, 7, 12))

        pay_jul = self._payroll_for_month(2025, 7)
        pay_jul.action_recompute_payroll_extras()
        self.assertFalse(pay_jul.line_ids.filtered(lambda l: l.category == "sales_commission"))

    def test_sales_commission_so_not_fully_invoiced(self):
        """T6: confirmed SO without invoice → not settled → 0."""
        product = self._make_comm_product("Not Invoiced", 5.0, 100000)
        so, _inv = self._make_so_with_invoice(
            [(product, 100000)],
            create_invoice=False,
        )
        self.assertFalse(self.company.currency_id.is_zero(float(so.amount_to_invoice or 0.0)))

        pay_jul = self._payroll_for_month(2025, 7)
        pay_jul.action_recompute_payroll_extras()
        self.assertFalse(pay_jul.line_ids.filtered(lambda l: l.category == "sales_commission"))
        proxy = self.env["spa.staff.payroll"].new({
            "company_id": self.company.id,
            "date_from": date(2025, 7, 1),
            "date_to": date(2025, 7, 31),
            "employee_id": self.employee.id,
        })
        self.assertFalse(proxy._spa_is_sale_order_settled(so))

    def test_sales_commission_no_double_count_so_invoice(self):
        """T7: SO-linked invoice must not also count as standalone HH."""
        product = self._make_comm_product("No Double", 5.0, 200000)
        so, inv = self._make_so_with_invoice([(product, 200000)])
        self._spa_pay_invoice(inv, date(2025, 7, 22))
        expected = self._expected_so_commission(so)

        pay_jul = self._payroll_for_month(2025, 7)
        pay_jul.action_recompute_payroll_extras()
        comm = pay_jul.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertEqual(len(comm), 1)
        self.assertAlmostEqual(comm.amount, expected)

        proxy = self.env["spa.staff.payroll"].new({
            "company_id": self.company.id,
            "date_from": date(2025, 7, 1),
            "date_to": date(2025, 7, 31),
            "employee_id": self.employee.id,
        })
        standalone = proxy._spa_standalone_paid_invoices_in_period(self.user)
        self.assertNotIn(inv.id, standalone.ids)

    def test_ranking_scores_sales_cash_basis_july(self):
        """T8: ranking sales score uses settle month (July), not invoice month."""
        product = self._make_comm_product("Rank Sales", 0.0, 300000)
        _so, inv = self._make_so_with_invoice([(product, 300000)])
        self._spa_pay_invoice(inv, date(2025, 7, 25))

        wiz = self.env["spa.payroll.ranking.compute.wizard"].create({
            "company_id": self.company.id,
            "period_type": "month",
            "year": 2025,
            "month": 7,
        })
        scores_jul = wiz._scores_sales(date(2025, 7, 1), date(2025, 7, 31))
        self.assertIn(self.employee.id, scores_jul)
        self.assertGreater(scores_jul[self.employee.id], 0)

        scores_may = wiz._scores_sales(date(2025, 5, 1), date(2025, 5, 31))
        self.assertNotIn(self.employee.id, scores_may)

    def test_sales_commission_zero_percent_no_line(self):
        """T9: category % = 0 → no sales_commission ledger line after settle."""
        product = self._make_comm_product("Zero Pct", 0.0, 200000)
        _so, inv = self._make_so_with_invoice([(product, 200000)])
        self._spa_pay_invoice(inv, date(2025, 7, 28))

        pay_jul = self._payroll_for_month(2025, 7)
        pay_jul.action_recompute_payroll_extras()
        self.assertFalse(pay_jul.line_ids.filtered(lambda l: l.category == "sales_commission"))
        self.assertFalse(pay_jul.commission_line_ids)

    def test_wage_prorate_mid_month_contract_start(self):
        """HĐ từ 15/7 → wage_fixed = wage × 17/31; không overlap tháng 6."""
        emp = self.env["hr.employee"].create({
            "name": "Mid Month Hire",
            "user_id": self.env["res.users"].create({
                "name": "Mid Month User",
                "login": "mid_month_hire_payroll",
                "email": "mid@test.local",
                "groups_id": [(6, 0, [self.env.ref("base.group_user").id])],
            }).id,
            "company_id": self.company.id,
        })
        contract = self.env["hr.contract"].create({
            "name": "Jul 15 Contract",
            "employee_id": emp.id,
            "date_start": date(2025, 7, 15),
            "wage": 31000000,
            "state": "open",
            "company_id": self.company.id,
        })
        Payroll = self.env["spa.staff.payroll"]
        vals = Payroll._spa_prorate_contract_wage(
            contract, date(2025, 7, 1), date(2025, 7, 31)
        )
        self.assertEqual(vals["wage_days"], 17)
        self.assertEqual(vals["wage_days_period"], 31)
        self.assertEqual(vals["wage_period_start"], date(2025, 7, 15))
        self.assertEqual(vals["wage_period_end"], date(2025, 7, 31))
        self.assertAlmostEqual(vals["wage_fixed"], 31000000.0 * 17.0 / 31.0)

        payroll = Payroll.create({
            "employee_id": emp.id,
            "company_id": self.company.id,
            "date_from": date(2025, 7, 1),
            "date_to": date(2025, 7, 31),
            "currency_id": self.company.currency_id.id,
        })
        payroll.action_load_contract_values()
        self.assertAlmostEqual(payroll.wage_fixed, 31000000.0 * 17.0 / 31.0)
        self.assertEqual(payroll.wage_days, 17)

        june_contracts = self.env["hr.contract"].search([
            ("employee_id", "=", emp.id),
            ("state", "=", "open"),
            ("date_start", "<=", date(2025, 6, 30)),
            "|",
            ("date_end", "=", False),
            ("date_end", ">=", date(2025, 6, 1)),
        ])
        self.assertFalse(june_contracts)

    def test_wage_prorate_contract_end_mid_month(self):
        """HĐ kết thúc 10/7 → 10/31 ngày."""
        emp = self.env["hr.employee"].create({
            "name": "Early End",
            "user_id": self.env["res.users"].create({
                "name": "Early End User",
                "login": "early_end_payroll",
                "email": "early@test.local",
                "groups_id": [(6, 0, [self.env.ref("base.group_user").id])],
            }).id,
            "company_id": self.company.id,
        })
        contract = self.env["hr.contract"].create({
            "name": "Ends Jul 10",
            "employee_id": emp.id,
            "date_start": date(2025, 1, 1),
            "date_end": date(2025, 7, 10),
            "wage": 31000000,
            "state": "open",
            "company_id": self.company.id,
        })
        vals = self.env["spa.staff.payroll"]._spa_prorate_contract_wage(
            contract, date(2025, 7, 1), date(2025, 7, 31)
        )
        self.assertEqual(vals["wage_days"], 10)
        self.assertAlmostEqual(vals["wage_fixed"], 31000000.0 * 10.0 / 31.0)

    def test_commission_and_kpi_detail_lines_match_ledger(self):
        """Detail tabs: sum(commission) == ledger; KPI detail sum == ledger KPI."""
        p5 = self._make_comm_product("Detail 5pct", 5.0, 100000)
        p10 = self._make_comm_product("Detail 10pct", 10.0, 200000)
        so, inv = self._make_so_with_invoice([(p5, 100000), (p10, 200000)])
        self._spa_pay_invoice(inv, date(2025, 7, 20))

        self.env["spa.payroll.kpi.revenue.tier"].create({
            "company_id": self.company.id,
            "name": "KPI test tier detail",
            "min_revenue": 0,
            "percent": 2.0,
        })

        pay = self._payroll_for_month(2025, 7)
        pay.action_recompute_payroll_extras()

        expected_comm = self._expected_so_commission(so)
        comm_ledger = pay.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertEqual(len(comm_ledger), 1)
        self.assertAlmostEqual(comm_ledger.amount, expected_comm)
        self.assertEqual(len(pay.commission_line_ids), 2)
        self.assertAlmostEqual(
            sum(pay.commission_line_ids.mapped("commission_amount")),
            expected_comm,
        )
        for dline in pay.commission_line_ids:
            self.assertTrue(dline.product_id)
            self.assertTrue(dline.categ_id)
            self.assertEqual(dline.sale_order_id, so)
            self.assertTrue(dline.commission_percent)

        kpi_ledger = pay.line_ids.filtered(lambda l: l.category == "kpi_revenue")
        self.assertTrue(kpi_ledger)
        self.assertAlmostEqual(
            sum(pay.kpi_line_ids.mapped("kpi_amount")),
            kpi_ledger.amount,
        )
        self.assertAlmostEqual(
            sum(pay.kpi_line_ids.mapped("revenue_amount")),
            pay.kpi_revenue_base,
        )

    def test_long_shift_bonus_aggregated_one_line(self):
        """Two long sessions → one long_shift_bonus ledger line; total = rate × count."""
        Session = self.env["spa.treatment.session"]
        start = fields.Datetime.to_datetime(date(2025, 8, 5))
        for _ in range(2):
            ses = Session.create({
                "card_id": self.card.id,
                "date": start,
                "therapist_id": self.user.id,
                "therapist_ids": [(6, 0, [self.user.id])],
                "duration_minutes": 60,
                "state": "done",
            })
            ses.write({"spa_payroll_shift_kind": "long"})

        payroll = self.env["spa.staff.payroll"].create({
            "employee_id": self.employee.id,
            "company_id": self.company.id,
            "date_from": date(2025, 8, 1),
            "date_to": date(2025, 8, 31),
            "wage_fixed": 0,
            "currency_id": self.company.currency_id.id,
        })
        payroll.action_recompute_payroll_extras()
        long_bonus = payroll.line_ids.filtered(lambda l: l.category == "long_shift_bonus")
        self.assertEqual(len(long_bonus), 1)
        self.assertAlmostEqual(long_bonus.amount, 10000.0)
        svc = payroll.line_ids.filtered(lambda l: l.category == "service_payout_session")
        self.assertEqual(len(svc), 1)
        self.assertEqual(len(payroll.service_line_ids), 2)
        self.assertAlmostEqual(
            sum(payroll.service_line_ids.mapped("amount_share")),
            svc.amount,
        )

    def _spa_pay_move(self, invoice, pay_date, amount=None):
        """Pay invoice or credit note (inbound vs outbound)."""
        journal = self.env["account.journal"].search([
            ("company_id", "=", self.company.id),
            ("type", "in", ("cash", "bank")),
        ], limit=1)
        if not journal:
            raise SkipTest("No cash/bank journal for payment test")
        pay_amt = float(amount if amount is not None else invoice.amount_residual)
        payment = self.env["account.payment"].create({
            "payment_type": "inbound" if invoice.move_type == "out_invoice" else "outbound",
            "partner_type": "customer",
            "partner_id": invoice.partner_id.id,
            "amount": abs(pay_amt),
            "date": pay_date,
            "journal_id": journal.id,
            "currency_id": invoice.currency_id.id,
        })
        payment.action_post()
        lines = (invoice.line_ids + payment.move_id.line_ids).filtered(
            lambda l: l.account_id.account_type == "asset_receivable" and not l.reconciled
        )
        lines.reconcile()
        invoice.invalidate_recordset()
        if hasattr(invoice, "_compute_spa_fully_paid_date"):
            invoice._compute_spa_fully_paid_date()
            invoice.flush_recordset(["spa_fully_paid_date"])
        return payment

    def test_late_refund_keeps_origin_month_and_claws_back_refund_month(self):
        """SO settled July; CN paid October → July stays; October draft KPI/HH negative."""
        product = self._make_comm_product("SO Late Refund", 5.0, 200000)
        so, inv = self._make_so_with_invoice(
            [(product, 200000)],
            invoice_date=date(2025, 5, 10),
        )
        self._spa_pay_move(inv, date(2025, 7, 15))
        so.invalidate_recordset()
        so._compute_spa_settlement()
        so.flush_recordset(["spa_is_settled", "spa_settled_date"])
        self.assertEqual(so.spa_settled_date, date(2025, 7, 15))

        pay_jul = self._payroll_for_month(2025, 7)
        pay_jul.action_recompute_payroll_extras()
        jul_comm = pay_jul.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertEqual(len(jul_comm), 1)
        jul_comm_amt = jul_comm.amount
        jul_kpi_base = pay_jul.kpi_revenue_base
        self.assertAlmostEqual(jul_kpi_base, so._spa_recognized_amount_untaxed())
        pay_jul.action_confirm()
        self.assertEqual(pay_jul.state, "done")

        refund = inv._reverse_moves(default_values_list=[{
            "date": date(2025, 10, 5),
            "invoice_date": date(2025, 10, 5),
            "ref": "Payroll late refund",
        }], cancel=False)
        refund = refund[0]
        if refund.state != "posted":
            refund.action_post()
        self._spa_pay_move(refund, date(2025, 10, 20))
        so.invalidate_recordset()
        so._compute_spa_settlement()
        so.flush_recordset(["spa_is_settled", "spa_settled_date"])
        self.assertEqual(so.spa_settled_date, date(2025, 7, 15))

        pay_jul.invalidate_recordset()
        self.assertEqual(pay_jul.state, "done")
        self.assertAlmostEqual(
            pay_jul.line_ids.filtered(lambda l: l.category == "sales_commission").amount,
            jul_comm_amt,
        )

        pay_oct = self._payroll_for_month(2025, 10)
        pay_oct.action_recompute_payroll_extras()
        self.assertAlmostEqual(pay_oct.kpi_revenue_base, float(refund.amount_untaxed_signed))
        oct_comm = pay_oct.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertTrue(oct_comm)
        self.assertLess(oct_comm.amount, 0)

        pay_jul_draft = self._payroll_for_month(2025, 7)
        pay_jul_draft.action_recompute_payroll_extras()
        self.assertAlmostEqual(pay_jul_draft.kpi_revenue_base, so._spa_recognized_amount_untaxed())
        self.assertAlmostEqual(
            pay_jul_draft.line_ids.filtered(lambda l: l.category == "sales_commission").amount,
            jul_comm_amt,
        )

    def test_kpi_recognized_amount_prorates_global_discount_into_commission(self):
        """KPI = SO net CK; HH KHÔNG tính dòng is_global_discount nhưng có TRỪ
        phần CK toàn đơn phân bổ (untaxed) khỏi subtotal của dòng SP còn lại
        (PAY-DEC-2026-09-03-01 / global-discount fix)."""
        self.env["spa.payroll.kpi.revenue.tier"].create({
            "company_id": False,
            "name": "KPI CK recognized",
            "min_revenue": 1.0,
            "percent": 1.0,
            "currency_id": self.company.currency_id.id,
        })
        product = self._make_comm_product("SO KPI CK prod", 5.0, 2000000)
        ck = self._make_comm_product("CK HH bait", 5.0, 0)
        so = self.env["sale.order"].create({
            "partner_id": self.partner.id,
            "user_id": self.user.id,
            "company_id": self.company.id,
            "order_line": [
                (0, 0, {
                    "product_id": product.id,
                    "product_uom_qty": 1,
                    "price_unit": 2000000,
                }),
                (0, 0, {
                    "product_id": ck.id,
                    "name": "CK",
                    "product_uom_qty": 1,
                    "price_unit": -200000,
                    "is_global_discount": True,
                }),
            ],
        })
        so.action_confirm()
        inv = so._create_invoices()
        inv.write({
            "invoice_date": date(2025, 7, 1),
            "invoice_user_id": self.user.id,
        })
        inv.action_post()
        self._spa_pay_invoice(inv, date(2025, 7, 15))

        pay = self._payroll_for_month(2025, 7)
        pay.action_recompute_payroll_extras()
        recognized = so._spa_recognized_amount_untaxed()
        self.assertAlmostEqual(pay.kpi_revenue_base, recognized)
        self.assertAlmostEqual(
            pay._spa_net_invoice_revenue_for_sales_user(self.user),
            recognized,
        )
        comm = pay.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertAlmostEqual(comm.amount, self._expected_so_commission(so))
        # CK toàn đơn -200,000đ được trừ vào subtotal dòng SP trước khi tính %.
        self.assertAlmostEqual(comm.amount, (2000000.0 - 200000.0) * 0.05)
        self.assertFalse(
            pay.commission_line_ids.filtered(lambda l: l.product_id == ck)
        )

    def test_commission_prorates_points_redemption_like_global_discount(self):
        """PAY-DEC-2026-09-16-01: đổi điểm phải cho HH giống hệt CK toàn đơn cùng số tiền
        (test trên chính khớp với test global-discount ở trên: 2,000,000 × 5% sau khi
        trừ 200,000 = (2,000,000 - 200,000) × 5% = 90,000)."""
        product = self._make_comm_product("SO Points prod", 5.0, 2000000)
        product.taxes_id = [(6, 0, [])]
        so, _inv = self._make_so_with_invoice(
            [(product, 2000000)], invoice_date=date(2025, 7, 1), create_invoice=False,
        )
        self._redeem_points(so, 200)  # 200 điểm × 1000đ = 200,000đ
        self.assertTrue(
            so.order_line.filtered(lambda l: l.spa_is_order_discount_line),
            "wizard phải gán cờ spa_is_order_discount_line",
        )
        inv = so._create_invoices()
        inv.write({"invoice_date": date(2025, 7, 1), "invoice_user_id": self.user.id})
        inv.action_post()
        self._spa_pay_invoice(inv, date(2025, 7, 15))

        pay = self._payroll_for_month(2025, 7)
        pay.action_recompute_payroll_extras()
        comm = pay.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertAlmostEqual(comm.amount, (2000000.0 - 200000.0) * 0.05)
        self.assertFalse(
            pay.commission_line_ids.filtered(lambda l: l.product_id == self.reward_product),
            "dòng điểm không được sinh commission line riêng",
        )
        self.assertAlmostEqual(
            pay.kpi_revenue_base, so._spa_recognized_amount_untaxed()
        )

    def test_commission_prorates_amount_off_voucher(self):
        """Voucher giảm tiền dùng cùng cơ chế reward-product — phải cho kết quả giống điểm."""
        product = self._make_comm_product("SO Voucher prod", 10.0, 1000000)
        product.taxes_id = [(6, 0, [])]
        so, _inv = self._make_so_with_invoice(
            [(product, 1000000)], invoice_date=date(2025, 7, 1), create_invoice=False,
        )
        self._apply_amount_off_voucher(so, 100000)
        self.assertTrue(so.order_line.filtered(lambda l: l.spa_is_order_discount_line))
        inv = so._create_invoices()
        inv.write({"invoice_date": date(2025, 7, 1), "invoice_user_id": self.user.id})
        inv.action_post()
        self._spa_pay_invoice(inv, date(2025, 7, 15))

        pay = self._payroll_for_month(2025, 7)
        pay.action_recompute_payroll_extras()
        comm = pay.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertAlmostEqual(comm.amount, (1000000.0 - 100000.0) * 0.10)

    def test_commission_global_discount_and_points_together(self):
        """CK toàn đơn + điểm cùng đơn: mẫu số phân bổ phải loại CẢ HAI loại giảm giá
        (hồi quy lỗi 'mẫu số gồm cả dòng điểm' — PAY-DEC-2026-09-16-01).

        2 SP × 1,000,000 (10%), CK -100,000, điểm -50,000.
        Đúng: (2,000,000 - 150,000) × 10% = 185,000.
        """
        p1 = self._make_comm_product("Combo A", 10.0, 1000000)
        p1.taxes_id = [(6, 0, [])]
        p2 = self._make_comm_product("Combo B", 10.0, 1000000)
        p2.taxes_id = [(6, 0, [])]
        ck = self._make_comm_product("Combo CK", 10.0, 0)
        ck.taxes_id = [(6, 0, [])]
        so = self.env["sale.order"].create({
            "partner_id": self.partner.id,
            "user_id": self.user.id,
            "company_id": self.company.id,
            "order_line": [
                (0, 0, {"product_id": p1.id, "product_uom_qty": 1, "price_unit": 1000000}),
                (0, 0, {"product_id": p2.id, "product_uom_qty": 1, "price_unit": 1000000}),
                (0, 0, {
                    "product_id": ck.id, "name": "CK HD", "product_uom_qty": 1,
                    "price_unit": -100000, "is_global_discount": True,
                }),
            ],
        })
        so.action_confirm()
        self._redeem_points(so, 50)  # -50,000đ
        inv = so._create_invoices()
        inv.write({"invoice_date": date(2025, 7, 1), "invoice_user_id": self.user.id})
        inv.action_post()
        self._spa_pay_invoice(inv, date(2025, 7, 15))

        pay = self._payroll_for_month(2025, 7)
        pay.action_recompute_payroll_extras()
        comm = pay.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertAlmostEqual(comm.amount, 185000.0)

    def test_reward_product_under_parent_category_not_double_deducted(self):
        """SP thưởng thuộc danh mục con của danh mục có % HH > 0 — dòng điểm vẫn phải bị
        skip tường minh, không được vừa cộng share vừa tự sinh commission âm."""
        parent_categ = self.env["product.category"].create({
            "name": "Reward parent categ (10%)",
            "spa_sales_commission_percent": 10.0,
        })
        self.reward_product.product_tmpl_id.categ_id = self.env["product.category"].create({
            "name": "Reward child categ (0%)",
            "parent_id": parent_categ.id,
            "spa_sales_commission_percent": 0.0,
        })
        product = self._make_comm_product("Parent categ prod", 10.0, 2000000)
        product.taxes_id = [(6, 0, [])]
        so, _inv = self._make_so_with_invoice(
            [(product, 2000000)], invoice_date=date(2025, 7, 1), create_invoice=False,
        )
        self._redeem_points(so, 200)
        inv = so._create_invoices()
        inv.write({"invoice_date": date(2025, 7, 1), "invoice_user_id": self.user.id})
        inv.action_post()
        self._spa_pay_invoice(inv, date(2025, 7, 15))

        pay = self._payroll_for_month(2025, 7)
        pay.action_recompute_payroll_extras()
        comm = pay.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertAlmostEqual(comm.amount, (2000000.0 - 200000.0) * 0.10)
        self.assertFalse(
            pay.commission_line_ids.filtered(lambda l: l.product_id == self.reward_product)
        )

    def test_invoice_branch_prorates_order_level_discount(self):
        """Hồi quy: helper phân bổ từng trả 0 trên nhánh hóa đơn (display_type='product'
        trên AML khác SOL) — hóa đơn lẻ có CK toàn đơn phải trừ đúng vào HH."""
        product = self._make_comm_product("Standalone CK prod", 5.0, 1000000)
        product.taxes_id = [(6, 0, [])]
        ck = self._make_comm_product("Standalone CK bait", 5.0, 0)
        ck.taxes_id = [(6, 0, [])]
        inv = self.env["account.move"].with_context(
            spa_allow_invoice_without_so=True
        ).create({
            "move_type": "out_invoice",
            "partner_id": self.partner.id,
            "invoice_date": date(2025, 7, 1),
            "invoice_user_id": self.user.id,
            "company_id": self.company.id,
            "invoice_line_ids": [
                (0, 0, {
                    "product_id": product.id, "quantity": 1, "price_unit": 1000000,
                    "name": product.display_name,
                }),
                (0, 0, {
                    "product_id": ck.id, "quantity": 1, "price_unit": -100000,
                    "name": "CK HD", "is_global_discount": True,
                }),
            ],
        })
        inv.action_post()
        self._spa_pay_invoice(inv, date(2025, 7, 18))

        pay = self._payroll_for_month(2025, 7)
        pay.action_recompute_payroll_extras()
        comm = pay.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertAlmostEqual(comm.amount, (1000000.0 - 100000.0) * 0.05)

    def test_partial_credit_note_of_points_order_claws_back_net(self):
        """Hoàn 1 phần đơn dùng điểm, sang tháng sau: số hoàn = giá net, HH thu hồi net."""
        p_a = self._make_policy_comm_product("Refund Points A", 5.0, 500000, "order")
        p_b = self._make_policy_comm_product("Refund Points B", 5.0, 500000, "order")
        so = self._make_so_lines([(p_a, 500000, 1), (p_b, 500000, 1)])
        self._redeem_points(so, 100)  # -100,000đ trên tổng 1,000,000đ = -10%
        inv = self._invoice_and_post(so, date(2025, 8, 3))
        self._spa_pay_invoice(inv, date(2025, 8, 5))
        self._flush_settlement(so)

        pay_aug = self._payroll_for_month(2025, 8)
        pay_aug.action_recompute_payroll_extras()
        # net = (500000+500000-100000) × 5% = 45,000
        self.assertAlmostEqual(
            pay_aug.line_ids.filtered(lambda l: l.category == "sales_commission").amount,
            45000.0,
        )

        refund = self._partial_refund(inv, date(2025, 9, 3), qty_by_product={p_a: 1})
        refund_reward_lines = refund.invoice_line_ids.filtered(
            lambda l: l.spa_is_order_discount_line
        )
        self.assertTrue(refund_reward_lines, "credit note phải giữ/ sinh lại dòng điểm phân bổ")
        # Hoàn 1/2 hàng → phân bổ 1/2 số điểm: -50,000đ. Refund net = 500000-50000=450000.
        self.assertAlmostEqual(sum(refund_reward_lines.mapped("price_subtotal")), -50000.0)
        self._pay_open(refund, date(2025, 9, 10))
        self._flush_settlement(so)

        pay_sep = self._payroll_for_month(2025, 9)
        pay_sep.action_recompute_payroll_extras()
        self.assertAlmostEqual(pay_sep.kpi_revenue_base, -450000.0)
        self.assertAlmostEqual(
            pay_sep.line_ids.filtered(lambda l: l.category == "sales_commission").amount,
            -450000.0 * 0.05,
        )

    def test_full_credit_note_of_points_order_nets_to_zero(self):
        """Hoàn toàn bộ đơn dùng điểm trong cùng tháng → HH và KPI net = 0."""
        product = self._make_policy_comm_product("Full Refund Points", 5.0, 1000000, "order")
        so = self._make_so_lines([(product, 1000000, 1)])
        self._redeem_points(so, 100)
        inv = self._invoice_and_post(so, date(2025, 8, 3))
        refund = self._partial_refund(inv, date(2025, 8, 4))  # full reversal
        self._pay_open(inv, date(2025, 8, 5))
        self._pay_open(refund, date(2025, 8, 6))
        self._flush_settlement(so)

        pay = self._payroll_for_month(2025, 8)
        pay.action_recompute_payroll_extras()
        self.assertAlmostEqual(pay.kpi_revenue_base, 0.0)
        comm = pay.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertFalse(comm, "gross net 900,000 × 5% cộng thu hồi -900,000 × 5% = 0")

    def _make_card_from_sale_order(self, product, so, total_sessions, used_sessions):
        """Get (or create) the treatment card for this SO line.

        Posting an invoice for a service product with ``spa_sessions_per_unit``
        auto-creates a card (`spa.treatment.card._spa_create_treatment_card_if_paid`,
        `custom_addons/spa/models/account_move.py`), so reuse it instead of
        creating a duplicate (unique constraint on ``sale_order_line_id``).
        """
        sol = so.order_line.filtered(lambda l: l.product_id == product)[:1]
        card = self.env["spa.treatment.card"].search(
            [("sale_order_line_id", "=", sol.id)], limit=1
        )
        if card:
            card.write({"total_sessions": total_sessions, "lifecycle_status": "active"})
        else:
            card = self.env["spa.treatment.card"].create({
                "partner_id": so.partner_id.id,
                "product_id": product.id,
                "total_sessions": total_sessions,
                "sale_order_line_id": sol.id,
                "lifecycle_status": "active",
            })
        if used_sessions:
            self.env["spa.treatment.session"].create([
                {"card_id": card.id, "state": "done", "date": fields.Datetime.now()}
                for _ in range(used_sessions)
            ])
        card.invalidate_recordset()
        return card

    def test_card_return_credit_defaults_to_sold_price(self):
        """default_get lấy giá đã bán gốc trên dòng SO, không phải giá pricelist
        hiện tại (PAY-DEC-2026-09-03-01, mục C)."""
        product = self._make_comm_card_product("Card Credit Default", 5.0, 1000000, sessions=10)
        so, inv = self._make_so_with_invoice(
            [(product, 1000000)],
            discount=10.0,  # đơn gốc có CK dòng 10% → giá thực bán = 900,000
            invoice_date=date(2025, 6, 1),
        )
        card = self._make_card_from_sale_order(product, so, total_sessions=10, used_sessions=2)

        # Đổi giá bảng giá hiện tại — default KHÔNG được lấy giá này.
        product.list_price = 1500000

        wiz = self.env["spa.card.upgrade.wizard"].with_context(
            default_card_id=card.id
        ).create({
            "card_id": card.id,
            "partner_id": so.partner_id.id,
            "upgrade_option": "return",
        })
        self.assertAlmostEqual(wiz.credit_old_card, 900000.0, delta=0.01)

    def test_card_return_uses_original_salesperson(self):
        """SO trả thẻ lấy user_id từ salesperson đã bán thẻ gốc, không phải
        người phụ trách khách hàng hiện tại (PAY-DEC-2026-09-03-01, mục D)."""
        product = self._make_comm_card_product("Card Salesperson", 5.0, 1000000, sessions=10)
        so, inv = self._make_so_with_invoice(
            [(product, 1000000)],
            invoice_date=date(2025, 6, 1),
        )
        card = self._make_card_from_sale_order(product, so, total_sessions=10, used_sessions=2)

        other_user = self.env["res.users"].create({
            "name": "Other Salesperson",
            "login": "other_salesperson@example.com",
            "email": "other_salesperson@example.com",
        })
        so.partner_id.write({"user_id": other_user.id})

        wiz = self.env["spa.card.upgrade.wizard"].with_context(
            default_card_id=card.id
        ).create({
            "card_id": card.id,
            "partner_id": so.partner_id.id,
            "upgrade_option": "return",
            "amount_per_session": 100000.0,
            "credit_old_card": 1000000.0,
        })
        action = wiz.action_confirm()
        return_so = self.env["sale.order"].browse(action["res_id"])
        self.assertEqual(return_so.user_id, so.user_id)
        self.assertNotEqual(return_so.user_id, other_user)
        self.assertTrue(return_so.is_card_return_order)

    def test_card_return_commission_no_double_clawback(self):
        """Trả thẻ: claw-back hoa hồng đúng 1 lần, không bị tính trùng qua
        credit note liên kết (PAY-DEC-2026-09-03-01, mục B)."""
        product = self._make_comm_card_product("Card No Double Clawback", 5.0, 1000000, sessions=10)
        so, inv = self._make_so_with_invoice(
            [(product, 1000000)],
            invoice_date=date(2025, 6, 1),
        )
        self._spa_pay_invoice(inv, date(2025, 6, 5))
        so.invalidate_recordset()
        so._compute_spa_settlement()
        so.flush_recordset(["spa_is_settled", "spa_settled_date"])

        pay_jun = self._payroll_for_month(2025, 6)
        pay_jun.action_recompute_payroll_extras()
        jun_comm = pay_jun.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertAlmostEqual(jun_comm.amount, 1000000.0 * 0.05)
        pay_jun.action_confirm()

        card = self._make_card_from_sale_order(product, so, total_sessions=10, used_sessions=2)
        wiz = self.env["spa.card.upgrade.wizard"].with_context(
            default_card_id=card.id
        ).create({
            "card_id": card.id,
            "partner_id": so.partner_id.id,
            "upgrade_option": "return",
        })
        action = wiz.action_confirm()
        return_so = self.env["sale.order"].browse(action["res_id"])
        self.assertTrue(return_so.is_card_return_order)
        return_so.action_confirm()

        used_value = 2 * (1000000.0 / 10)
        refund_value = 1000000.0 - used_value

        # Trả thẻ được thanh toán ở THÁNG KHÁC (7) để đo riêng claw-back, tách
        # khỏi hoa hồng tháng 6. SO trả thẻ (tổng âm) tính qua nhánh "orders";
        # credit note của chính SO bị bỏ ở so_refunds (PAY-DEC-2026-09-17-01).
        return_inv = return_so._create_invoices(final=True)
        return_inv.write({"invoice_date": date(2025, 7, 1), "invoice_user_id": so.user_id.id})
        return_inv.action_post()
        self.assertEqual(return_inv.move_type, "out_refund")
        self._spa_pay_move(return_inv, date(2025, 7, 5))
        return_so.invalidate_recordset()

        # Tháng 6 (đã khóa) không tự sửa: hoa hồng gốc giữ nguyên.
        pay_jun.invalidate_recordset()
        self.assertAlmostEqual(
            pay_jun.line_ids.filtered(lambda l: l.category == "sales_commission").amount,
            1000000.0 * 0.05,
        )

        pay_jul = self._payroll_for_month(2025, 7)
        pay_jul.action_recompute_payroll_extras()
        comm2 = pay_jul.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertTrue(comm2)
        # Đúng 1 lần claw-back = -pct * refund_value, KHÔNG phải gấp đôi.
        self.assertAlmostEqual(comm2.amount, -refund_value * 0.05, delta=1.0)

    def _make_negative_upgrade_so(self, product, price, credit, pay_date, is_return=False):
        """SO kiểu «Nâng cấp thẻ» (wizard nhánh replace): gói mới + CK toàn đơn
        lớn hơn giá gói → tổng âm, HĐ duy nhất là credit note tự chuyển."""
        so = self.env["sale.order"].create({
            "partner_id": self.partner.id,
            "user_id": self.user.id,
            "company_id": self.company.id,
            "origin": "Nâng cấp thẻ TEST",
            "is_card_return_order": is_return,
        })
        line = self.env["sale.order.line"].create({
            "order_id": so.id,
            "product_id": product.id,
            "product_uom_qty": 1,
            "price_unit": price,
            "tax_id": [(6, 0, [])],
        })
        disc_wiz = self.env["sale.order.discount"].new({"sale_order_id": so.id})
        disc_vals = disc_wiz._prepare_discount_line_values(
            product=disc_wiz._get_discount_product(),
            amount=credit,
            taxes=line.tax_id,
            description="Khấu trừ giá trị quy đổi thẻ cũ",
        )
        disc_vals["is_global_discount"] = True
        self.env["sale.order.line"].create(disc_vals)
        so.action_confirm()
        self.assertLess(so.amount_total, 0)
        refund = so._create_invoices(final=True)
        refund.write({"invoice_date": pay_date, "invoice_user_id": self.user.id})
        refund.action_post()
        self.assertEqual(refund.move_type, "out_refund")
        self.assertTrue(refund._spa_is_order_self_refund())
        self._spa_pay_move(refund, pay_date)
        self._flush_settlement(so)
        self.assertEqual(so.spa_settled_date, pay_date)
        return so, refund

    def _assert_negative_so_counted_once(self, so, refund, payroll):
        self.assertAlmostEqual(payroll.kpi_revenue_base, so._spa_recognized_amount_untaxed())
        self.assertAlmostEqual(payroll.kpi_revenue_base, float(so.amount_untaxed))
        self.assertEqual(payroll.kpi_line_ids.sale_order_id, so)
        self.assertFalse(payroll.kpi_line_ids.move_id)
        self.assertFalse(payroll.commission_line_ids.filtered(lambda l: l.move_id == refund))
        comm = payroll.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertAlmostEqual(comm.amount, self._expected_so_commission(so), delta=1.0)
        self.assertLess(comm.amount, 0)

    def test_negative_upgrade_order_counted_once(self):
        """Nâng cấp thẻ xuống gói rẻ hơn: SO âm tính 1 lần, credit note của
        chính SO không bị trừ thêm (PAY-DEC-2026-09-17-01)."""
        product = self._make_comm_product("Upgrade Negative", 5.0, 400000)
        so, refund = self._make_negative_upgrade_so(
            product, 400000, 1000000, date(2025, 8, 2)
        )
        pay = self._payroll_for_month(2025, 8)
        pay.action_recompute_payroll_extras()
        self._assert_negative_so_counted_once(so, refund, pay)
        self.assertAlmostEqual(
            pay.line_ids.filtered(lambda l: l.category == "sales_commission").amount,
            -600000 * 0.05,
            delta=1.0,
        )

    def test_negative_card_return_order_counted_once_any_flag(self):
        """Đơn trả thẻ cũ có cờ NULL/False hay mới có cờ True đều tính 1 lần."""
        product = self._make_comm_product("Return Negative", 5.0, 300000)
        for month, flag in ((8, False), (9, True)):
            so, refund = self._make_negative_upgrade_so(
                product, 300000, 800000, date(2025, month, 3), is_return=flag
            )
            pay = self._payroll_for_month(2025, month)
            pay.action_recompute_payroll_extras()
            self._assert_negative_so_counted_once(so, refund, pay)

    def test_self_refund_no_late_refund_activity(self):
        product = self._make_comm_product("Upgrade No Activity", 5.0, 400000)
        pay_aug = self._payroll_for_month(2025, 8)
        pay_aug.action_confirm()
        self.assertEqual(pay_aug.state, "done")
        self._make_negative_upgrade_so(product, 400000, 1000000, date(2025, 8, 2))
        acts = self.env["mail.activity"].search([
            ("res_model", "=", "spa.staff.payroll"),
            ("res_id", "=", pay_aug.id),
            ("summary", "=", "Hoàn tiền sau khi phiếu lương đã khóa"),
        ])
        self.assertFalse(acts)

    def test_late_refund_activity_on_done_refund_month_payslip(self):
        product = self._make_comm_product("SO Refund Activity", 5.0, 100000)
        so, inv = self._make_so_with_invoice(
            [(product, 100000)],
            invoice_date=date(2025, 5, 10),
        )
        self._spa_pay_move(inv, date(2025, 7, 15))
        so.invalidate_recordset()
        so._compute_spa_settlement()
        so.flush_recordset(["spa_settled_date"])

        pay_oct = self._payroll_for_month(2025, 10)
        pay_oct.action_confirm()
        self.assertEqual(pay_oct.state, "done")

        refund = inv._reverse_moves(default_values_list=[{
            "date": date(2025, 10, 5),
            "invoice_date": date(2025, 10, 5),
            "ref": "Payroll refund activity",
        }], cancel=False)
        refund = refund[0]
        if refund.state != "posted":
            refund.action_post()
        self._spa_pay_move(refund, date(2025, 10, 20))

        acts = self.env["mail.activity"].search([
            ("res_model", "=", "spa.staff.payroll"),
            ("res_id", "=", pay_oct.id),
            ("summary", "=", "Hoàn tiền sau khi phiếu lương đã khóa"),
        ])
        self.assertTrue(acts, "Locked October payslip should get a todo activity")

    def test_late_refund_notify_failure_does_not_rollback_payment(self):
        """PAY-DEC-2026-09-03-01 follow-up: a bug in
        _spa_notify_late_refund_on_locked_payslips (called from the
        account.partial.reconcile hook, custom_addons/spa_staff_payroll/
        models/account_move_payroll.py) must never roll back the accounting
        reconcile/payment that triggered it — it's a best-effort side-effect,
        not part of the settlement itself."""
        product = self._make_comm_product("SO Refund Notify Isolation", 5.0, 100000)
        so, inv = self._make_so_with_invoice(
            [(product, 100000)],
            invoice_date=date(2025, 5, 10),
        )
        self._spa_pay_move(inv, date(2025, 7, 15))
        so.invalidate_recordset()
        so._compute_spa_settlement()
        so.flush_recordset(["spa_settled_date"])

        pay_oct = self._payroll_for_month(2025, 10)
        pay_oct.action_confirm()

        refund = inv._reverse_moves(default_values_list=[{
            "date": date(2025, 10, 5),
            "invoice_date": date(2025, 10, 5),
            "ref": "Payroll refund notify isolation",
        }], cancel=False)
        refund = refund[0]
        if refund.state != "posted":
            refund.action_post()

        with patch.object(
            type(self.env["spa.staff.payroll"]),
            "_spa_notify_late_refund_on_locked_payslips",
            side_effect=RuntimeError("boom - simulated bug in notify"),
        ):
            # Must NOT raise: the exception is isolated via cr.savepoint()
            # inside _spa_mark_settlement_recompute.
            self._spa_pay_move(refund, date(2025, 10, 20))

        self.assertTrue(
            self.company.currency_id.is_zero(refund.amount_residual),
            "Refund payment must still be fully reconciled despite the notify bug",
        )
        # No activity created this time (notify failed before creating any),
        # but the payment/reconcile itself committed successfully — that's
        # the property under test, not the activity.

    def test_late_refund_notify_flush_uses_narrow_model_flush(self):
        """The reconcile hook must flush only sale.order/account.move (the
        models the notify query reads), not env.flush_all() — regression
        guard so a future edit doesn't reintroduce the broad flush on this
        hot path (bulk bank-statement reconcile)."""
        from odoo.addons.spa_staff_payroll.models import account_move_payroll as amp

        source = inspect.getsource(amp.AccountMove._spa_mark_settlement_recompute)
        self.assertNotIn("flush_all", source)
        self.assertIn("flush_model", source)

    # --- Partial / full return closes SO for KPI + HH (PAY-DEC-2026-09-09-01) ---

    def _make_policy_comm_product(self, name, pct, price, policy):
        categ = self.env["product.category"].create({
            "name": "Categ %s" % name,
            "spa_sales_commission_percent": pct,
        })
        tmpl = self.env["product.template"].create({
            "name": name,
            "detailed_type": "service",
            "invoice_policy": policy,
            "categ_id": categ.id,
            "list_price": price,
            "taxes_id": [(5, 0, 0)],  # no tax -> hard-coded expected amounts
        })
        return tmpl.product_variant_id

    def _make_so_lines(self, line_specs, user=None):
        """line_specs: [(product, price_unit, qty), ...] -> confirmed SO."""
        user = user or self.user
        so = self.env["sale.order"].create({
            "partner_id": self.partner.id,
            "user_id": user.id,
            "company_id": self.company.id,
            "order_line": [
                (0, 0, {
                    "product_id": p.id,
                    "product_uom_qty": q,
                    "price_unit": pu,
                })
                for (p, pu, q) in line_specs
            ],
        })
        so.action_confirm()
        return so

    def _invoice_and_post(self, so, invoice_date, user=None):
        user = user or self.user
        inv = so._create_invoices()
        inv.write({"invoice_date": invoice_date, "invoice_user_id": user.id})
        inv.action_post()
        return inv

    def _partial_refund(self, inv, refund_date, qty_by_product=None, user=None):
        """Credit-note ``inv``; optionally cut line quantities before posting."""
        user = user or self.user
        moves = inv._reverse_moves(default_values_list=[{
            "date": refund_date,
            "invoice_date": refund_date,
            "invoice_user_id": user.id,
            "ref": "Return %s" % inv.name,
        }], cancel=False)
        refund = moves[0]
        if qty_by_product is not None:
            # Keep only the returned products, at the requested quantities.
            for line in refund.invoice_line_ids.filtered(
                lambda l: l.display_type == "product"
            ):
                if line.product_id in qty_by_product:
                    line.quantity = qty_by_product[line.product_id]
                else:
                    line.unlink()
        if refund.state != "posted":
            refund.action_post()
        return refund

    def _flush_settlement(self, so):
        so.invalidate_recordset()
        so._compute_spa_settlement()
        so.flush_recordset(["spa_is_settled", "spa_settled_date", "spa_settled_month"])

    def _pay_open(self, move, pay_date):
        """Pay a move only if it still has a residual.

        ``account.move._reverse_moves`` (spa override + core) reconciles a fresh
        credit note against an open source invoice on post, so a return posted
        *before* the invoice is paid already has residual 0 — the S53056 shape.
        Then only the invoice's remaining net residual needs a real payment.
        """
        move.invalidate_recordset()
        if not self.company.currency_id.is_zero(move.amount_residual):
            if move.move_type == "out_invoice":
                self._spa_pay_invoice(move, pay_date)
            else:
                self._spa_pay_move(move, pay_date)
        if hasattr(move, "_compute_spa_fully_paid_date"):
            move.invalidate_recordset()
            move._compute_spa_fully_paid_date()
            move.flush_recordset(["spa_fully_paid_date"])

    def test_partial_return_same_month_counts_net_kpi_hh(self):
        """Order-policy SO, 2 lines (no tax); a partial credit note is posted
        BEFORE the invoice is fully paid, so it reconciles against the open
        invoice, amount_to_invoice never returns to 0 and the SO is never even
        transiently _spa_order_is_settled() — the S53056 shape. All money lands
        the SAME month → draft payslip recognises net (gross − refund) for KPI
        and the sales_commission ledger line, with HARD-CODED expected amounts.

        gross = 2·100000 + 2·200000 = 600000 ; refund = 1·100000 = 100000
        net KPI = 500000
        net HH  = (2·100000·5%) + (2·200000·10%) − (1·100000·5%)
                = 10000 + 40000 − 5000 = 45000
        """
        p_a = self._make_policy_comm_product("Ret A", 5.0, 100000, "order")
        p_b = self._make_policy_comm_product("Ret B", 10.0, 200000, "order")
        so = self._make_so_lines([(p_a, 100000, 2), (p_b, 200000, 2)])
        inv = self._invoice_and_post(so, date(2025, 8, 3))

        refund = self._partial_refund(
            inv, date(2025, 8, 4), qty_by_product={p_a: 1}
        )
        self._pay_open(inv, date(2025, 8, 5))
        self._pay_open(refund, date(2025, 8, 6))
        self._flush_settlement(so)

        self.assertFalse(so._spa_order_is_settled())
        self.assertFalse(so.spa_is_settled)
        self.assertTrue(so._spa_order_is_closed_by_returns())
        self.assertEqual(so.spa_settled_date, date(2025, 8, 5))

        pay = self._payroll_for_month(2025, 8)
        pay.action_recompute_payroll_extras()

        self.assertAlmostEqual(pay.kpi_revenue_base, 500000.0)
        comm = pay.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertEqual(len(comm), 1)
        self.assertAlmostEqual(comm.amount, 45000.0)

    def test_full_return_same_month_nets_to_zero(self):
        """100% returned same month → net KPI and net HH exactly 0; SO gets a
        spa_settled_date but stays spa_is_settled = False."""
        product = self._make_policy_comm_product("Ret Full", 5.0, 200000, "order")
        so = self._make_so_lines([(product, 200000, 2)])
        inv = self._invoice_and_post(so, date(2025, 8, 3))

        refund = self._partial_refund(inv, date(2025, 8, 4))  # full reversal
        self._pay_open(inv, date(2025, 8, 5))
        self._pay_open(refund, date(2025, 8, 6))
        self._flush_settlement(so)

        self.assertFalse(so.spa_is_settled)
        self.assertTrue(so._spa_order_is_closed_by_returns())
        self.assertTrue(so.spa_settled_date)

        pay = self._payroll_for_month(2025, 8)
        pay.action_recompute_payroll_extras()

        self.assertAlmostEqual(pay.kpi_revenue_base, 0.0)
        comm = pay.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertFalse(comm, "gross 20000 − refund 20000 = 0 → no ledger line")

    def test_multi_partial_return_across_months(self):
        """Invoice fully paid Aug (settled, date frozen) then two returns paid to
        the customer in Sep and Oct. Each return is clawed back as an exact
        negative amount in its own paid month; the origin month never moves.

        line: 3·100000 = 300000 gross ; each return = 1·100000 = 100000
        Sep clawback KPI = −100000, HH = −100000·5% = −5000
        Oct clawback KPI = −100000, HH = −5000
        """
        product = self._make_policy_comm_product("Ret Multi", 5.0, 100000, "order")
        so = self._make_so_lines([(product, 100000, 3)])
        inv = self._invoice_and_post(so, date(2025, 8, 2))
        self._spa_pay_invoice(inv, date(2025, 8, 5))
        self._flush_settlement(so)
        self.assertEqual(so.spa_settled_date, date(2025, 8, 5))

        r1 = self._partial_refund(inv, date(2025, 9, 3), qty_by_product={product: 1})
        self._pay_open(r1, date(2025, 9, 10))
        r2 = self._partial_refund(inv, date(2025, 10, 3), qty_by_product={product: 1})
        self._pay_open(r2, date(2025, 10, 10))
        self._flush_settlement(so)

        self.assertEqual(so.spa_settled_date, date(2025, 8, 5), "origin frozen")
        self.assertFalse(so.spa_is_settled)

        pay_aug = self._payroll_for_month(2025, 8)
        pay_aug.action_recompute_payroll_extras()
        self.assertAlmostEqual(pay_aug.kpi_revenue_base, 300000.0)

        pay_sep = self._payroll_for_month(2025, 9)
        pay_sep.action_recompute_payroll_extras()
        self.assertAlmostEqual(pay_sep.kpi_revenue_base, -100000.0)
        self.assertAlmostEqual(
            pay_sep.line_ids.filtered(lambda l: l.category == "sales_commission").amount,
            -5000.0,
        )

        pay_oct = self._payroll_for_month(2025, 10)
        pay_oct.action_recompute_payroll_extras()
        self.assertAlmostEqual(pay_oct.kpi_revenue_base, -100000.0)
        self.assertAlmostEqual(
            pay_oct.line_ids.filtered(lambda l: l.category == "sales_commission").amount,
            -5000.0,
        )

        # Origin month unchanged after both clawbacks exist.
        pay_aug2 = self._payroll_for_month(2025, 8)
        pay_aug2.action_recompute_payroll_extras()
        self.assertAlmostEqual(pay_aug2.kpi_revenue_base, 300000.0)

    def test_return_delivery_policy_product(self):
        """invoice_policy='delivery', fully delivered + invoiced before the
        partial return: the amount check (amount_to_invoice - Σ refunds) behaves
        exactly like an order-policy line → closed-by-returns, net KPI/HH.

        The amount check keys off ``amount_total - invoiced_gross`` (0 once the
        order was fully invoiced), so a later physical restock (qty_delivered
        dropping) does NOT reopen it. The genuinely-limited case is a
        delivery-policy order that was NEVER fully delivered and then abandoned
        after returning its delivered portion — there really are undelivered
        units, so it correctly stays open (see
        test_over_return_not_closed_by_returns) and would fall back to the
        original bug only if the order is abandoned. SPA-DEC-2026-09-09-08.
        """
        product = self._make_policy_comm_product("Ret Deliv", 5.0, 200000, "delivery")
        so = self._make_so_lines([(product, 200000, 2)])
        so.order_line.qty_delivered = 2
        inv = self._invoice_and_post(so, date(2025, 8, 3))

        refund = self._partial_refund(inv, date(2025, 8, 4), qty_by_product={product: 1})
        self._pay_open(inv, date(2025, 8, 5))
        self._pay_open(refund, date(2025, 8, 6))
        self._flush_settlement(so)

        self.assertTrue(so._spa_order_is_closed_by_returns())
        self.assertFalse(so.spa_is_settled)
        self.assertEqual(so.spa_settled_date, date(2025, 8, 5))

        pay = self._payroll_for_month(2025, 8)
        pay.action_recompute_payroll_extras()
        # gross 2·200000 = 400000 ; refund 1·200000 = 200000 ; net 200000
        self.assertAlmostEqual(pay.kpi_revenue_base, 200000.0)

        # Physical restock of the returned unit: accounting unchanged -> still closed.
        so.order_line.qty_delivered = 1
        so.invalidate_recordset()
        self.assertTrue(so._spa_order_is_closed_by_returns())

    def test_return_mixed_policy_order(self):
        """One order-policy line + one fully-delivered delivery-policy line
        (no tax); partial return on the order-policy line. HARD-CODED amounts.

        gross = 2·100000 + 2·200000 = 600000 ; refund = 1·100000 = 100000
        net KPI = 500000
        net HH  = 2·100000·5% + 2·200000·10% − 1·100000·5%
                = 10000 + 40000 − 5000 = 45000
        """
        p_order = self._make_policy_comm_product("Ret Mix O", 5.0, 100000, "order")
        p_deliv = self._make_policy_comm_product("Ret Mix D", 10.0, 200000, "delivery")
        so = self._make_so_lines([(p_order, 100000, 2), (p_deliv, 200000, 2)])
        so.order_line.filtered(lambda l: l.product_id == p_deliv).qty_delivered = 2
        inv = self._invoice_and_post(so, date(2025, 8, 3))

        refund = self._partial_refund(
            inv, date(2025, 8, 4), qty_by_product={p_order: 1}
        )
        self._pay_open(inv, date(2025, 8, 5))
        self._pay_open(refund, date(2025, 8, 6))
        self._flush_settlement(so)

        self.assertTrue(so._spa_order_is_closed_by_returns())
        self.assertFalse(so.spa_is_settled)
        self.assertEqual(so.spa_settled_date, date(2025, 8, 5))

        pay = self._payroll_for_month(2025, 8)
        pay.action_recompute_payroll_extras()
        self.assertAlmostEqual(pay.kpi_revenue_base, 500000.0)
        comm = pay.line_ids.filtered(lambda l: l.category == "sales_commission")
        self.assertAlmostEqual(comm.amount, 45000.0)

    def test_action_draft_requires_payroll_manager(self):
        """Mở lại nháp phiếu đã duyệt chỉ dành cho quản lý lương.

        Nút "Xác nhận" vốn đã giới hạn nhóm ``group_spa_payroll_manager`` nhưng
        ``action_draft`` trước đây không guard gì — nhân viên không duyệt được phiếu
        nhưng lại mở lại nháp được phiếu ``done``, rồi bấm "Tính các khoản" sẽ dựng
        lại toàn bộ dòng hoa hồng từ dữ liệu hiện tại (giá / % có thể đã đổi).
        """
        payroll = self.env["spa.staff.payroll"].create({
            "employee_id": self.employee.id,
            "date_from": date(2026, 1, 1),
            "date_to": date(2026, 1, 31),
        })
        payroll.action_confirm()
        self.assertEqual(payroll.state, "done")

        staff = self.env["res.users"].create({
            "name": "Nhan vien khong phai quan ly luong",
            "login": "payroll_non_manager_test",
            "email": "payroll_non_manager@test.local",
            "groups_id": [(6, 0, [self.env.ref("base.group_user").id])],
        })
        with self.assertRaises(UserError):
            payroll.with_user(staff).action_draft()
        self.assertEqual(payroll.state, "done")

        manager = self.env["res.users"].create({
            "name": "Quan ly luong",
            "login": "payroll_manager_test",
            "email": "payroll_manager@test.local",
            "groups_id": [(6, 0, [
                self.env.ref("base.group_user").id,
                self.env.ref("spa_staff_payroll.group_spa_payroll_manager").id,
            ])],
        })
        payroll.with_user(manager).action_draft()
        self.assertEqual(payroll.state, "draft")
@tagged("post_install", "-at_install", "spa_security")
class TestSpaPayrollOperatorBookingForm(TransactionCase):
    def test_operator_form_payroll_flags_readonly(self):
        view = self.env.ref("booking_calendar.view_spa_service_booking_form_operator")
        arch = self.env["spa.service.booking"].get_view(view_id=view.id, view_type="form")["arch"]
        self.assertIn('name="spa_payroll_shift_kind"', arch)
        self.assertIn('name="spa_payroll_customer_requested"', arch)
        self.assertIn('readonly="1"', arch)


@tagged("post_install", "-at_install")
class TestSpaPayrollProductFormLayout(TransactionCase):
    def test_payroll_payout_table_before_composite_sub_services(self):
        view = self.env.ref("product.product_template_only_form_view")
        arch = self.env["product.template"].get_view(view_id=view.id, view_type="form")["arch"]
        payout_pos = arch.find('name="spa_payroll_profile_payout_line_ids"')
        sub_pos = arch.find('name="spa_sub_service_ids"')
        self.assertNotEqual(payout_pos, -1)
        self.assertNotEqual(sub_pos, -1)
        self.assertLess(payout_pos, sub_pos)
        self.assertIn('name="group_spa_payroll_profile"', arch)
        self.assertIn('name="group_spa_composite_service"', arch)




@tagged("post_install", "-at_install")
class TestSpaPayrollMultiBranch(TransactionCase):
    """1 người làm 2 chi nhánh = 2 hr.employee chung 1 res.users (PAY-DEC-2026-09-25-01/02)."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        if "hr.employee" not in cls.env.registry:
            raise SkipTest("hr is not installed in this database")
        cls.company_a = cls.env.company
        cls.company_b = cls.env["res.company"].create({
            "name": "Chi nhanh B (payroll test)",
            "currency_id": cls.company_a.currency_id.id,
        })
        both = [(6, 0, [cls.company_a.id, cls.company_b.id])]
        cls.user = cls.env["res.users"].create({
            "name": "Therapist Two Branches",
            "login": "therapist_two_branches_test",
            "email": "t2b@test.local",
            "company_id": cls.company_a.id,
            "company_ids": both,
            "groups_id": [(6, 0, [cls.env.ref("base.group_user").id])],
        })
        cls.emp_a = cls.env["hr.employee"].create({
            "name": "Therapist Two Branches",
            "user_id": cls.user.id,
            "company_id": cls.company_a.id,
        })
        cls.emp_b = cls.env["hr.employee"].create({
            "name": "Therapist Two Branches",
            "user_id": cls.user.id,
            "company_id": cls.company_b.id,
        })
        cls.partner = cls.env["res.partner"].create({"name": "Multi-branch Partner"})
        cls.product = cls.env["product.product"].create({
            "name": "Multi-branch Service",
            "detailed_type": "service",
            "list_price": 1000000,
            "service_employee_salary": 200000,
        })
        manager_group = cls.env.ref("spa_staff_payroll.group_spa_payroll_manager")
        cls.manager_both = cls.env["res.users"].create({
            "name": "Payroll Manager Both",
            "login": "payroll_manager_both_test",
            "email": "pmb@test.local",
            "company_id": cls.company_a.id,
            "company_ids": both,
            "groups_id": [(6, 0, [cls.env.ref("base.group_user").id, manager_group.id])],
        })
        cls.manager_a = cls.env["res.users"].create({
            "name": "Payroll Manager A",
            "login": "payroll_manager_a_test",
            "email": "pma@test.local",
            "company_id": cls.company_a.id,
            "company_ids": [(6, 0, [cls.company_a.id])],
            "groups_id": [(6, 0, [cls.env.ref("base.group_user").id, manager_group.id])],
        })

    def _payroll(self, employee, company):
        return self.env["spa.staff.payroll"].create({
            "employee_id": employee.id,
            "company_id": company.id,
            "currency_id": company.currency_id.id,
            "date_from": date(2026, 8, 1),
            "date_to": date(2026, 8, 31),
        })

    def _comm_line(self, payroll, amount, branch="spa"):
        return self.env["spa.staff.payroll.commission.line"].create({
            "payroll_id": payroll.id,
            "product_id": self.product.id,
            "product_branch": branch,
            "partner_id": self.partner.id,
            "settle_date": date(2026, 8, 10),
            "quantity": 1.0,
            "price_unit": amount * 10,
            "price_subtotal": amount * 10,
            "commission_percent": 10.0,
            "commission_amount": amount,
        })

    def _export_workbook(self, payslips, user):
        import base64
        import io
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise SkipTest("openpyxl is not installed")

        action = payslips.with_user(user).with_context(
            allowed_company_ids=[self.company_a.id]
        ).action_export_commission_xlsx_merged()
        export = self.env["spa.staff.payroll.commission.export"].browse(action["res_id"])
        self.assertTrue(export.file_name.endswith(".xlsx"))
        return load_workbook(io.BytesIO(base64.b64decode(export.file_data)))

    @staticmethod
    def _cells(ws):
        return [c for row in ws.iter_rows(values_only=True) for c in row if c not in (None, "")]

    def _total_commission(self, ws):
        for row in ws.iter_rows(values_only=True):
            if row and row[0] == "TỔNG CỘNG":
                return row[14]
        return None

    def test_session_pay_split_by_branch_no_double_count(self):
        so_b = self.env["sale.order"].create({
            "partner_id": self.partner.id,
            "company_id": self.company_b.id,
            "order_line": [(0, 0, {"product_id": self.product.id, "product_uom_qty": 1})],
        })
        card_b = self.env["spa.treatment.card"].create({
            "partner_id": self.partner.id,
            "product_id": self.product.id,
            "total_sessions": 10,
            "sale_order_line_id": so_b.order_line[:1].id,
        })
        card_no_so = self.env["spa.treatment.card"].create({
            "partner_id": self.partner.id,
            "product_id": self.product.id,
            "total_sessions": 10,
        })
        Session = self.env["spa.treatment.session"]
        for card in (card_b, card_no_so):
            Session.create({
                "card_id": card.id,
                "date": fields.Datetime.to_datetime(date(2026, 8, 15)),
                "therapist_id": self.user.id,
                "therapist_ids": [(6, 0, [self.user.id])],
                "duration_minutes": 60,
                "state": "done",
            })
        pay_a = self._payroll(self.emp_a, self.company_a)
        pay_b = self._payroll(self.emp_b, self.company_b)
        (pay_a | pay_b).action_recompute_service_lines()
        # Buổi thẻ có SO chi nhánh B → phiếu B; buổi không suy được → công ty mặc định user (A).
        self.assertEqual(pay_a.service_line_ids.session_id.card_id, card_no_so)
        self.assertEqual(pay_b.service_line_ids.session_id.card_id, card_b)

    def _done_session(self, card):
        return self.env["spa.treatment.session"].create({
            "card_id": card.id,
            "date": fields.Datetime.to_datetime(date(2026, 8, 15)),
            "therapist_id": self.user.id,
            "therapist_ids": [(6, 0, [self.user.id])],
            "duration_minutes": 60,
            "state": "done",
        })

    def _card_for_company(self, company):
        so = self.env["sale.order"].create({
            "partner_id": self.partner.id,
            "company_id": company.id,
            "order_line": [(0, 0, {"product_id": self.product.id, "product_uom_qty": 1})],
        })
        return self.env["spa.treatment.card"].create({
            "partner_id": self.partner.id,
            "product_id": self.product.id,
            "total_sessions": 10,
            "sale_order_line_id": so.order_line[:1].id,
        })

    def test_session_from_foreign_company_goes_to_fallback_not_lost(self):
        company_c = self.env["res.company"].create({
            "name": "Chi nhanh C (payroll test)",
            "currency_id": self.company_a.currency_id.id,
        })
        card_c = self._card_for_company(company_c)
        self._done_session(card_c)
        pay_a = self._payroll(self.emp_a, self.company_a)
        pay_b = self._payroll(self.emp_b, self.company_b)
        (pay_a | pay_b).action_recompute_service_lines()
        self.assertEqual(pay_a.service_line_ids.session_id.card_id, card_c)
        self.assertFalse(pay_b.service_line_ids)

    def test_single_branch_payslip_keeps_all_sessions(self):
        """Chưa có phiếu chi nhánh B → phiếu A giữ mọi buổi (không để buổi B không ai trả)."""
        self._done_session(self._card_for_company(self.company_b))
        self._done_session(self._card_for_company(self.company_a))
        pay_a = self._payroll(self.emp_a, self.company_a)
        pay_a.action_recompute_service_lines()
        self.assertEqual(len(pay_a.service_line_ids), 2)
        # Tạo phiếu B rồi tính lại cả hai → chia đúng, không trùng.
        pay_b = self._payroll(self.emp_b, self.company_b)
        (pay_a | pay_b).action_recompute_service_lines()
        self.assertEqual(len(pay_a.service_line_ids), 1)
        self.assertEqual(len(pay_b.service_line_ids), 1)

    def test_payslip_outside_employee_companies_keeps_old_behaviour(self):
        company_c = self.env["res.company"].create({
            "name": "Chi nhanh C2 (payroll test)",
            "currency_id": self.company_a.currency_id.id,
        })
        self._done_session(self._card_for_company(self.company_b))
        pay_c = self._payroll(self.emp_a, company_c)
        pay_c.action_recompute_service_lines()
        self.assertEqual(len(pay_c.service_line_ids), 1)

    def test_export_merges_commission_of_both_branches(self):
        pay_a = self._payroll(self.emp_a, self.company_a)
        pay_b = self._payroll(self.emp_b, self.company_b)
        self._comm_line(pay_a, 30000)
        self._comm_line(pay_a, 5000, branch="clinic")
        self._comm_line(pay_b, 70000)
        # Chỉ chọn phiếu A, switcher chỉ bật A — vẫn phải kéo phiếu B cùng user/kỳ.
        wb = self._export_workbook(pay_a, self.manager_both)
        self.assertEqual(len(wb.worksheets), 1)
        ws = wb.worksheets[0]
        cells = self._cells(ws)
        self.assertIn(self.company_a.name, cells)
        self.assertIn(self.company_b.name, cells)
        self.assertIn(pay_b.name, cells)
        self.assertEqual(self._total_commission(ws), 105000)
        self.assertFalse(any(isinstance(c, str) and c.startswith("Lưu ý") for c in cells))

    def test_export_respects_company_access(self):
        pay_a = self._payroll(self.emp_a, self.company_a)
        pay_b = self._payroll(self.emp_b, self.company_b)
        self._comm_line(pay_a, 30000)
        self._comm_line(pay_b, 70000)
        wb = self._export_workbook(pay_a, self.manager_a)
        ws = wb.worksheets[0]
        cells = self._cells(ws)
        self.assertNotIn(self.company_b.name, cells)
        self.assertEqual(self._total_commission(ws), 30000)
        self.assertTrue(any(isinstance(c, str) and c.startswith("Lưu ý: có 1 phiếu") for c in cells))

    def test_export_requires_payroll_manager(self):
        pay_a = self._payroll(self.emp_a, self.company_a)
        from odoo.exceptions import AccessError
        with self.assertRaises(AccessError):
            pay_a.with_user(self.user).action_export_commission_xlsx_merged()

    def test_export_payslip_without_user_only_itself(self):
        emp = self.env["hr.employee"].create({
            "name": "No User [Emp]/Test:*?",
            "company_id": self.company_a.id,
        })
        pay = self._payroll(emp, self.company_a)
        self._comm_line(pay, 12000)
        other = self._payroll(self.emp_a, self.company_a)
        self._comm_line(other, 99000)
        wb = self._export_workbook(pay | other, self.manager_both)
        self.assertEqual(len(wb.worksheets), 2)
        no_user_ws = [ws for ws in wb.worksheets if "No User" in ws.title]
        self.assertEqual(len(no_user_ws), 1)
        self.assertEqual(self._total_commission(no_user_ws[0]), 12000)
