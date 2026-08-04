# -*- coding: utf-8 -*-
from datetime import date
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from odoo.addons.rental.helper.transport_matrix_export import build_transport_matrix_xlsx_bytes
from odoo.addons.rental.helper.xlsx_template_utils import (
    find_transport_matrix_footer_row,
    find_transport_matrix_layout,
)
from odoo.tests.common import TransactionCase


class TestTransportMatrixXlsxExport(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        company = cls.env.company
        cls._a_company = cls.env["res.partner"].create({
            "name": "Matrix Export Renter",
            "is_company": True,
            "customer_type": "company",
            "is_rental_customer": True,
        })
        cls._a_party = cls.env["res.partner"].create({
            "name": "Matrix Export A rep",
            "parent_id": cls._a_company.id,
        })
        cls._b_party = cls.env["res.partner"].create({
            "name": "Matrix Export B rep",
            "parent_id": company.partner_id.id,
        })
        cls._driver = cls.env["res.partner"].create({
            "name": "Matrix Export Driver",
            "customer_type": "driver",
        })
        cls._truck = cls.env["transport.truck"].create({
            "plate": "29H-90001",
            "name": "29H-90001",
            "company_id": company.id,
        })
        cls._product = cls.env["product.product"].create({
            "name": "Khóa giáo export",
            "type": "product",
        })
        cls._contract = cls.env["rental.contract"].create({
            "a_company_party": cls._a_company.id,
            "a_party": cls._a_party.id,
            "b_company_party": company.partner_id.id,
            "b_party": cls._b_party.id,
        })
        cls.env["rental.contract.line"].create({
            "contract_id": cls._contract.id,
            "product_tmpl_id": cls._product.product_tmpl_id.id,
            "price_unit": 100.0,
        })

    def _create_transport(
        self, when, qty, plate="29H-90001", transport_type="delivery", product=None,
        state="done",
    ):
        return self.env["rr.transport"].create({
            "rental_contract_id": self._contract.id,
            "type": transport_type,
            "state": state,
            "start_rental_or_return_date": when,
            "driver_id": self._driver.id,
            "transport_truck_id": self._truck.id,
            "vehicle_start_time": when,
            "plate": plate,
            "transport_line_ids": [(0, 0, {
                "product_id": (product or self._product).id,
                "qty": qty,
            })],
        })

    def _export_workbook(self, start_date, end_date):
        data = build_transport_matrix_xlsx_bytes(
            self.env, self._contract, start_date, end_date,
        )
        return load_workbook(BytesIO(data)), data

    def test_period_placeholders_are_replaced(self):
        start_date = date(2026, 6, 1)
        end_date = date(2026, 6, 30)
        self._create_transport(date(2026, 6, 2), 10)
        wb, _data = self._export_workbook(start_date, end_date)
        ws = wb.active
        period_text = ws.cell(6, 1).value or ""
        self.assertIn("01/06/2026", period_text)
        self.assertIn("30/06/2026", period_text)
        self.assertNotIn("{{from_date}}", period_text)
        self.assertNotIn("{{to_date}}", period_text)

    def test_footer_rows_remain_visible_and_compact(self):
        start_date = date(2026, 6, 1)
        end_date = date(2026, 6, 30)
        self._create_transport(date(2026, 6, 2), 10)
        self._create_transport(date(2026, 6, 7), 20)
        wb, _data = self._export_workbook(start_date, end_date)
        ws = wb.active

        footer_row = find_transport_matrix_footer_row(ws, min_row=17)
        self.assertTrue(footer_row, "Footer row should be detected")
        self.assertFalse(getattr(ws.row_dimensions[footer_row], "hidden", False))
        self.assertFalse(getattr(ws.row_dimensions[footer_row + 1], "hidden", False))

        footer_values = []
        for row in (footer_row, footer_row + 1):
            for cell in ws[row]:
                if isinstance(cell.value, str):
                    footer_values.append(cell.value)
        joined = " ".join(footer_values)
        self.assertIn("Hà Nội", joined)
        self.assertIn("Đơn vị", joined)

        # Footer sits immediately below data (2 transports + total = 3 rows).
        _title, _h2, _h3, start_row = find_transport_matrix_layout(ws)
        self.assertEqual(footer_row, start_row + 3)

    def test_product_columns_keep_fixed_width(self):
        start_date = date(2026, 6, 1)
        end_date = date(2026, 6, 30)
        self._create_transport(date(2026, 6, 2), 2000)
        wb, _data = self._export_workbook(start_date, end_date)
        ws = wb.active
        dim = ws.column_dimensions[get_column_letter(4)]
        self.assertTrue(dim.width)
        self.assertFalse(dim.bestFit)

    def test_data_rows_have_no_product_column_merges(self):
        """Inserted/cloned template rows must not keep horizontal merges on qty cols."""
        start_date = date(2026, 6, 1)
        end_date = date(2026, 6, 30)
        # Enough trips to force insert_rows_below (template data area is small).
        for day, qty in ((2, 10), (5, 20), (8, 30), (12, 40), (15, 50), (20, 60)):
            self._create_transport(date(2026, 6, day), qty)
        # Opening balance row as well.
        self._create_transport(date(2026, 5, 20), 100)

        wb, _data = self._export_workbook(start_date, end_date)
        ws = wb.active
        _title, _h2, _h3, start_row = find_transport_matrix_layout(ws)
        # opening + 6 transports + total
        data_end = start_row + 7
        for rng in list(ws.merged_cells.ranges):
            if rng.min_row < start_row or rng.min_row > data_end:
                continue
            if rng.min_row != rng.max_row:
                continue
            # Only B:C label merges are allowed on data/total rows.
            if rng.min_col >= 4 or rng.max_col >= 4:
                self.fail(
                    "Unexpected product-column merge %s on data row %s"
                    % (rng, rng.min_row)
                )

    def test_return_quantities_are_negative_in_rows_and_balances(self):
        start_date = date(2026, 6, 1)
        end_date = date(2026, 6, 30)
        self._create_transport(date(2026, 5, 10), 100)
        self._create_transport(
            date(2026, 5, 20), 30, transport_type="return"
        )
        self._create_transport(date(2026, 6, 5), 20)
        self._create_transport(
            date(2026, 6, 10), 10, transport_type="return"
        )

        wb, _data = self._export_workbook(start_date, end_date)
        ws = wb.active
        _title, _h2, _h3, start_row = find_transport_matrix_layout(ws)
        self.assertEqual(ws.cell(start_row, 4).value, 70)
        self.assertEqual(ws.cell(start_row + 1, 4).value, 20)
        self.assertEqual(ws.cell(start_row + 2, 4).value, -10)
        self.assertEqual(ws.cell(start_row + 3, 4).value, 80)

        matrix = self.env["rental.transport.matrix"].create({
            "rental_contract_id": self._contract.id,
            "start_date": start_date,
            "end_date": end_date,
        })
        self.assertIn(">-10</td>", matrix.matrix_html)

    def test_return_quantity_keeps_negative_sign_in_total_md(self):
        md_category = self.env["uom.category"].create({
            "name": "MD matrix return export",
        })
        md_uom = self.env["uom.uom"].create({
            "name": "Mét dài matrix return export",
            "category_id": md_category.id,
            "uom_type": "reference",
            "is_linear_meter_variant": True,
        })
        md_product = self.env["product.template"].create({
            "name": "Hộp MD return export",
            "type": "product",
            "uom_id": md_uom.id,
            "uom_po_id": md_uom.id,
        }).product_variant_ids[0]
        self._create_transport(
            date(2026, 6, 10),
            7,
            transport_type="return",
            product=md_product,
        )

        wb, _data = self._export_workbook(date(2026, 6, 1), date(2026, 6, 30))
        ws = wb.active
        _title, _h2, header_third_row, start_row = find_transport_matrix_layout(ws)
        md_col = next(
            col
            for col in range(4, (ws.max_column or 4) + 1)
            if ws.cell(header_third_row, col).value == "Tổng MD"
        )
        self.assertEqual(ws.cell(start_row, md_col).value, -7)
        self.assertEqual(ws.cell(start_row + 1, md_col).value, -7)

    def test_draft_transport_before_period_excluded_from_opening(self):
        """DEC-17: draft (and non-done) trips must not inflate KLCT opening."""
        self._create_transport(date(2026, 5, 20), 40, state="draft")
        self._create_transport(date(2026, 5, 25), 10, state="done")
        self._create_transport(date(2026, 6, 5), 5, state="done")

        wb, _data = self._export_workbook(date(2026, 6, 1), date(2026, 6, 30))
        ws = wb.active
        _title, _h2, _header_third_row, start_row = find_transport_matrix_layout(ws)
        # Opening = done-only before period (10), not draft 40.
        self.assertEqual(ws.cell(start_row, 2).value, "Tồn đầu kỳ")
        self.assertEqual(ws.cell(start_row, 4).value, 10)
        # In-period done delivery only.
        self.assertEqual(ws.cell(start_row + 1, 4).value, 5)

        matrix = self.env["rental.transport.matrix"].create({
            "rental_contract_id": self._contract.id,
            "start_date": date(2026, 6, 1),
            "end_date": date(2026, 6, 30),
        })
        self.assertIn(">10</td>", matrix.matrix_html)
        self.assertNotIn(">40</td>", matrix.matrix_html)
        self.assertNotIn(">50</td>", matrix.matrix_html)
