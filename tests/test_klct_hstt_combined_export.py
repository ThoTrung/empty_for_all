# -*- coding: utf-8 -*-
from datetime import date
from io import BytesIO
import base64

from openpyxl import Workbook, load_workbook

from odoo.addons.rental.services import rental_contract_services as rcs
from odoo.addons.rental.helper.xlsx_template_utils import find_transport_matrix_layout
from odoo.exceptions import UserError
from odoo.tests import tagged
from odoo.tests.common import TransactionCase


@tagged("post_install", "-at_install")
class TestKlctHsttCombinedExport(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        company = cls.env.company
        cls._a_company = cls.env["res.partner"].create({
            "name": "KLCT HSTT Renter",
            "is_company": True,
            "customer_type": "company",
            "is_rental_customer": True,
        })
        cls._a_party = cls.env["res.partner"].create({
            "name": "KLCT HSTT A rep",
            "parent_id": cls._a_company.id,
        })
        cls._b_party = cls.env["res.partner"].create({
            "name": "KLCT HSTT B rep",
            "parent_id": company.partner_id.id,
        })
        cls._driver = cls.env["res.partner"].create({
            "name": "KLCT HSTT Driver",
            "customer_type": "driver",
        })
        cls._truck = cls.env["transport.truck"].create({
            "plate": "29H-91001",
            "name": "29H-91001",
            "company_id": company.id,
        })
        cls._product = cls.env["product.product"].create({
            "name": "Khóa giáo KLCT-HSTT",
            "type": "product",
            "list_price": 100.0,
        })
        # Ensure income account so invoice creation works in tests.
        income = cls.env["account.account"].search([
            ("company_id", "=", company.id),
            ("account_type", "=", "income"),
        ], limit=1)
        if not income:
            income = cls.env["account.account"].create({
                "name": "KLCT HSTT Income",
                "code": "KLCTHSTTINC",
                "account_type": "income",
                "company_id": company.id,
            })
        cls._product.property_account_income_id = income.id
        cls._product.product_tmpl_id.property_account_income_id = income.id

        cls._contract = cls.env["rental.contract"].create({
            "a_company_party": cls._a_company.id,
            "a_party": cls._a_party.id,
            "b_company_party": company.partner_id.id,
            "b_party": cls._b_party.id,
            "rental_billing_mode": "day",
            "include_start_day_current": False,
            "include_start_day_bob": True,
        })
        cls.env["rental.contract.line"].create({
            "contract_id": cls._contract.id,
            "product_tmpl_id": cls._product.product_tmpl_id.id,
            "price_unit": 100.0,
        })
        # Avoid DB-uploaded HSTT layouts (may merge row 13); use module static + start_row 13.
        cls.env["rental.template"].with_context(active_test=False).search([
            ("template_type", "=", "rental_invoice_xlsx"),
        ]).write({"active": False, "is_default": False})

    def _create_transport(self, when, qty, transport_type="delivery"):
        vals = {
            "rental_contract_id": self._contract.id,
            "type": transport_type,
            "state": "done",
            "start_rental_or_return_date": when,
            "driver_id": self._driver.id,
            "transport_truck_id": self._truck.id,
            "vehicle_start_time": when,
            "plate": "29H-91001",
            "transport_line_ids": [(0, 0, {
                "product_id": self._product.id,
                "qty": qty,
            })],
        }
        if "source_type" in self.env["rr.transport"]._fields:
            vals["source_type"] = "owned"
        return self.env["rr.transport"].create(vals)

    def test_draft_only_transports_raise_user_error_not_merge_crash(self):
        """No done product movements → UserError (RC00038-style), not openpyxl ValueError."""
        start_date = date(2026, 7, 1)
        end_date = date(2026, 7, 31)
        vals = {
            "rental_contract_id": self._contract.id,
            "type": "delivery",
            "state": "draft",
            "start_rental_or_return_date": date(2026, 7, 11),
            "driver_id": self._driver.id,
            "transport_truck_id": self._truck.id,
            "vehicle_start_time": date(2026, 7, 11),
            "plate": "29H-91001",
            "transport_line_ids": [(0, 0, {
                "product_id": self._product.id,
                "qty": 10,
            })],
        }
        if "source_type" in self.env["rr.transport"]._fields:
            vals["source_type"] = "owned"
        self.env["rr.transport"].create(vals)

        with self.assertRaises(UserError) as cm:
            self._contract.action_export_klct_hstt_excel(start_date, end_date)
        self.assertIn("hoàn thành", str(cm.exception).lower())

        with self.assertRaises(UserError):
            self._contract._build_klct_hstt_xlsx_buffer(start_date, end_date)

        self.assertFalse(
            self.env["rental.transport.matrix"].search_count([
                ("rental_contract_id", "=", self._contract.id),
                ("start_date", "=", start_date),
                ("end_date", "=", end_date),
            ]),
            "must not create matrix when guard blocks export",
        )
        self.assertFalse(
            self.env["account.move"].search_count([
                ("rental_contract_id", "=", self._contract.id),
                ("rental_start_date", "=", start_date),
                ("rental_end_date", "=", end_date),
            ]),
            "must not create invoice when guard blocks export",
        )

    def test_combined_workbook_has_two_named_sheets_with_formulas(self):
        start_date = date(2026, 4, 1)
        end_date = date(2026, 4, 30)
        self._create_transport(date(2026, 4, 10), 50)

        buffer, _subtotal = self._contract._build_klct_hstt_xlsx_buffer(start_date, end_date)
        wb = load_workbook(BytesIO(buffer.getvalue()))
        self.assertEqual(wb.sheetnames[0], "KLCT 04-2026")
        self.assertEqual(wb.sheetnames[1], "HSTT 04-2026")
        self.assertIn("ĐCCN 04-2026", wb.sheetnames)

        hstt = wb["HSTT 04-2026"]
        # Data starts at row 13; first billing line should use Excel formulas.
        qty_cell = hstt.cell(13, 6).value
        days_cell = hstt.cell(13, 7).value
        amount_cell = hstt.cell(13, 9).value
        self.assertIsInstance(qty_cell, str)
        self.assertTrue(str(qty_cell).startswith("="))
        self.assertIn("KLCT 04-2026", str(qty_cell))
        self.assertEqual(days_cell, "=C13-B13")
        self.assertEqual(amount_cell, "=F13*G13*H13")
        # Unit price: no forced decimal (template used to be #,##0.0 → 9,200.0).
        self.assertEqual(hstt.cell(13, 8).number_format, "#,##0")

        # Dates must be real date values for C-B formulas.
        self.assertEqual(hstt.cell(13, 2).value.date(), date(2026, 4, 10))
        self.assertEqual(hstt.cell(13, 3).value.date(), date(2026, 4, 30))

    def test_orphan_return_is_negative_in_klct_and_hstt_reference(self):
        start_date = date(2026, 4, 1)
        end_date = date(2026, 4, 30)
        self._create_transport(date(2026, 4, 10), 12, transport_type="return")
        self._create_transport(date(2026, 4, 11), 20, transport_type="delivery")

        buffer, _subtotal = self._contract._build_klct_hstt_xlsx_buffer(
            start_date, end_date
        )
        wb = load_workbook(BytesIO(buffer.getvalue()))
        klct = wb["KLCT 04-2026"]
        _title, _h2, _h3, start_row = find_transport_matrix_layout(klct)
        self.assertEqual(klct.cell(start_row, 4).value, -12)
        self.assertEqual(klct.cell(start_row + 1, 4).value, 20)
        self.assertEqual(klct.cell(start_row + 2, 4).value, 8)

        hstt = wb["HSTT 04-2026"]
        self.assertEqual(
            hstt.cell(13, 6).value,
            f"='KLCT 04-2026'!D{start_row}",
        )
        self.assertEqual(
            hstt.cell(14, 6).value,
            f"='KLCT 04-2026'!D{start_row + 1}",
        )

    def test_opening_qty_with_period_return_is_not_double_counted(self):
        """DEC-19 credit layout: full BOB present + credit return with F/G/I formulas."""
        self._contract.rental_billing_mode = "month"
        self._contract.minimum_rental_months = 0
        start_date = date(2026, 4, 1)
        end_date = date(2026, 4, 30)
        self._create_transport(date(2026, 3, 10), 100)
        self._create_transport(date(2026, 4, 10), 30, transport_type="return")

        buffer, _subtotal = self._contract._build_klct_hstt_xlsx_buffer(
            start_date, end_date
        )
        wb = load_workbook(BytesIO(buffer.getvalue()))
        klct = wb["KLCT 04-2026"]
        _title, _h2, _h3, start_row = find_transport_matrix_layout(klct)
        self.assertEqual(klct.cell(start_row, 4).value, 100)
        self.assertEqual(klct.cell(start_row + 1, 4).value, -30)
        self.assertEqual(klct.cell(start_row + 2, 4).value, 70)

        hstt = wb["HSTT 04-2026"]
        # Credit layout: BOB billed as full opening (100), then credit row for return.
        self.assertEqual(
            hstt.cell(13, 6).value,
            f"='KLCT 04-2026'!D{start_row}",
        )
        self.assertAlmostEqual(hstt.cell(13, 8).value, 100.0, places=6)
        self.assertEqual(hstt.cell(13, 9).value, "=F13*G13*H13/30")

        credit_row = None
        for row in range(13, 40):
            content = hstt.cell(row, 4).value or ""
            if "trả hàng" in str(content):
                credit_row = row
                break
        self.assertIsNotNone(credit_row, "expected credit return row on HSTT")
        self.assertEqual(
            hstt.cell(credit_row, 6).value,
            f"='KLCT 04-2026'!D{start_row + 1}",
        )
        self.assertEqual(hstt.cell(credit_row, 7).value, f"=C{credit_row}-B{credit_row}")
        self.assertAlmostEqual(hstt.cell(credit_row, 8).value, 100.0, places=6)
        self.assertEqual(
            hstt.cell(credit_row, 9).value,
            f"=F{credit_row}*G{credit_row}*H{credit_row}/30",
        )

    def test_draft_transport_excluded_from_klct_opening_and_hstt_bob(self):
        """DEC-17: draft before period must not enter KLCT opening / HSTT BOB."""
        self._contract.rental_billing_mode = "month"
        start_date = date(2026, 5, 1)
        end_date = date(2026, 5, 31)
        self._create_transport(date(2026, 4, 10), 232)  # done (helper default)
        draft_vals = {
            "rental_contract_id": self._contract.id,
            "type": "delivery",
            "state": "draft",
            "start_rental_or_return_date": date(2026, 4, 15),
            "driver_id": self._driver.id,
            "transport_truck_id": self._truck.id,
            "vehicle_start_time": date(2026, 4, 15),
            "plate": "29H-91001",
            "transport_line_ids": [(0, 0, {
                "product_id": self._product.id,
                "qty": 1000,
            })],
        }
        if "source_type" in self.env["rr.transport"]._fields:
            draft_vals["source_type"] = "owned"
        draft = self.env["rr.transport"].create(draft_vals)
        self.assertEqual(draft.state, "draft")

        buffer, _subtotal = self._contract._build_klct_hstt_xlsx_buffer(
            start_date, end_date
        )
        wb = load_workbook(BytesIO(buffer.getvalue()))
        klct = wb["KLCT 05-2026"]
        _title, _h2, _h3, start_row = find_transport_matrix_layout(klct)
        self.assertEqual(klct.cell(start_row, 2).value, "Tồn đầu kỳ")
        self.assertEqual(klct.cell(start_row, 4).value, 232)

        hstt = wb["HSTT 05-2026"]
        self.assertEqual(hstt.cell(13, 6).value, f"='KLCT 05-2026'!D{start_row}")

    def test_same_day_delivery_and_return_use_direction_specific_rows(self):
        start_date = date(2026, 4, 1)
        end_date = date(2026, 4, 30)
        # Return first makes it an orphan; the later delivery remains a normal lot.
        self._create_transport(date(2026, 4, 10), 5, transport_type="return")
        self._create_transport(date(2026, 4, 10), 9, transport_type="delivery")

        buffer, _subtotal = self._contract._build_klct_hstt_xlsx_buffer(
            start_date, end_date
        )
        wb = load_workbook(BytesIO(buffer.getvalue()))
        klct = wb["KLCT 04-2026"]
        _title, _h2, _h3, start_row = find_transport_matrix_layout(klct)
        self.assertEqual(klct.cell(start_row, 4).value, -5)
        self.assertEqual(klct.cell(start_row + 1, 4).value, 9)
        self.assertEqual(klct.cell(start_row + 2, 4).value, 4)

        hstt = wb["HSTT 04-2026"]
        self.assertEqual(
            hstt.cell(13, 6).value,
            (
                f"=SUM('KLCT 04-2026'!D{start_row},"
                f"'KLCT 04-2026'!D{start_row + 1})"
            ),
        )
        self.assertEqual(hstt.cell(14, 6).value, 4)

    def test_standalone_invoice_export_still_writes_literal_values(self):
        start_date = date(2026, 4, 1)
        end_date = date(2026, 4, 30)
        self._create_transport(date(2026, 4, 10), 50)

        buffer = self._contract._build_rental_payment_xlsx_buffer(start_date, end_date)
        wb = load_workbook(BytesIO(buffer.getvalue()))
        ws = wb.active
        self.assertEqual(ws.cell(13, 6).value, 50)
        self.assertIsInstance(ws.cell(13, 7).value, (int, float))
        self.assertIsInstance(ws.cell(13, 9).value, (int, float))
        self.assertIsInstance(ws.cell(13, 2).value, str)

    def test_include_start_day_current_adds_plus_one_in_formula(self):
        self._contract.include_start_day_current = True
        start_date = date(2026, 4, 1)
        end_date = date(2026, 4, 30)
        self._create_transport(date(2026, 4, 10), 20)

        buffer, _subtotal = self._contract._build_klct_hstt_xlsx_buffer(start_date, end_date)
        wb = load_workbook(BytesIO(buffer.getvalue()))
        hstt = wb["HSTT 04-2026"]
        self.assertEqual(hstt.cell(13, 7).value, "=C13-B13+1")

    def test_hstt_days_formula_subtracts_holidays(self):
        """Excel G must subtract holiday days so SUM(I) matches invoice Python total."""
        self._contract.include_start_day_bob = True
        self._contract.include_start_day_current = False
        # Opening stock before period → bob line [01/02 → 28/02] with include start.
        self._create_transport(date(2026, 1, 15), 10)
        self.env["rental.holiday"].create({
            "name": "Nghỉ Tết KLCT-HSTT",
            "date_from": date(2026, 2, 14),
            "date_to": date(2026, 2, 20),
            "company_id": self.env.company.id,
        })
        start_date = date(2026, 2, 1)
        end_date = date(2026, 2, 28)

        buffer, hstt_subtotal = self._contract._build_klct_hstt_xlsx_buffer(start_date, end_date)
        wb = load_workbook(BytesIO(buffer.getvalue()))
        hstt = wb["HSTT 02-2026"]
        # 7 holiday days on [B,C] with include_start → MAX(0,C-B+1-7)
        self.assertEqual(hstt.cell(13, 7).value, "=MAX(0,C13-B13+1-7)")

        bob_map, map_map = rcs.calc_rental_payment_table_grouped_by_template(
            self.env, self._contract, start_date, end_date
        )
        fee_lines = self._contract._rental_period_transport_fee_lines(start_date, end_date)
        expected = float(int(round(
            self._contract._payment_table_subtotal(bob_map, map_map, fee_lines)
        )))
        self.assertAlmostEqual(float(int(round(hstt_subtotal))), expected, places=2)

        # Invoice from export must match the same holiday-aware subtotal.
        self.env["rental.transport.matrix"].search([
            ("rental_contract_id", "=", self._contract.id),
        ]).unlink()
        action = self._contract.action_export_klct_hstt_excel(start_date, end_date)
        self.assertEqual(action["type"], "ir.actions.act_url")
        move = self.env["account.move"].search([
            ("rental_contract_id", "=", self._contract.id),
        ], limit=1)
        self.assertTrue(move)
        self.assertAlmostEqual(move.invoice_line_ids.price_unit, expected, places=2)

    def test_action_export_klct_hstt_excel_creates_matrix_and_one_line_invoice(self):
        start_date = date(2026, 4, 1)
        end_date = date(2026, 4, 30)
        self._create_transport(date(2026, 4, 10), 10)

        bob_map, map_map = rcs.calc_rental_payment_table_grouped_by_template(
            self.env, self._contract, start_date, end_date
        )
        fee_lines = self._contract._rental_period_transport_fee_lines(start_date, end_date)
        expected = float(int(round(
            self._contract._payment_table_subtotal(bob_map, map_map, fee_lines)
        )))

        action = self._contract.action_export_klct_hstt_excel(start_date, end_date)
        self.assertEqual(action["type"], "ir.actions.act_url")
        self.assertIn("/web/content/", action["url"])
        self.assertIn("download=1", action["url"])

        matrices = self.env["rental.transport.matrix"].search([
            ("rental_contract_id", "=", self._contract.id),
            ("start_date", "=", start_date),
            ("end_date", "=", end_date),
        ])
        self.assertEqual(len(matrices), 1)

        moves = self.env["account.move"].search([
            ("rental_contract_id", "=", self._contract.id),
        ])
        self.assertEqual(len(moves), 1)
        self.assertEqual(moves.state, "posted")
        lines = moves.invoice_line_ids
        self.assertEqual(len(lines), 1)
        self.assertEqual(lines.product_id.name, "Tổng thanh toán: 04-2026")
        self.assertIn("Tổng thanh toán: 04-2026", lines.name)
        self.assertEqual(lines.quantity, 1)
        self.assertAlmostEqual(lines.price_unit, expected, places=2)
        self.assertAlmostEqual(moves.amount_untaxed, expected, places=2)
        self.assertTrue(lines.tax_ids)
        self.assertAlmostEqual(lines.tax_ids[0].amount, 8.0, places=1)

        # Combined attachment includes ĐCCN with this period as phát sinh.
        att = moves.business_xlsx_attachment_id
        self.assertTrue(att)
        wb = load_workbook(BytesIO(base64.b64decode(att.datas)))
        self.assertIn("ĐCCN 04-2026", wb.sheetnames)
        dccn = wb["ĐCCN 04-2026"]
        self.assertAlmostEqual(float(dccn.cell(17, 8).value or 0), moves.amount_total, places=2)
        self.assertAlmostEqual(float(dccn.cell(16, 8).value or 0), 0.0, places=2)

    def test_hstt_total_product_replaces_stale_tax_with_eight_percent(self):
        tax_10 = self.env["account.tax"].search([
            ("company_id", "=", self.env.company.id),
            ("type_tax_use", "=", "sale"),
            ("amount", "=", 10.0),
            ("amount_type", "=", "percent"),
        ], limit=1)
        if not tax_10:
            tax_10 = self.env["account.tax"].create({
                "name": "10% HSTT stale test",
                "amount": 10.0,
                "amount_type": "percent",
                "type_tax_use": "sale",
                "company_id": self.env.company.id,
            })
        product = self._contract._get_or_create_hstt_total_product(date(2026, 4, 30))
        product.taxes_id = [(6, 0, tax_10.ids)]

        product = self._contract._get_or_create_hstt_total_product(date(2026, 4, 30))

        self.assertEqual(product.taxes_id.amount, 8.0)

    def test_action_export_klct_hstt_excel_overlap_returns_wizard(self):
        start_date = date(2026, 4, 1)
        end_date = date(2026, 4, 30)
        self._create_transport(date(2026, 4, 10), 10)
        self.env["rental.transport.matrix"].create({
            "rental_contract_id": self._contract.id,
            "start_date": start_date,
            "end_date": end_date,
            "name": "Existing matrix",
        })
        moves_before = self.env["account.move"].search_count([
            ("rental_contract_id", "=", self._contract.id),
        ])
        action = self._contract.action_export_klct_hstt_excel(start_date, end_date)
        self.assertEqual(action["type"], "ir.actions.act_window")
        self.assertEqual(action["res_model"], "rental.transport.matrix.overlap.wizard")
        moves_after = self.env["account.move"].search_count([
            ("rental_contract_id", "=", self._contract.id),
        ])
        self.assertEqual(moves_before, moves_after)

    def _create_transport_with_fee(self, when, qty, fee):
        transport = self._create_transport(when, qty)
        transport.fee = fee
        return transport

    def test_unbilled_fee_filter_by_cutoff_and_marks_on_export(self):
        start_date = date(2026, 6, 1)
        end_date = date(2026, 6, 30)
        cutoff = date(2026, 6, 30)
        t1 = self._create_transport_with_fee(date(2026, 6, 2), 5, 1_000_000)
        t2 = self._create_transport_with_fee(date(2026, 6, 7), 5, 1_000_000)
        t3 = self._create_transport_with_fee(date(2026, 6, 29), 5, 1_000_000)
        # Future trip beyond cutoff — must not be billed yet.
        t_future = self._create_transport_with_fee(date(2026, 7, 5), 5, 500_000)

        fee_lines = self._contract._rental_unbilled_transport_fee_lines(cutoff)
        self.assertEqual(len(fee_lines), 3)
        self.assertAlmostEqual(sum(fl["amount"] for fl in fee_lines), 3_000_000)

        empty_lines = self._contract._rental_unbilled_transport_fee_lines(False)
        self.assertEqual(empty_lines, [])

        bob_map, map_map = rcs.calc_rental_payment_table_grouped_by_template(
            self.env, self._contract, start_date, end_date
        )
        expected = float(int(round(
            self._contract._payment_table_subtotal(bob_map, map_map, fee_lines)
        )))

        action = self._contract.action_export_klct_hstt_excel(
            start_date, end_date, transport_fee_until_date=cutoff
        )
        self.assertEqual(action["type"], "ir.actions.act_url")

        for t in (t1, t2, t3):
            self.assertEqual(t.fee_billed_date, cutoff)
            self.assertTrue(t.fee_invoice_id)
        self.assertFalse(t_future.fee_billed_date)
        self.assertFalse(t_future.fee_invoice_id)

        move = self.env["account.move"].search([
            ("rental_contract_id", "=", self._contract.id),
        ], limit=1)
        self.assertEqual(move.transport_fee_until_date, cutoff)
        self.assertAlmostEqual(move.invoice_line_ids.price_unit, expected, places=2)

        # Second period must not pick up already-billed fees.
        later_fees = self._contract._rental_unbilled_transport_fee_lines(date(2026, 7, 31))
        self.assertEqual(len(later_fees), 1)
        self.assertEqual(later_fees[0]["transport_id"], t_future.id)

    def test_hstt_fee_section_one_row_plus_detail_sheet(self):
        start_date = date(2026, 6, 1)
        end_date = date(2026, 6, 30)
        self._create_transport_with_fee(date(2026, 6, 2), 10, 1_000_000)
        self._create_transport_with_fee(date(2026, 6, 7), 10, 1_000_000)
        self._create_transport_with_fee(date(2026, 6, 29), 10, 1_000_000)

        buffer, subtotal = self._contract._build_klct_hstt_xlsx_buffer(
            start_date, end_date, transport_fee_until_date=end_date
        )
        wb = load_workbook(BytesIO(buffer.getvalue()))
        self.assertIn("Chi tiết phí VC", wb.sheetnames)
        detail = wb["Chi tiết phí VC"]
        self.assertEqual(detail.cell(1, 1).value, "Ngày")
        # 3 fee rows + header + total row
        self.assertEqual(detail.cell(2, 4).value, 1_000_000)
        self.assertEqual(detail.cell(4, 4).value, 1_000_000)
        self.assertEqual(detail.cell(5, 3).value, "Tổng")
        self.assertEqual(detail.cell(5, 4).value, 3_000_000)

        hstt = wb["HSTT 06-2026"]
        fee_header_row = None
        for r in range(13, 40):
            if hstt.cell(r, 4).value == "Phí vận chuyển":
                fee_header_row = r
                break
        self.assertIsNotNone(fee_header_row)
        summary_row = fee_header_row + 1
        self.assertEqual(hstt.cell(summary_row, 4).value, "Phí vận chuyển (3 chuyến)")
        self.assertEqual(hstt.cell(summary_row, 9).value, 3_000_000)
        # No per-transport fee rows after the summary.
        self.assertNotEqual(
            hstt.cell(summary_row + 1, 4).value or "",
            "Phí vận chuyển TR",
        )
        self.assertGreaterEqual(subtotal, 3_000_000)

    def test_regenerate_klct_hstt_keeps_invoice_fees(self):
        start_date = date(2026, 6, 1)
        end_date = date(2026, 6, 30)
        t1 = self._create_transport_with_fee(date(2026, 6, 2), 10, 800_000)
        self._contract.action_export_klct_hstt_excel(
            start_date, end_date, transport_fee_until_date=end_date
        )
        move = t1.fee_invoice_id
        self.assertTrue(move)

        action = move.action_regenerate_business_xlsx()
        self.assertEqual(action["type"], "ir.actions.act_url")
        att = move.business_xlsx_attachment_id
        self.assertTrue(att)
        wb = load_workbook(BytesIO(base64.b64decode(att.datas)))
        self.assertIn("Chi tiết phí VC", wb.sheetnames)
        detail = wb["Chi tiết phí VC"]
        self.assertEqual(detail.cell(2, 4).value, 800_000)

        # Unbilled query must be empty for this trip after billing.
        self.assertEqual(
            self._contract._rental_unbilled_transport_fee_lines(end_date),
            [],
        )

    def test_download_business_xlsx_uses_cached_attachment(self):
        """«Tải» must not rebuild — attachment datas stay unchanged."""
        start_date = date(2026, 5, 1)
        end_date = date(2026, 5, 31)
        self._create_transport(date(2026, 5, 10), 12)
        self._contract.action_export_klct_hstt_excel(start_date, end_date)
        move = self.env["account.move"].search([
            ("rental_contract_id", "=", self._contract.id),
            ("rental_start_date", "=", start_date),
            ("rental_end_date", "=", end_date),
        ], limit=1)
        self.assertTrue(move.business_xlsx_attachment_id)
        before = move.business_xlsx_attachment_id.datas
        before_write_date = move.business_xlsx_attachment_id.write_date

        action = move.action_download_business_xlsx()
        self.assertEqual(action["type"], "ir.actions.act_url")
        self.assertIn(
            f"/web/content/{move.business_xlsx_attachment_id.id}",
            action["url"],
        )
        self.assertEqual(move.business_xlsx_attachment_id.datas, before)
        self.assertEqual(move.business_xlsx_attachment_id.write_date, before_write_date)

    def _make_hstt_template_bytes_start_row_15(self):
        """Minimal HSTT workbook: confirmation merge on row 13, header 14, data from 15."""
        wb = Workbook()
        ws = wb.active
        ws.title = "GT thuê"
        ws["A13"] = (
            "Hai bên cùng nhau xác nhận khối lượng từ ngày {{start_date}} "
            "đến {{end_date}} với số liệu cụ thể như sau:"
        )
        ws.merge_cells("A13:I13")
        headers = [
            "STT",
            "Ngày thuê",
            "Đến ngày",
            "Nội dung",
            "ĐVT",
            "Số lượng",
            "Thời gian thuê (ngày)",
            "Đơn giá thuê/\n1 ngày (chưa VAT)",
            "Thành tiền",
        ]
        for col, header in enumerate(headers, start=1):
            ws.cell(14, col).value = header
        for r in range(15, 215):
            ws.cell(r, 1).value = r - 14
        ws.cell(215, 5).value = "Cộng tiền thuê trước thuế"
        ws.cell(215, 9).value = "=SUM(I15:I214)"
        ws.cell(216, 5).value = "Thuế GTGT 8%"
        ws.cell(217, 5).value = "Tổng Cộng tiền thuê sau thuế"
        ws.cell(218, 1).value = "(Bằng chữ: {{total_after_tax_string}}./.)"
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_data_start_row_15_writes_below_merged_confirmation(self):
        """Uploaded template with data_start_row=15 must not write into A13:I13 merge."""
        start_date = date(2026, 4, 1)
        end_date = date(2026, 4, 30)
        self._create_transport(date(2026, 4, 10), 50)

        self.env["rental.template"].with_context(active_test=False).search([
            ("template_type", "=", "rental_invoice_xlsx"),
        ]).write({"active": False, "is_default": False})
        self.env["rental.template"].create({
            "name": "HSTT start row 15",
            "company_id": self._contract.company_id.id,
            "template_type": "rental_invoice_xlsx",
            "data_start_row": 15,
            "is_default": True,
            "file_name": "hstt_row15.xlsx",
            "file_data": base64.b64encode(self._make_hstt_template_bytes_start_row_15()),
        })

        buffer, _subtotal = self._contract._build_klct_hstt_xlsx_buffer(start_date, end_date)
        wb = load_workbook(BytesIO(buffer.getvalue()))
        hstt = wb["HSTT 04-2026"]

        # Confirmation row kept (placeholders = period dates); not overwritten by billing.
        conf = hstt.cell(13, 1).value or ""
        self.assertIn("01/04/2026", conf)
        self.assertIn("30/04/2026", conf)
        # Non-anchor merged cells stay empty (read-only MergedCell).
        from openpyxl.cell.cell import MergedCell
        self.assertTrue(isinstance(hstt.cell(13, 2), MergedCell) or hstt.cell(13, 2).value in (None, ""))

        # First billing line starts at configured row 15 (delivery 10/04).
        self.assertEqual(hstt.cell(15, 2).value.date(), date(2026, 4, 10))
        self.assertEqual(hstt.cell(15, 3).value.date(), date(2026, 4, 30))
        self.assertEqual(hstt.cell(15, 7).value, "=C15-B15")
        self.assertEqual(hstt.cell(15, 9).value, "=F15*G15*H15")

    def test_hstt_placeholders_include_contract_number_and_date(self):
        """HSTT XLSX replaces {{contract_number}} / {{contract_date}} from contract."""
        self._contract.write({
            "contract_number": "125 /HĐKT/XDMT – TLP",
            "contract_date": "2026-07-05",
        })
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Số HĐ: {{contract_number}}"
        ws["A2"] = "Ngày HĐ: {{contract_date}}"
        self._contract._rental_invoice_xlsx_apply_placeholders(
            ws, date(2026, 4, 1), date(2026, 4, 30)
        )
        self.assertEqual(ws["A1"].value, "Số HĐ: 125 /HĐKT/XDMT – TLP")
        self.assertEqual(ws["A2"].value, "Ngày HĐ: ngày 05 tháng 07 năm 2026")

    def test_export_rejects_second_posted_invoice_same_period(self):
        start_date = date(2026, 4, 1)
        end_date = date(2026, 4, 30)
        self._create_transport(date(2026, 4, 10), 10)
        self._contract.action_export_klct_hstt_excel(start_date, end_date)
        # Clear matrix so second call hits invoice-period guard (not matrix overlap wizard).
        self.env["rental.transport.matrix"].search([
            ("rental_contract_id", "=", self._contract.id),
        ]).unlink()
        with self.assertRaises(Exception) as ctx:
            self._contract.action_export_klct_hstt_excel(start_date, end_date)
        self.assertIn("đã đăng sổ", str(ctx.exception))

    def test_unlink_rental_invoice_blocked_until_cancelled(self):
        start_date = date(2026, 5, 1)
        end_date = date(2026, 5, 31)
        self._create_transport(date(2026, 5, 5), 10)
        self._contract.action_export_klct_hstt_excel(start_date, end_date)
        move = self.env["account.move"].search([
            ("rental_contract_id", "=", self._contract.id),
        ], limit=1)
        with self.assertRaises(Exception) as ctx:
            move.unlink()
        self.assertIn("Không được xóa", str(ctx.exception))

        move.button_cancel()
        move.unlink()
        self.assertFalse(move.exists())

    def test_cancel_releases_transport_fees_for_reexport(self):
        start_date = date(2026, 6, 1)
        end_date = date(2026, 6, 30)
        t1 = self._create_transport_with_fee(date(2026, 6, 2), 10, 500_000)
        self._contract.action_export_klct_hstt_excel(
            start_date, end_date, transport_fee_until_date=end_date
        )
        move = t1.fee_invoice_id
        self.assertTrue(move)
        self.assertEqual(t1.fee_billed_date, end_date)

        move.button_cancel()
        self.assertFalse(t1.fee_billed_date)
        self.assertFalse(t1.fee_invoice_id)
        self.assertEqual(move.state, "cancel")

        self.env["rental.transport.matrix"].search([
            ("rental_contract_id", "=", self._contract.id),
        ]).unlink()

        self._contract.action_export_klct_hstt_excel(
            start_date, end_date, transport_fee_until_date=end_date
        )
        self.assertEqual(t1.fee_billed_date, end_date)
        self.assertTrue(t1.fee_invoice_id)
        self.assertEqual(t1.fee_invoice_id.state, "posted")

    def test_debt_confirmation_summary_buckets(self):
        """ĐCCN collect: beginning residual + in-period totals."""
        start_date = date(2026, 4, 1)
        end_date = date(2026, 4, 30)
        self._create_transport(date(2026, 3, 10), 10)
        self._contract.action_export_klct_hstt_excel(
            date(2026, 3, 1), date(2026, 3, 31)
        )
        prior = self.env["account.move"].search([
            ("rental_contract_id", "=", self._contract.id),
            ("rental_end_date", "=", date(2026, 3, 31)),
        ], limit=1)
        self.assertEqual(prior.state, "posted")

        self._create_transport(date(2026, 4, 10), 5)
        self.env["rental.transport.matrix"].search([
            ("rental_contract_id", "=", self._contract.id),
        ]).unlink()
        self._contract.action_export_klct_hstt_excel(start_date, end_date)
        current = self.env["account.move"].search([
            ("rental_contract_id", "=", self._contract.id),
            ("rental_end_date", "=", end_date),
        ], limit=1)

        data = self._contract._collect_debt_confirmation_data(start_date, end_date)
        self.assertAlmostEqual(data["beginning_debit"], prior.amount_residual, places=2)
        self.assertAlmostEqual(data["total_amount_in_period"], current.amount_total, places=2)
        self.assertAlmostEqual(
            data["total_remain"],
            prior.amount_residual + current.amount_residual,
            places=2,
        )

        buffer, _ = self._contract._build_klct_hstt_xlsx_buffer(
            start_date, end_date, fee_invoice=current
        )
        wb = load_workbook(BytesIO(buffer.getvalue()))
        dccn = wb["ĐCCN 04-2026"]
        self.assertAlmostEqual(
            float(dccn.cell(16, 8).value or 0), prior.amount_residual, places=2
        )
        self.assertAlmostEqual(
            float(dccn.cell(17, 8).value or 0), current.amount_total, places=2
        )

    def test_md_cross_variant_return_bob_refs_klct_tong_md_and_invoice_matches(self):
        """MD multi-variant: giao chỉ 2m, trả 1m+2m → HSTT BOB F→Tổng MD; HĐ = HSTT subtotal."""
        from openpyxl.utils import get_column_letter

        md_categ = self.env["uom.category"].create({"name": "KLCT HSTT MD"})
        uom_md = self.env["uom.uom"].create({
            "name": "Mét dài KLCT-HSTT",
            "category_id": md_categ.id,
            "uom_type": "reference",
            "is_linear_meter_variant": True,
        })
        attr = self.env["product.attribute"].create({
            "name": "Length KLCT-HSTT MD",
            "create_variant": "always",
        })
        val_1 = self.env["product.attribute.value"].create({
            "name": "1m",
            "attribute_id": attr.id,
            "default_price_multiplier": 1.0,
        })
        val_2 = self.env["product.attribute.value"].create({
            "name": "2m",
            "attribute_id": attr.id,
            "default_price_multiplier": 2.0,
        })
        tmpl = self.env["product.template"].create({
            "name": "Hộp MD KLCT-HSTT",
            "type": "product",
            "list_price": 2760.0,  # /30 → 92/day at factor 1 after variant scale
            "uom_id": uom_md.id,
            "uom_po_id": uom_md.id,
            "rental_price_day": 92.0,
            "attribute_line_ids": [
                (0, 0, {"attribute_id": attr.id, "value_ids": [(6, 0, [val_1.id, val_2.id])]})
            ],
        })
        by_mult = {
            v.product_template_attribute_value_ids.price_multiplier: v
            for v in tmpl.product_variant_ids
        }
        v1, v2 = by_mult[1.0], by_mult[2.0]
        for v in (v1, v2):
            v.rental_price_day = 92.0 * (
                v.product_template_attribute_value_ids.price_multiplier or 1.0
            )

        contract = self.env["rental.contract"].create({
            "a_company_party": self._a_company.id,
            "a_party": self._a_party.id,
            "b_company_party": self.env.company.partner_id.id,
            "b_party": self._b_party.id,
            "rental_billing_mode": "day",
            "minimum_rental_months": 0,
            "minimum_penalty_current_period_only": True,
            "include_start_day_bob": True,
            "include_start_day_current": False,
        })
        self.env["rental.contract.line"].create({
            "contract_id": contract.id,
            "product_tmpl_id": tmpl.id,
            "price_unit": 2760.0,
        })

        def _tr(when, ttype, lines):
            vals = {
                "rental_contract_id": contract.id,
                "type": ttype,
                "state": "done",
                "start_rental_or_return_date": when,
                "driver_id": self._driver.id,
                "transport_truck_id": self._truck.id,
                "vehicle_start_time": when,
                "plate": "29H-91001",
                "transport_line_ids": [
                    (0, 0, {"product_id": p.id, "qty": q}) for p, q in lines
                ],
            }
            if "source_type" in self.env["rr.transport"]._fields:
                vals["source_type"] = "owned"
            return self.env["rr.transport"].create(vals)

        # Opening: 146 cây × 2m = 292 MD. Return 12×1m + 65×2m = 142 MD on 11/02.
        _tr(date(2026, 1, 10), "delivery", [(v2, 146)])
        _tr(date(2026, 2, 11), "return", [(v1, 12), (v2, 65)])
        start_date, end_date = date(2026, 2, 1), date(2026, 2, 28)
        opening_md, returned_md = 292.0, 142.0

        buffer, subtotal = contract._build_klct_hstt_xlsx_buffer(start_date, end_date)
        wb = load_workbook(BytesIO(buffer.getvalue()))
        klct = wb["KLCT 02-2026"]
        hstt = wb["HSTT 02-2026"]
        _t, _h2, header_third, start_row = find_transport_matrix_layout(klct)
        md_col = next(
            col
            for col in range(4, 20)
            if klct.cell(header_third, col).value == "Tổng MD"
        )
        letter = get_column_letter(md_col)
        self.assertEqual(klct.cell(start_row, md_col).value, int(opening_md))
        self.assertEqual(klct.cell(start_row + 1, md_col).value, -int(returned_md))

        bob_qty = hstt.cell(13, 6).value
        self.assertEqual(bob_qty, f"='KLCT 02-2026'!{letter}{start_row}")
        credit_row = None
        for row in range(13, 40):
            if "trả hàng" in str(hstt.cell(row, 4).value or ""):
                credit_row = row
                break
        self.assertIsNotNone(credit_row)
        self.assertEqual(
            hstt.cell(credit_row, 6).value,
            f"='KLCT 02-2026'!{letter}{start_row + 1}",
        )

        bob_map, map_map = rcs.calc_rental_payment_table_grouped_by_template(
            self.env, contract, start_date, end_date
        )
        hstt_subtotal = contract._payment_table_subtotal(bob_map, map_map)
        self.assertAlmostEqual(subtotal, hstt_subtotal, places=2)

        # Create invoice the same way as Xuất KLCT+HSTT (1 line = round HSTT subtotal).
        self.env["rental.transport.matrix"].search([
            ("rental_contract_id", "=", contract.id),
        ]).unlink()
        contract.action_export_klct_hstt_excel(start_date, end_date)
        move = self.env["account.move"].search([
            ("rental_contract_id", "=", contract.id),
            ("rental_end_date", "=", end_date),
        ], limit=1)
        self.assertEqual(move.state, "posted")
        inv_line = move.invoice_line_ids.filtered(lambda l: l.product_id)
        self.assertEqual(len(inv_line), 1)
        self.assertEqual(inv_line.price_unit, float(int(round(hstt_subtotal))))

        blocks = rcs.calc_rental_payment_blocks_by_template(
            self.env, contract, start_date, end_date
        )
        block = next(b for b in blocks if b["tmpl_id"] == tmpl.id)
        self.assertEqual(sum(l["qty"] for l in block["normal_lines"]), opening_md)
        self.assertEqual(block["present_total_qty"], opening_md - returned_md)
