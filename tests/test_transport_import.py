# -*- coding: utf-8 -*-
import base64
from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook

from odoo.addons.rental.helper.import_transport_matrix import (
    _company_product_groups,
    build_import_template_bytes,
    parse_transport_matrix_xlsx,
)
from odoo.addons.rental.helper.xlsx_template_utils import find_transport_matrix_layout
from odoo.exceptions import UserError
from odoo.tests.common import TransactionCase


class TestTransportImport(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        company = cls.env.company
        cls._a_company = cls.env["res.partner"].create({
            "name": "Import Test Renter",
            "is_company": True,
            "customer_type": "company",
            "is_rental_customer": True,
        })
        cls._a_party = cls.env["res.partner"].create({
            "name": "Import A rep",
            "parent_id": cls._a_company.id,
        })
        cls._b_party = cls.env["res.partner"].create({
            "name": "Import B rep",
            "parent_id": company.partner_id.id,
        })
        cls._driver = cls.env["res.partner"].create({
            "name": "Import Driver",
            "customer_type": "driver",
        })
        cls._truck = cls.env["transport.truck"].create({
            "plate": "29H-80228",
            "name": "29H-80228",
            "company_id": company.id,
        })
        cls._product = cls.env["product.product"].create({
            "name": "Khóa giáo",
            "type": "product",
        })
        cls._product.product_tmpl_id.company_id = company
        cls._contract = cls.env["rental.contract"].create({
            "a_company_party": cls._a_company.id,
            "a_party": cls._a_party.id,
            "b_company_party": company.partner_id.id,
            "b_party": cls._b_party.id,
        })
        cls.env["rental.contract.line"].create({
            "contract_id": cls._contract.id,
            "product_tmpl_id": cls._product.product_tmpl_id.id,
            "price_unit": 100.0,
        })

    def _build_workbook_bytes(self, rows, product_header="Khóa giáo"):
        wb = Workbook()
        ws = wb.active
        ws.cell(14, 4).value = "Chủng loại/ Khối lượng"
        ws.cell(15, 4).value = product_header
        ws.cell(16, 4).value = product_header
        ws.merge_cells(start_row=15, start_column=4, end_row=16, end_column=4)
        start_row = 17
        for idx, row in enumerate(rows):
            r = start_row + idx
            ws.cell(r, 1).value = idx + 1
            ws.cell(r, 2).value = row["date"]
            ws.cell(r, 3).value = row["plate"]
            ws.cell(r, 4).value = row.get("qty")
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _build_user_style_workbook_bytes(self, rows, products=None, extra_rows=None):
        """Simulate manual matrix Excel (headers row 7-8, col A=STT, B=date, C=plate)."""
        wb = Workbook()
        ws = wb.active
        ws.cell(7, 1).value = "STT"
        ws.cell(7, 2).value = "Ngày tháng"
        ws.cell(7, 3).value = "Vận chuyển xe công"
        products = products or [{"col": 4, "group": "Khóa giáo", "variant": "Khóa giáo"}]
        for prod in products:
            col = prod["col"]
            ws.cell(7, col).value = prod.get("group") or prod["variant"]
            ws.cell(8, col).value = prod.get("variant") or prod.get("group")
        if extra_rows:
            for spec in extra_rows:
                r = spec["row"]
                for col, val in spec.get("cells", {}).items():
                    ws.cell(r, col).value = val
        for idx, row in enumerate(rows):
            r = 14 + idx
            ws.cell(r, 1).value = idx + 1
            ws.cell(r, 2).value = row["date"]
            ws.cell(r, 3).value = row["plate"]
            for col, qty in (row.get("qty_by_col") or {4: row.get("qty")}).items():
                if qty is not None:
                    ws.cell(r, col).value = qty
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _build_compact_workbook_bytes(self, rows, product_header="Khóa giáo"):
        """Compact import shell: title+note, Chủng loại row 3, headers 4-5, data from 6."""
        wb = Workbook()
        ws = wb.active
        ws.cell(1, 1).value = "BẢNG XÁC NHẬN KHỐI LƯỢNG CHO THUÊ GIÀN GIÁO SẮT, THÉP HỘP VÀ PHỤ KIỆN"
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
        ws.cell(2, 1).value = (
            "Ghi chú: số lượng dương = xuất (giao), số âm = nhập (trả). Cột C = biển số xe."
        )
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
        ws.cell(3, 1).value = "STT"
        ws.cell(3, 2).value = "Ngày tháng"
        ws.cell(3, 3).value = "BKS Xe"
        ws.merge_cells(start_row=3, start_column=1, end_row=5, end_column=1)
        ws.merge_cells(start_row=3, start_column=2, end_row=5, end_column=2)
        ws.merge_cells(start_row=3, start_column=3, end_row=5, end_column=3)
        ws.cell(3, 4).value = "Chủng loại/ Khối lượng ( cây/cái/chân/cặp)"
        ws.merge_cells(start_row=3, start_column=4, end_row=3, end_column=8)
        ws.cell(4, 4).value = product_header
        ws.merge_cells(start_row=4, start_column=4, end_row=5, end_column=4)
        for idx, row in enumerate(rows):
            r = 6 + idx
            ws.cell(r, 1).value = idx + 1
            ws.cell(r, 2).value = row["date"]
            ws.cell(r, 3).value = row["plate"]
            ws.cell(r, 4).value = row.get("qty")
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_parse_delivery_row(self):
        data = self._build_workbook_bytes([{
            "date": date(2026, 3, 22),
            "plate": "29H-80228",
            "qty": 600,
        }])
        parsed = parse_transport_matrix_xlsx(data, self.env, contract=self._contract)
        self.assertEqual(len(parsed["rows"]), 1)
        row = parsed["rows"][0]
        self.assertEqual(row["transport_type"], "delivery")
        self.assertEqual(row["lines"], [(self._product.id, 600)])
        self.assertFalse(row["errors"])

    def test_parse_compact_layout_chung_loai_row_3(self):
        """Compact shell: Chủng loại row 3, product headers 4-5, data from row 6."""
        data = self._build_compact_workbook_bytes([{
            "date": date(2026, 1, 20),
            "plate": "29H-00000",
            "qty": 50,
        }])
        parsed = parse_transport_matrix_xlsx(data, self.env, contract=self._contract)
        self.assertFalse(parsed["errors"])
        self.assertFalse(parsed["product_mapping_errors"])
        self.assertEqual(len(parsed["rows"]), 1)
        row = parsed["rows"][0]
        self.assertEqual(row["row_number"], 6)
        self.assertEqual(row["transport_date"], date(2026, 1, 20))
        self.assertEqual(row["plate"], "29H-00000")
        self.assertEqual(row["transport_type"], "delivery")
        self.assertEqual(row["lines"], [(self._product.id, 50)])
        self.assertFalse(row["errors"])

    def test_build_import_template_keeps_compact_note_out_of_data(self):
        """Uploaded compact shell already has Ghi chú above table — do not write into data rows."""
        shell = self._build_compact_workbook_bytes([])
        self.env["rental.template"].create({
            "name": "Compact import shell",
            "company_id": self.env.company.id,
            "template_type": "transport_import_xlsx",
            "file_data": base64.b64encode(shell),
            "file_name": "compact_import.xlsx",
            "is_default": True,
        })
        data = build_import_template_bytes(self._contract, self.env)
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        title_row, _h2, _h3, data_start_row = find_transport_matrix_layout(ws)
        self.assertEqual(title_row, 3)
        self.assertEqual(data_start_row, 6)
        note_above = any(
            isinstance(ws.cell(r, c).value, str)
            and str(ws.cell(r, c).value).casefold().startswith("ghi chú")
            for r in range(1, data_start_row)
            for c in range(1, 5)
        )
        self.assertTrue(note_above)
        for r in range(data_start_row, data_start_row + 3):
            for c in range(1, 5):
                val = ws.cell(r, c).value
                if isinstance(val, str):
                    self.assertFalse(
                        val.casefold().startswith("ghi chú"),
                        "Ghi chú must not be written into data rows",
                    )

    def test_parse_return_row_with_negative_qty(self):
        data = self._build_workbook_bytes([{
            "date": date(2026, 3, 23),
            "plate": "29H-80228",
            "qty": -100,
        }])
        parsed = parse_transport_matrix_xlsx(data, self.env, contract=self._contract)
        row = parsed["rows"][0]
        self.assertEqual(row["transport_type"], "return")
        self.assertEqual(row["lines"], [(self._product.id, 100)])

    def test_product_mismatch_shows_detailed_error(self):
        data = self._build_workbook_bytes(
            [{
                "date": date(2026, 3, 22),
                "plate": "29H-80228",
                "qty": 10,
            }],
            product_header="SP Không Tồn Tại",
        )
        parsed = parse_transport_matrix_xlsx(data, self.env, contract=self._contract)
        self.assertTrue(parsed["product_mapping_errors"])
        detail = parsed["product_mapping_errors"][0]["detail"]
        self.assertIn("SP Không Tồn Tại", detail)
        self.assertIn("không có sản phẩm tương ứng", detail.lower())

    def test_product_in_other_company_is_mapping_error(self):
        other_company = self.env["res.company"].create({"name": "Công ty khác Import"})
        self.env["product.product"].with_company(other_company).create({
            "name": "SP Công Ty Khác",
            "type": "product",
            "company_id": other_company.id,
        })
        data = self._build_workbook_bytes(
            [{
                "date": date(2026, 3, 22),
                "plate": "29H-80228",
                "qty": 10,
            }],
            product_header="SP Công Ty Khác",
        )
        parsed = parse_transport_matrix_xlsx(data, self.env, contract=self._contract)
        self.assertTrue(parsed["product_mapping_errors"])
        detail = parsed["product_mapping_errors"][0]["detail"]
        self.assertIn("SP Công Ty Khác", detail)
        self.assertIn("không có sản phẩm tương ứng", detail.lower())
        for row in parsed["rows"]:
            self.assertFalse(row["lines"])

    def test_same_day_same_plate_multiple_trips_allowed(self):
        row = {
            "date": date(2026, 3, 22),
            "plate": "29H-80228",
            "qty": 10,
        }
        data = self._build_workbook_bytes([row, row])
        parsed = parse_transport_matrix_xlsx(data, self.env, contract=self._contract)
        self.assertEqual(len(parsed["rows"]), 2)
        self.assertFalse(parsed["errors"])
        for row in parsed["rows"]:
            self.assertFalse(row["errors"])

    def test_import_allows_existing_duplicate_on_contract(self):
        self.env["rr.transport"].create({
            "rental_contract_id": self._contract.id,
            "type": "delivery",
            "start_rental_or_return_date": date(2026, 3, 22),
            "driver_id": self._driver.id,
            "transport_truck_id": self._truck.id,
            "vehicle_start_time": date(2026, 3, 22),
            "transport_line_ids": [(0, 0, {
                "product_id": self._product.id,
                "qty": 5,
            })],
        })
        data = self._build_workbook_bytes([{
            "date": date(2026, 3, 22),
            "plate": "29H-80228",
            "qty": 20,
        }])
        wizard = self.env["rental.transport.import.wizard"].create({
            "rental_contract_id": self._contract.id,
            "default_driver_id": self._driver.id,
            "validate_picking": False,
            "import_file": base64.b64encode(data),
            "import_filename": "test.xlsx",
        })
        wizard.action_import_transports()
        transports = self._contract.rr_transport_ids.filtered(
            lambda t: t.start_rental_or_return_date == date(2026, 3, 22)
        )
        self.assertEqual(len(transports), 2)

    def test_import_creates_transport(self):
        data = self._build_workbook_bytes([{
            "date": date(2026, 4, 1),
            "plate": "30A-11111",
            "qty": 50,
        }])
        wizard = self.env["rental.transport.import.wizard"].create({
            "rental_contract_id": self._contract.id,
            "default_driver_id": self._driver.id,
            "auto_create_truck": True,
            "validate_picking": False,
            "import_file": base64.b64encode(data),
            "import_filename": "test.xlsx",
        })
        wizard.action_import_transports()
        transport = self._contract.rr_transport_ids.filtered(
            lambda t: t.start_rental_or_return_date == date(2026, 4, 1)
        )
        self.assertEqual(len(transport), 1)
        self.assertEqual(transport.type, "delivery")
        self.assertEqual(transport.transport_line_ids.qty, 50)

    def test_build_import_template_bytes(self):
        data = build_import_template_bytes(self._contract, self.env)
        self.assertTrue(data)
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        _title, header_second_row, _header_third, _data_start = find_transport_matrix_layout(ws)
        header_values = [
            ws.cell(header_second_row, col).value
            for col in range(4, (ws.max_column or 4) + 1)
            if ws.cell(header_second_row, col).value
        ]
        self.assertTrue(
            any(self._product.display_name in str(v) or "Khóa giáo" in str(v) for v in header_values),
            "Import template should include company product columns",
        )
        groups = _company_product_groups(self._contract, self.env)
        self.assertTrue(any(g["tmpl_id"] == self._product.product_tmpl_id.id for g in groups))

    def test_import_template_columns_follow_product_order(self):
        first = self.env["product.template"].create({
            "name": "ZZ Order First Template",
            "type": "product",
            "company_id": self.env.company.id,
            "order": 1,
        })
        second = self.env["product.template"].create({
            "name": "AA Order Second Template",
            "type": "product",
            "company_id": self.env.company.id,
            "order": 50,
        })
        # Not on contract quotation lines — still must appear, ordered by `order`.
        groups = _company_product_groups(self._contract, self.env)
        tmpl_ids = [g["tmpl_id"] for g in groups]
        self.assertIn(first.id, tmpl_ids)
        self.assertIn(second.id, tmpl_ids)
        self.assertLess(tmpl_ids.index(first.id), tmpl_ids.index(second.id))

        data = build_import_template_bytes(self._contract, self.env)
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        _title, header_second_row, _header_third, _data_start = find_transport_matrix_layout(ws)
        headers = [
            str(ws.cell(header_second_row, col).value or "")
            for col in range(4, (ws.max_column or 4) + 1)
        ]
        first_name = first.product_variant_ids[:1].display_name
        second_name = second.product_variant_ids[:1].display_name
        first_cols = [i for i, h in enumerate(headers) if first_name in h or first.name in h]
        second_cols = [i for i, h in enumerate(headers) if second_name in h or second.name in h]
        self.assertTrue(first_cols)
        self.assertTrue(second_cols)
        self.assertLess(first_cols[0], second_cols[0])

    def test_import_template_variants_sorted_by_lst_price(self):
        md_categ = self.env["uom.category"].create({"name": "MD import template test"})
        uom_md = self.env["uom.uom"].create({
            "name": "Mét dài (import template)",
            "category_id": md_categ.id,
            "uom_type": "reference",
            "is_linear_meter_variant": True,
        })
        attr = self.env["product.attribute"].create({
            "name": "Length Import Order",
            "create_variant": "always",
        })
        val_expensive = self.env["product.attribute.value"].create({
            "name": "3m",
            "attribute_id": attr.id,
            "default_price_multiplier": 3.0,
        })
        val_cheap = self.env["product.attribute.value"].create({
            "name": "1m",
            "attribute_id": attr.id,
            "default_price_multiplier": 1.0,
        })
        tmpl = self.env["product.template"].create({
            "name": "Hộp Order By Price",
            "type": "product",
            "company_id": self.env.company.id,
            "list_price": 1000.0,
            "order": 5,
            "uom_id": uom_md.id,
            "uom_po_id": uom_md.id,
            "attribute_line_ids": [(0, 0, {
                "attribute_id": attr.id,
                "value_ids": [(6, 0, [val_expensive.id, val_cheap.id])],
            })],
        })
        groups = _company_product_groups(self._contract, self.env)
        group = next(g for g in groups if g["tmpl_id"] == tmpl.id)
        self.assertTrue(group["needs_md"])
        prod_order = group["prod_order"]
        prices = [group["p_tmpl"]["products"][pid]["lst_price"] for pid in prod_order]
        self.assertEqual(prices, sorted(prices))
        labels = [
            group["p_tmpl"]["products"][pid]["variant_name"]
            for pid in prod_order
        ]
        self.assertEqual(labels, ["1m", "3m"])

        data = build_import_template_bytes(self._contract, self.env)
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        _title, header_second_row, header_third_row, _data_start = find_transport_matrix_layout(ws)
        group_col = None
        for col in range(4, (ws.max_column or 4) + 1):
            if ws.cell(header_second_row, col).value == tmpl.display_name:
                group_col = col
                break
        self.assertIsNotNone(group_col)
        self.assertEqual(ws.cell(header_third_row, group_col).value, "1m")
        self.assertEqual(ws.cell(header_third_row, group_col + 1).value, "3m")
        self.assertEqual(ws.cell(header_third_row, group_col + 2).value, "Tổng MD")

    def test_import_template_styles_extend_beyond_template_width(self):
        """Product columns past the static template width keep header/data borders."""
        for idx in range(20):
            self.env["product.template"].create({
                "name": f"Style Col Product {idx:02d}",
                "type": "product",
                "company_id": self.env.company.id,
                "order": 100 + idx,
            })
        data = build_import_template_bytes(self._contract, self.env)
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        title_row, header_second_row, header_third_row, data_start_row = find_transport_matrix_layout(ws)
        ref_col = 4
        far_col = 20  # column T — beyond typical pre-styled template width (~O)
        ref_cell = ws.cell(header_third_row, ref_col)
        far_cell = ws.cell(header_third_row, far_col)
        self.assertTrue(
            (ref_cell.border.left and ref_cell.border.left.style)
            or (ref_cell.fill and ref_cell.fill.fill_type),
            "Reference product column should have style in the static template",
        )
        self.assertEqual(
            far_cell.border.left.style if far_cell.border.left else None,
            ref_cell.border.left.style if ref_cell.border.left else None,
        )
        self.assertEqual(
            far_cell.border.right.style if far_cell.border.right else None,
            ref_cell.border.right.style if ref_cell.border.right else None,
        )
        self.assertEqual(far_cell.fill.fill_type, ref_cell.fill.fill_type)
        self.assertEqual(
            getattr(far_cell.fill.fgColor, "rgb", None),
            getattr(ref_cell.fill.fgColor, "rgb", None),
        )
        data_ref = ws.cell(data_start_row, ref_col)
        data_far = ws.cell(data_start_row, far_col)
        self.assertEqual(
            data_far.border.left.style if data_far.border.left else None,
            data_ref.border.left.style if data_ref.border.left else None,
        )
        # Title merge should span through the far product column.
        title_merged = False
        for rng in ws.merged_cells.ranges:
            if (
                rng.min_row == title_row == rng.max_row
                and rng.min_col <= ref_col
                and rng.max_col >= far_col
            ):
                title_merged = True
                break
        self.assertTrue(title_merged)

    def test_import_validate_picking_allows_negative_stock(self):
        if not self.env["stock.warehouse"].search([("company_id", "=", self.env.company.id)], limit=1):
            self.env["stock.warehouse"].create({
                "name": "Import WH",
                "code": "IMP",
                "company_id": self.env.company.id,
            })
        data = self._build_workbook_bytes([{
            "date": date(2026, 5, 10),
            "plate": "99X-99999",
            "qty": 25,
        }])
        wizard = self.env["rental.transport.import.wizard"].create({
            "rental_contract_id": self._contract.id,
            "default_driver_id": self._driver.id,
            "auto_create_truck": True,
            "validate_picking": True,
            "import_file": base64.b64encode(data),
            "import_filename": "test.xlsx",
        })
        wizard.action_import_transports()
        transport = self._contract.rr_transport_ids.filtered(
            lambda t: t.start_rental_or_return_date == date(2026, 5, 10)
        )
        self.assertEqual(len(transport), 1)
        self.assertEqual(transport.state, "done")
        picking = transport.picking_ids.filtered(lambda p: p.state == "done" and not p.return_id)
        self.assertTrue(picking)

    def test_user_style_layout_rows_7_8(self):
        data = self._build_user_style_workbook_bytes([{
            "date": date(2026, 3, 16),
            "plate": "30B-02898",
            "qty": 120,
        }])
        parsed = parse_transport_matrix_xlsx(data, self.env, contract=self._contract)
        self.assertEqual(len(parsed["rows"]), 1)
        row = parsed["rows"][0]
        self.assertEqual(row["transport_type"], "delivery")
        self.assertEqual(row["lines"], [(self._product.id, 120)])

    def test_skip_opening_balance_section(self):
        data = self._build_user_style_workbook_bytes(
            [{
                "date": date(2026, 3, 16),
                "plate": "30B-02898",
                "qty": 50,
            }],
            extra_rows=[
                {"row": 9, "cells": {2: "DƯ ĐẦU KỲ - PHỐ CÀ LANMAK C SANG"}},
                {"row": 10, "cells": {4: 999}},
                {"row": 11, "cells": {4: 888}},
                {"row": 13, "cells": {2: "CỘNG CHUYỂN DƯ ĐẦU T3", 4: 1887}},
            ],
        )
        parsed = parse_transport_matrix_xlsx(data, self.env, contract=self._contract)
        self.assertEqual(len(parsed["rows"]), 1)
        self.assertEqual(parsed["rows"][0]["lines"], [(self._product.id, 50)])

    def test_skip_monthly_total_row(self):
        data = self._build_user_style_workbook_bytes(
            [
                {"date": date(2026, 3, 16), "plate": "30B-02898", "qty": 10},
                {"date": "Cộng tháng 03/2026", "plate": "", "qty": 999},
                {"date": date(2026, 4, 1), "plate": "29H-80228", "qty": 20},
            ],
        )
        parsed = parse_transport_matrix_xlsx(data, self.env, contract=self._contract)
        self.assertEqual(len(parsed["rows"]), 2)
        self.assertEqual(parsed["rows"][0]["lines"], [(self._product.id, 10)])
        self.assertEqual(parsed["rows"][1]["lines"], [(self._product.id, 20)])

    def test_product_name_comma_decimal_matches(self):
        product = self.env["product.product"].create({
            "name": "Giáo nêm 2.5m",
            "type": "product",
        })
        self.env["rental.contract.line"].create({
            "contract_id": self._contract.id,
            "product_tmpl_id": product.product_tmpl_id.id,
            "price_unit": 10.0,
        })
        data = self._build_user_style_workbook_bytes(
            [{"date": date(2026, 3, 18), "plate": "29H-80228", "qty": 7}],
            products=[{"col": 4, "group": "Giáo nêm 2,5m", "variant": "Giáo nêm 2,5m"}],
        )
        parsed = parse_transport_matrix_xlsx(data, self.env, contract=self._contract)
        self.assertEqual(len(parsed["rows"]), 1)
        self.assertEqual(parsed["rows"][0]["lines"], [(product.id, 7)])

    def test_variant_size_columns_map_to_distinct_variants(self):
        attr = self.env["product.attribute"].create({
            "name": "Mét Import Test",
            "create_variant": "always",
        })
        values = {}
        for label in ("0,9m", "2m", "6m"):
            values[label] = self.env["product.attribute.value"].create({
                "name": label,
                "attribute_id": attr.id,
            })
        tmpl = self.env["product.template"].create({
            "name": "Hộp Test 5*5",
            "type": "product",
            "attribute_line_ids": [(0, 0, {
                "attribute_id": attr.id,
                "value_ids": [(6, 0, [v.id for v in values.values()])],
            })],
        })
        self.env["rental.contract.line"].create({
            "contract_id": self._contract.id,
            "product_tmpl_id": tmpl.id,
            "price_unit": 10.0,
        })
        data = self._build_user_style_workbook_bytes(
            [{
                "date": date(2026, 7, 1),
                "plate": "29H-80228",
                "qty_by_col": {4: 200, 5: 150, 6: 100},
            }],
            products=[
                {"col": 4, "group": "Hộp Test 5*5", "variant": "0.9"},
                {"col": 5, "group": "Hộp Test 5*5", "variant": "2.0"},
                {"col": 6, "group": "Hộp Test 5*5", "variant": "6.0"},
            ],
        )
        parsed = parse_transport_matrix_xlsx(data, self.env, contract=self._contract)
        self.assertFalse(parsed["product_mapping_errors"])
        self.assertEqual(len(parsed["rows"]), 1)
        lines = dict(parsed["rows"][0]["lines"])

        def variant_id(label):
            variant = tmpl.product_variant_ids.filtered(
                lambda p: label in p.product_template_variant_value_ids.mapped("name")
            )
            return variant.id

        self.assertEqual(len(lines), 3)
        self.assertEqual(lines[variant_id("0,9m")], 200)
        self.assertEqual(lines[variant_id("2m")], 150)
        self.assertEqual(lines[variant_id("6m")], 100)

    def test_import_blocks_on_product_mapping_error(self):
        data = self._build_workbook_bytes(
            [{
                "date": date(2026, 6, 5),
                "plate": "29H-80228",
                "qty": 10,
            }],
            product_header="SP Không Tồn Tại",
        )
        wizard = self.env["rental.transport.import.wizard"].create({
            "rental_contract_id": self._contract.id,
            "default_driver_id": self._driver.id,
            "validate_picking": False,
            "import_file": base64.b64encode(data),
            "import_filename": "test.xlsx",
        })
        before = len(self._contract.rr_transport_ids)
        with self.assertRaises(UserError) as ctx:
            wizard.action_import_transports()
        self.assertIn("SP Không Tồn Tại", str(ctx.exception))
        self.assertEqual(len(self._contract.rr_transport_ids), before)

    def test_import_skips_error_rows(self):
        row_ok = {"date": date(2026, 6, 1), "plate": "29H-80228", "qty": 10}
        row_bad = {"date": date(2026, 6, 2), "plate": "", "qty": 5}
        data = self._build_workbook_bytes([row_ok, row_bad])
        wizard = self.env["rental.transport.import.wizard"].create({
            "rental_contract_id": self._contract.id,
            "default_driver_id": self._driver.id,
            "validate_picking": False,
            "import_file": base64.b64encode(data),
            "import_filename": "test.xlsx",
        })
        before = len(self._contract.rr_transport_ids)
        wizard.action_import_transports()
        self.assertEqual(len(self._contract.rr_transport_ids), before + 1)
