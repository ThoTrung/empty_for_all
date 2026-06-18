from pathlib import Path
from io import BytesIO
from openpyxl import load_workbook
from odoo.addons.rental.helper.import_transport_matrix import (
    _find_import_worksheet,
    _detect_import_layout,
    parse_transport_matrix_xlsx,
)

path = Path("/mnt/extra-addons/odoo_modules/rental/tests/fixtures/import delivery.xlsx")
wb = load_workbook(BytesIO(path.read_bytes()), data_only=True)
print("SHEETS", wb.sheetnames)
ws = _find_import_worksheet(wb)
print("PICKED", repr(ws.title))
print("LAYOUT", _detect_import_layout(ws))

contract = env["rental.contract"].search([("code", "=", "RC00052")], limit=1)
parsed = parse_transport_matrix_xlsx(path.read_bytes(), env, contract=contract)
ok = [r for r in parsed["rows"] if not r.get("errors")]
err = [r for r in parsed["rows"] if r.get("errors")]
print("TOTAL", len(parsed["rows"]), "OK", len(ok), "ERR", len(err))
print("MAPPING", len(parsed.get("product_mapping_errors") or []))
for row in err:
    print("ERR", row["row_number"], row.get("transport_date"), row.get("plate"))
    for e in (row.get("errors") or [])[:2]:
        print(" ", str(e).split("\n")[0][:120])
