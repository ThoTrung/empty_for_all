# -*- coding: utf-8 -*-
"""Parse transport-matrix Excel (inverse of export) into rr.transport payloads."""
import re
from collections import OrderedDict
from datetime import date, datetime
from difflib import SequenceMatcher
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from odoo import _, fields

from .xlsx_template_utils import (
    find_transport_matrix_layout,
    replace_placeholders_in_sheet,
    transport_matrix_date_replacements,
)

_SKIP_ROW_MARKERS = (
    "tồn đầu kỳ",
    "dư đầu kỳ",
    "dư đầu",
    "cộng chuyển dư đầu",
    "cộng chuyển",
    "tổng kl cuối",
    "tổng kl",
    "cộng tháng",
)
_OPENING_SECTION_START = ("dư đầu kỳ",)
_OPENING_SECTION_END = ("cộng chuyển dư đầu",)
_MD_HEADER = "tổng md"
_DEFAULT_DATE_COL = 2
_DEFAULT_PLATE_COL = 3
_DEFAULT_PRODUCT_START_COL = 4


def _normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _normalize_key(value):
    return _normalize_text(value).casefold()


def _normalize_product_key(value):
    """Normalize product labels for matching (comma decimal, spacing)."""
    text = _normalize_key(value)
    text = text.replace(",", ".")
    text = re.sub(r"\s+", " ", text)
    return text


def _normalize_variant_token(value):
    text = _normalize_text(value)
    if not text:
        return ""
    text = text.casefold()
    text = re.sub(r"\s*m\b", "", text)
    text = text.replace(",", ".").strip()
    # Canonicalize pure numbers so Excel «2.0»/«1.20»/«6.0» match Odoo «2m»/«1,2m»/«6m».
    if re.fullmatch(r"\d+(?:\.\d+)?", text):
        num = float(text)
        if num == int(num):
            return str(int(num))
        return ("%f" % num).rstrip("0").rstrip(".")
    return text


def _parse_excel_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _normalize_text(value)
    if not text:
        return None
    for fmt in ("%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:
        return fields.Date.from_string(text)
    except (ValueError, TypeError):
        return None


def _parse_qty(value):
    if value is None or value == "":
        return None
    try:
        qty = float(value)
    except (TypeError, ValueError):
        return None
    if abs(qty) < 1e-12:
        return None
    return int(round(qty))


def _cell_value(ws, row, col):
    return ws.cell(row, col).value


def _merged_top_left_value(ws, row, col):
    """Return the value stored on the merged range top-left for (row, col)."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col).value
    return _cell_value(ws, row, col)


def _walk_left_value(ws, row, col, min_col=1):
    """For horizontally merged headers, walk left until a value is found."""
    for c in range(col, min_col - 1, -1):
        val = _merged_top_left_value(ws, row, c)
        if val is not None and str(val).strip():
            return val
    return None


def _is_skip_row_label(value):
    key = _normalize_key(value)
    return any(marker in key for marker in _SKIP_ROW_MARKERS)


def _is_variant_dimension_token(value):
    """True for size tokens like 3.0, 1.5 — not full product names."""
    token = _normalize_variant_token(value)
    if not token:
        return False
    return bool(re.match(r"^[\d.]+$", token))


def _row_label_text(ws, row_idx, date_col, plate_col):
    parts = []
    for col in (1, date_col, plate_col):
        val = _normalize_text(_merged_top_left_value(ws, row_idx, col))
        if val:
            parts.append(val)
    return " ".join(parts).casefold()


def _should_skip_row(ws, row_idx, date_col, plate_col):
    for col in (1, date_col, plate_col, plate_col + 1):
        if _is_skip_row_label(_merged_top_left_value(ws, row_idx, col)):
            return True
    return False


def _sheet_has_import_layout(ws):
    for row_idx in range(1, 41):
        for col in range(1, 8):
            val = _merged_top_left_value(ws, row_idx, col)
            if isinstance(val, str):
                key = _normalize_key(val)
                if "chủng loại" in key or "ngày tháng" in key:
                    return True
    return False


def _find_import_worksheet(wb):
    """Pick worksheet with transport matrix headers (skip empty active sheet)."""
    for name in wb.sheetnames:
        ws = wb[name]
        for row_idx in range(1, 41):
            for col in range(1, 8):
                val = _merged_top_left_value(ws, row_idx, col)
                if isinstance(val, str) and "chủng loại" in val:
                    return ws
    for name in wb.sheetnames:
        ws = wb[name]
        if _sheet_has_import_layout(ws):
            return ws
    return wb.active


def _row_has_variant_header_row(ws, row_idx, product_start_col):
    max_col = min(ws.max_column or product_start_col, product_start_col + 80)
    for col in range(product_start_col, max_col + 1):
        text = _normalize_text(_cell_value(ws, row_idx, col))
        if not text or _normalize_key(text) == _MD_HEADER:
            continue
        if _is_variant_dimension_token(text):
            return True
        group = _normalize_text(_walk_left_value(ws, row_idx - 1, col, min_col=product_start_col))
        if group and _normalize_product_key(text) != _normalize_product_key(group):
            return True
    return False


def _detect_import_layout(ws):
    """
    Detect column positions and header rows for matrix / manual Excel files.

    Supports:
    - Standard template (row «Chủng loại», headers at +1/+2, data at +3)
    - Manual sheets (row «Ngày tháng» at col B/C with product headers same row or below)
    """
    for row in ws.iter_rows(min_row=1, max_row=40):
        for cell in row:
            val = cell.value
            if isinstance(val, str) and "Chủng loại" in val:
                title_row = cell.row
                return {
                    "date_col": _DEFAULT_DATE_COL,
                    "plate_col": _DEFAULT_PLATE_COL,
                    "product_start_col": _DEFAULT_PRODUCT_START_COL,
                    "header_second_row": title_row + 1,
                    "header_third_row": title_row + 2,
                    "data_start_row": title_row + 3,
                }

    for row_idx in range(1, 41):
        for col in range(1, 6):
            val = _cell_value(ws, row_idx, col)
            if not isinstance(val, str):
                continue
            if "ngày tháng" not in _normalize_key(val):
                continue
            date_col = col
            plate_col = col + 1
            product_start_col = col + 2
            group_on_same_row = bool(
                _normalize_text(
                    _walk_left_value(ws, row_idx, product_start_col, min_col=product_start_col)
                )
            )
            if group_on_same_row:
                variant_row = row_idx + 1
                if _row_has_variant_header_row(ws, variant_row, product_start_col):
                    header_second_row = row_idx
                    header_third_row = variant_row
                else:
                    header_second_row = header_third_row = row_idx
                return {
                    "date_col": date_col,
                    "plate_col": plate_col,
                    "product_start_col": product_start_col,
                    "header_second_row": header_second_row,
                    "header_third_row": header_third_row,
                    "data_start_row": header_third_row + 1,
                }
            return {
                "date_col": date_col,
                "plate_col": plate_col,
                "product_start_col": product_start_col,
                "header_second_row": row_idx + 1,
                "header_third_row": row_idx + 2,
                "data_start_row": row_idx + 3,
            }

    _, header_second_row, header_third_row, data_start_row = find_transport_matrix_layout(ws)
    return {
        "date_col": _DEFAULT_DATE_COL,
        "plate_col": _DEFAULT_PLATE_COL,
        "product_start_col": _DEFAULT_PRODUCT_START_COL,
        "header_second_row": header_second_row,
        "header_third_row": header_third_row,
        "data_start_row": data_start_row,
    }


def _resolve_plate(ws, row_idx, plate_col, max_row, window=10):
    """Use cell plate or nearest non-empty plate in following rows (same shipment block)."""
    plate = _normalize_text(_merged_top_left_value(ws, row_idx, plate_col))
    if plate:
        return plate, False
    for r in range(row_idx + 1, min(max_row, row_idx + window) + 1):
        below = _normalize_text(_merged_top_left_value(ws, r, plate_col))
        if below:
            return below, True
    return "", False


def _normalize_md_header():
    return _MD_HEADER


def _excel_label(group_name, variant_label=None):
    group = _normalize_text(group_name)
    variant = _normalize_text(variant_label)
    if variant and _normalize_key(variant) != _MD_HEADER:
        return f"{group} - {variant}"
    return group


def _excel_name_aliases(group_name, variant_label=None):
    """Alternate labels used on manual Excel sheets vs Odoo product names."""
    group = _normalize_text(group_name)
    variant = _normalize_text(variant_label)
    labels = []
    seen = set()

    def add(text):
        text = _normalize_text(text)
        if not text:
            return
        key = _normalize_product_key(text)
        if key in seen:
            return
        seen.add(key)
        labels.append(text)

    def synonymize(text):
        result = _normalize_key(text)
        rules = (
            (r"giáo\s*ht", "giáo hoàn thiện"),
            (r"giằng\s*ht", "giằng hoàn thiện"),
            (r"bát\s*kích", "kích đầu"),
            (r"chân\s*kích", "kích chân"),
            (r"ống\s*nối\s*nêm", "ống nối"),
            (r"u\s*chống\s*truyền", "u chống chuyền"),
            (r"tuýp\s*mạ\s*kẽm\s*d48", "tuýp d48"),
        )
        for pattern, repl in rules:
            result = re.sub(pattern, repl, result)
        result = re.sub(r"1[,.]7\b", "1.7m", result)
        return result

    add(_excel_label(group, variant))
    if variant:
        add(f"{group} ({variant})")
    add(group)

    for label in list(labels):
        add(synonymize(label))
        if variant:
            add(f"{synonymize(group)} - {variant}")

    if _normalize_product_key(group) == _normalize_product_key("Mâm"):
        add("Mâm giáo")

    return labels


def _compact_product_key(value):
    return _normalize_product_key(value).replace(" ", "").replace("*", "")


def _loose_product_match(needle, haystack):
    """Match Excel shorthand to Odoo names, e.g. kích đầu L500 vs Kích đầu D38* L500."""
    if not needle or not haystack:
        return False
    if needle in haystack or haystack in needle:
        return True
    n_len = re.search(r"l(\d+)", needle)
    h_len = re.search(r"l(\d+)", haystack)
    if not n_len or not h_len or n_len.group(1) != h_len.group(1):
        return False
    for token in ("đầu", "chân"):
        if token in needle and token in haystack:
            return "kích" in needle and "kích" in haystack
    return False


class TransportMatrixProductResolver:
    """Resolve Excel column labels to product.product records by name."""

    def __init__(self, env, contract=None):
        self.env = env
        self.contract = contract
        self._by_display_name = {}
        self._by_template_name = {}
        self._all_template_names = []
        self._errors = {}

    def _contract_company_id(self):
        if self.contract and self.contract.company_id:
            return self.contract.company_id.id
        return None

    def _product_models(self):
        Product = self.env["product.product"]
        Template = self.env["product.template"]
        if self.contract and self.contract.company_id:
            company = self.contract.company_id
            Product = Product.with_company(company)
            Template = Template.with_company(company)
        return Product, Template

    def _product_domain(self):
        domain = [("active", "in", [True, False])]
        cid = self._contract_company_id()
        if cid:
            domain += ["|", ("company_id", "=", False), ("company_id", "=", cid)]
        return domain

    def _product_allowed_for_contract(self, product):
        cid = self._contract_company_id()
        if not cid or not product:
            return True
        pc = product.company_id
        return not pc or pc.id == cid

    def _index_products(self):
        if self._by_display_name:
            return
        Product, Template = self._product_models()
        domain = self._product_domain()
        for product in Product.search(domain):
            labels = [product.display_name]
            # product.name is the bare template name, shared by every variant.
            # Only index it when the template has a single variant, otherwise the
            # bare group name would collapse all columns onto one arbitrary variant.
            if len(product.product_tmpl_id.product_variant_ids) <= 1:
                labels.append(product.name)
            for label in labels:
                self._by_display_name.setdefault(_normalize_product_key(label), product)
        for tmpl in Template.search(domain):
            key = _normalize_product_key(tmpl.name)
            self._by_template_name.setdefault(key, tmpl)
            self._all_template_names.append(tmpl.name)

    def _similar_template_names(self, group_name, limit=5):
        self._index_products()
        needle = _normalize_product_key(group_name)
        if not needle:
            return []
        scored = []
        for name in self._all_template_names:
            key = _normalize_product_key(name)
            ratio = SequenceMatcher(None, needle, key).ratio()
            if needle in key or key in needle:
                scored.append((0.0, -ratio, name))
            elif ratio >= 0.55:
                scored.append((1.0, -ratio, name))
            elif needle[:4] and needle[:4] in key:
                scored.append((2.0, -ratio, name))
        scored.sort()
        return [name for _a, _b, name in scored[:limit]]

    def _match_contract_product(self, group_name, variant_label=None, threshold=0.72):
        if not self.contract:
            return None
        needle = _normalize_product_key(_excel_label(group_name, variant_label))
        if not needle:
            return None
        best_product = None
        best_ratio = 0.0
        for label in _excel_name_aliases(group_name, variant_label):
            needle = _normalize_product_key(label)
            for line in self.contract.rental_contract_line_ids:
                tmpl = line.product_tmpl_id
                if not tmpl:
                    continue
                for prod in tmpl.product_variant_ids:
                    for name in (prod.display_name, tmpl.name):
                        ratio = SequenceMatcher(
                            None, needle, _normalize_product_key(name)
                        ).ratio()
                        if ratio > best_ratio:
                            best_ratio = ratio
                            best_product = prod
        if best_product and best_ratio >= threshold:
            return best_product
        return None

    def _contract_product_names(self):
        if not self.contract:
            return []
        names = []
        for line in self.contract.rental_contract_line_ids:
            tmpl = line.product_tmpl_id
            if not tmpl:
                continue
            if len(tmpl.product_variant_ids) <= 1:
                names.append(tmpl.display_name)
            else:
                names.extend(tmpl.product_variant_ids.mapped("display_name"))
        return names

    def _format_resolution_error(self, group_name, variant_label=None, col_letter=None, header_rows=None):
        group = _normalize_text(group_name)
        variant = _normalize_text(variant_label)
        excel_label = _excel_label(group, variant)
        header_hint = ""
        if header_rows:
            header_hint = _("dòng %(r1)s–%(r2)s") % {
                "r1": header_rows[0],
                "r2": header_rows[1],
            }
        parts = [
            _("Tên trên Excel: «%(label)s»") % {"label": excel_label},
        ]
        if col_letter:
            parts.append(
                _("Cột Excel: %(col)s (%(header)s)") % {
                    "col": col_letter,
                    "header": header_hint or _("header sản phẩm"),
                }
            )

        tmpl = self._by_template_name.get(_normalize_product_key(group))
        if tmpl:
            variant_names = tmpl.product_variant_ids.mapped("display_name")
            if variant:
                parts.append(
                    _("Đã tìm thấy mẫu SP «%(tmpl)s» nhưng không khớp biến thể «%(variant)s».") % {
                        "tmpl": tmpl.display_name,
                        "variant": variant,
                    }
                )
            else:
                parts.append(
                    _("Đã tìm thấy mẫu SP «%(tmpl)s» nhưng không xác định được biến thể.") % {
                        "tmpl": tmpl.display_name,
                    }
                )
            if variant_names:
                parts.append(
                    _("Tên biến thể đúng trong Odoo: %(names)s") % {
                        "names": "; ".join(variant_names),
                    }
                )
            parts.append(
                _("Hãy sửa tên ở header Excel (%(header)s) cho khớp một tên trong Odoo ở trên.") % {
                    "header": header_hint or _("dòng header sản phẩm"),
                }
            )
        else:
            parts.append(
                _("Không có sản phẩm tương ứng «%(label)s» trong hệ thống.") % {
                    "label": excel_label,
                }
            )
        return "\n".join(parts)

    def _finalize_product(self, product, cache_key, group, variant, col_letter, header_rows):
        if product and self._product_allowed_for_contract(product):
            return product, None
        msg = self._format_resolution_error(
            group, variant, col_letter=col_letter, header_rows=header_rows,
        )
        self._errors[cache_key] = msg
        return None, msg

    def _find_template(self, group, variant=None):
        tmpl = self._by_template_name.get(_normalize_product_key(group))
        if tmpl:
            return tmpl
        for label in _excel_name_aliases(group, variant):
            tmpl = self._by_template_name.get(_normalize_product_key(label.split(" - ")[0]))
            if tmpl:
                return tmpl
        return None

    def _resolve_in_template(self, group, variant):
        """Resolve to a specific variant of the matching template.

        When a size/variant column is given, match it against the template's
        variant values (e.g. Excel «2.0» → variant «2m»). Never collapse onto an
        arbitrary variant when the size cannot be matched on a multi-variant
        template.
        """
        tmpl = self._find_template(group, variant)
        if not tmpl:
            return None
        variants = tmpl.product_variant_ids
        if not variant:
            return variants[0] if len(variants) == 1 else None

        token = _normalize_variant_token(variant)
        for prod in variants:
            ptav_tokens = [
                _normalize_variant_token(name)
                for name in prod.product_template_variant_value_ids.mapped("name")
            ]
            if token and token in ptav_tokens:
                return prod
        target = _normalize_product_key(f"{group} - {variant}")
        for prod in variants:
            if _normalize_product_key(prod.display_name) == target:
                return prod
        if len(variants) == 1:
            return variants[0]
        return None

    def resolve(self, group_name, variant_label=None, col_index=None, header_rows=None):
        self._index_products()
        group = _normalize_text(group_name)
        variant = _normalize_text(variant_label)
        cache_key = (group, variant, col_index)
        if cache_key in self._errors:
            return None, self._errors[cache_key]

        col_letter = get_column_letter(col_index) if col_index else None

        # When a size/variant is given, resolve within the matching template's
        # variants first so the bare group name can never collapse every column
        # onto a single arbitrary variant.
        if variant:
            prod = self._resolve_in_template(group, variant)
            if prod is not None:
                return self._finalize_product(
                    prod, cache_key, group, variant, col_letter, header_rows,
                )

        for label in _excel_name_aliases(group, variant):
            product = self._by_display_name.get(_normalize_product_key(label))
            if product:
                return self._finalize_product(
                    product, cache_key, group, variant, col_letter, header_rows,
                )

        needles = [_compact_product_key(label) for label in _excel_name_aliases(group, variant)]
        for needle in needles:
            if len(needle) < 8:
                continue
            for key, product in self._by_display_name.items():
                compact = _compact_product_key(key)
                if _loose_product_match(needle, compact):
                    return self._finalize_product(
                        product, cache_key, group, variant, col_letter, header_rows,
                    )

        if not variant:
            prod = self._resolve_in_template(group, variant)
            if prod is not None:
                return self._finalize_product(
                    prod, cache_key, group, variant, col_letter, header_rows,
                )

        contract_product = self._match_contract_product(group, variant)
        if contract_product:
            return self._finalize_product(
                contract_product, cache_key, group, variant, col_letter, header_rows,
            )

        msg = self._format_resolution_error(
            group, variant, col_letter=col_letter, header_rows=header_rows,
        )
        self._errors[cache_key] = msg
        return None, msg


def _build_column_map(ws, header_second_row, header_third_row, product_start_col):
    """Map 1-based column index -> (group_name, variant_label, column_label)."""
    max_col = ws.max_column or product_start_col
    columns = {}
    for col in range(product_start_col, max_col + 1):
        group_name = _walk_left_value(ws, header_second_row, col, min_col=product_start_col)
        variant_label = _cell_value(ws, header_third_row, col)
        group_text = _normalize_text(group_name)
        variant_text = _normalize_text(variant_label)

        if not group_text and not variant_text:
            continue
        if _normalize_key(variant_text) == _MD_HEADER or _normalize_key(group_text) == _MD_HEADER:
            continue

        if not group_text and variant_text:
            if _is_variant_dimension_token(variant_text):
                continue
            group_text = variant_text
            variant_text = ""
        elif variant_text and _normalize_key(variant_text) == _normalize_key(group_text):
            variant_text = ""
        elif variant_text and not _is_variant_dimension_token(variant_text) and group_text:
            if _normalize_product_key(variant_text) == _normalize_product_key(group_text):
                variant_text = ""

        label = _excel_label(group_text, variant_text)
        columns[col] = {
            "group_name": group_text,
            "variant_label": variant_text,
            "label": label,
            "col_letter": get_column_letter(col),
        }
    return columns


def _contract_product_groups(contract, env):
    """Build product groups from contract quotation lines (template + variants)."""
    from odoo.addons.rental.models import rental_transport_matrix as rtm

    product_tmpls = OrderedDict()
    for line in contract.rental_contract_line_ids:
        tmpl = line.product_tmpl_id
        if not tmpl:
            continue
        entry = product_tmpls.get(tmpl.id)
        if not entry:
            entry = {
                "product_tmpl_id": tmpl.id,
                "product_tmpl_name": tmpl.display_name,
                "products": OrderedDict(),
            }
            product_tmpls[tmpl.id] = entry
        for product in tmpl.product_variant_ids:
            entry["products"][product.id] = {
                "id": product.id,
                "name": product.display_name,
                "variant_name": product.product_template_variant_value_ids.mapped("name")[:1] or "",
                "price_multiplier": rtm._variant_price_multiplier(product),
            }

    groups = []
    for tmpl_id, p_tmpl in product_tmpls.items():
        p_tmpl["products"] = rtm._sort_products_odict(OrderedDict(p_tmpl["products"]))
        prod_order = list(p_tmpl["products"].keys())
        prods = p_tmpl["products"]
        needs_md = rtm._group_needs_md_column(env, prods, prod_order)
        groups.append({
            "tmpl_id": tmpl_id,
            "p_tmpl": p_tmpl,
            "prod_order": prod_order,
            "needs_md": needs_md,
        })
    return groups


def build_import_template_bytes(contract, env):
    """Generate an empty import template for a rental contract."""
    try:
        data, _source = env["rental.template"].sudo().get_template_bytes(
            contract.company_id,
            "transport_matrix_xlsx",
        )
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        replacements = {
            "{{b_company}}": contract.b_party.parent_id.name or "",
            "{{b_address}}": contract.b_address or "",
            "{{b_representative}}": contract.b_name or "",
            "{{b_function}}": contract.b_function or "",
            "{{a_representative}}": contract.a_name or "",
            "{{a_company}}": contract.a_party.parent_id.name or "",
            "{{a_function}}": contract.a_function or "",
            "{{construction_work_project}}": contract.construction_work_project_id.name or "",
            "{{construction_work_name}}": contract.construction_work_id.name or "",
            "{{construction_work_address}}": contract.construction_work_address or "",
        }
        replacements.update(transport_matrix_date_replacements(None, None))
        replace_placeholders_in_sheet(ws, replacements)
        title_row, header_second_row, header_third_row, data_start_row = find_transport_matrix_layout(ws)
    except Exception:
        wb = Workbook()
        ws = wb.active
        title_row, header_second_row, header_third_row, data_start_row = 14, 15, 16, 17
        ws.cell(1, 1).value = "BẢNG TỔNG HỢP KHỐI LƯỢNG THUÊ THIẾT BỊ"
        ws.cell(2, 1).value = f"Dự án: {contract.construction_work_project_id.name or ''}"
        ws.cell(title_row, 4).value = "Chủng loại/ Khối lượng"

    ws.cell(header_second_row - 1, 1).value = "STT"
    ws.cell(header_second_row - 1, 2).value = "Ngày tháng"
    ws.cell(header_second_row - 1, 3).value = "Biển số xe"

    groups = _contract_product_groups(contract, env)
    cur_col = _DEFAULT_PRODUCT_START_COL
    for group in groups:
        p_tmpl = group["p_tmpl"]
        if group["needs_md"] and len(group["prod_order"]) > 1:
            span = len(group["prod_order"]) + 1
            ws.merge_cells(
                start_row=header_second_row,
                start_column=cur_col,
                end_row=header_second_row,
                end_column=cur_col + span - 1,
            )
            ws.cell(header_second_row, cur_col).value = p_tmpl["product_tmpl_name"]
            offset = 0
            for prod_id in group["prod_order"]:
                prod = p_tmpl["products"][prod_id]
                label = prod.get("variant_name") or prod.get("name")
                ws.cell(header_third_row, cur_col + offset).value = label
                offset += 1
            ws.cell(header_third_row, cur_col + offset).value = "Tổng MD"
            cur_col += span
        else:
            prod_id = group["prod_order"][0]
            prod = p_tmpl["products"][prod_id]
            ws.merge_cells(
                start_row=header_second_row,
                start_column=cur_col,
                end_row=header_third_row,
                end_column=cur_col,
            )
            ws.cell(header_second_row, cur_col).value = prod["name"]
            cur_col += 1

    ws.cell(data_start_row, 1).value = 1
    ws.cell(data_start_row, 2).value = "dd/mm/yyyy"
    ws.cell(data_start_row, 3).value = "29H-00000"
    ws.cell(data_start_row + 1, 2).value = _(
        "Ghi chú: số lượng dương = xuất (giao), số âm = nhập (trả). Cột C = biển số xe."
    )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def parse_transport_matrix_xlsx(file_bytes, env, contract=None):
    """
    Parse an Excel workbook into transport import rows.

    Returns dict with keys: rows, errors, warnings, product_mapping_errors.
    """
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    ws = _find_import_worksheet(wb)
    layout = _detect_import_layout(ws)
    date_col = layout["date_col"]
    plate_col = layout["plate_col"]
    product_start_col = layout["product_start_col"]
    header_second_row = layout["header_second_row"]
    header_third_row = layout["header_third_row"]
    data_start_row = layout["data_start_row"]
    header_rows = (header_second_row, header_third_row)

    column_map = _build_column_map(ws, header_second_row, header_third_row, product_start_col)
    if not column_map:
        return {
            "rows": [],
            "errors": [_("Không đọc được cột sản phẩm từ file Excel.")],
            "warnings": [],
            "product_mapping_errors": [],
        }

    resolver = TransportMatrixProductResolver(env, contract=contract)
    parsed_rows = []
    global_errors = []
    warnings = []
    product_mapping_errors = []
    seen_product_errors = set()
    in_opening_section = False

    max_row = ws.max_row or data_start_row
    for row_idx in range(data_start_row, max_row + 1):
        row_label = _row_label_text(ws, row_idx, date_col, plate_col)
        if in_opening_section:
            if any(marker in row_label for marker in _OPENING_SECTION_END):
                in_opening_section = False
            continue
        if any(marker in row_label for marker in _OPENING_SECTION_START):
            in_opening_section = True
            continue
        if _should_skip_row(ws, row_idx, date_col, plate_col):
            continue

        date_val = _merged_top_left_value(ws, row_idx, date_col)
        date_text = _normalize_text(date_val)
        if not date_text or _is_skip_row_label(date_text):
            continue
        if date_text.startswith("Ghi chú:"):
            continue
        if _normalize_key(date_text).startswith("đại diện"):
            continue

        transport_date = _parse_excel_date(date_val)
        plate, plate_inferred = _resolve_plate(ws, row_idx, plate_col, max_row)
        row_errors = []
        row_warnings = []
        if plate_inferred:
            row_warnings.append(
                _("Dòng %(row)s: biển số trống, dùng biển số %(plate)s từ dòng phía dưới.") % {
                    "row": row_idx,
                    "plate": plate,
                }
            )

        if not transport_date:
            row_errors.append(_("Dòng %(row)s: không đọc được ngày '%(val)s'.") % {
                "row": row_idx,
                "val": date_val,
            })
        if not plate:
            row_errors.append(_("Dòng %(row)s: thiếu biển số xe (cột %(col)s).") % {
                "row": row_idx,
                "col": get_column_letter(plate_col),
            })

        lines = []
        has_negative = False
        has_positive = False
        for col, col_info in column_map.items():
            qty = _parse_qty(_cell_value(ws, row_idx, col))
            if qty is None:
                continue
            if qty < 0:
                has_negative = True
            elif qty > 0:
                has_positive = True

            product, err = resolver.resolve(
                col_info["group_name"],
                col_info["variant_label"],
                col_index=col,
                header_rows=header_rows,
            )
            if err:
                detail = _("Dòng %(row)s, cột %(col)s:\n%(detail)s") % {
                    "row": row_idx,
                    "col": col_info["col_letter"],
                    "detail": err,
                }
                row_errors.append(detail)
                err_key = (col_info["col_letter"], col_info["label"])
                if err_key not in seen_product_errors:
                    seen_product_errors.add(err_key)
                    product_mapping_errors.append({
                        "col_letter": col_info["col_letter"],
                        "excel_label": col_info["label"],
                        "detail": err,
                    })
                continue
            lines.append((product.id, abs(qty)))

        if has_negative and has_positive:
            row_errors.append(
                _("Dòng %(row)s: không được trộn số lượng dương (xuất) và âm (nhập) trên cùng một chuyến.") % {
                    "row": row_idx,
                }
            )

        if not lines and not row_errors:
            warnings.append(_("Dòng %(row)s: không có sản phẩm, bỏ qua.") % {"row": row_idx})
            continue

        transport_type = "return" if has_negative else "delivery"
        parsed_rows.append({
            "row_number": row_idx,
            "transport_date": transport_date,
            "plate": plate,
            "transport_type": transport_type,
            "lines": lines,
            "errors": row_errors,
            "warnings": row_warnings,
        })
        global_errors.extend(row_errors)
        warnings.extend(row_warnings)

    return {
        "rows": parsed_rows,
        "errors": global_errors,
        "warnings": warnings,
        "product_mapping_errors": product_mapping_errors,
    }
