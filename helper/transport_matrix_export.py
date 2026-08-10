# -*- coding: utf-8 -*-
"""Build transport matrix (volume confirmation) XLSX exports."""
import io
from collections import OrderedDict, defaultdict

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from odoo import _
from odoo.exceptions import UserError

from odoo.addons.rental.models import rental_transport_matrix as rtm

from .xlsx_template_utils import (
    adjust_transport_matrix_data_rows,
    apply_product_column_styles,
    apply_product_data_cell_alignment,
    find_transport_matrix_layout,
    lock_product_column_widths,
    replace_placeholders_in_sheet,
    transport_matrix_date_replacements,
    unmerge_row_span,
)


def build_transport_matrix_into_workbook(env, contract, start_date, end_date, wb=None):
    """Fill the volume-confirmation sheet and return (workbook, qty-ref metadata).

    Metadata keys used by the combined KLCT+HSTT export:
    - sheet_title
    - opening_row (or None)
    - opening_qty_by_tmpl_id: opening quantity represented by each HSTT qty column
    - total_row
    - qty_col_by_tmpl_id: template → product col or Tổng MD col
    - data_rows_by_date: date → [row, ...] (multiple trips same day)
    - data_rows_by_date_and_direction: date → direction → [row, ...]
    """
    # DEC-17: official volume confirmation uses completed transports only (same as billing).
    rr_transport_ids = contract.rr_transport_ids.filtered_domain([
        ('state', '=', 'done'),
        ('start_rental_or_return_date', '>=', start_date),
        ('start_rental_or_return_date', '<=', end_date),
    ])
    rr_transport_ids_before = contract.rr_transport_ids.filtered_domain([
        ('state', '=', 'done'),
        ('start_rental_or_return_date', '<', start_date),
    ])
    # Prefetch lines + products for cell aggregation / MD factors.
    (rr_transport_ids | rr_transport_ids_before).mapped('transport_line_ids.product_id')

    def _cells_from_transports(transport_set):
        cell_qty = {}
        for transport in transport_set:
            for line in transport.transport_line_ids:
                if not line.product_id:
                    continue
                key = (line.product_tmpl_id.id, line.product_id.id)
                cell_qty[key] = (
                    cell_qty.get(key, 0)
                    + rtm._signed_transport_line_qty(line)
                )
        return cell_qty

    opening_cells = _cells_from_transports(rr_transport_ids_before)
    has_opening = any(v != 0 for v in opening_cells.values())

    product_tmpls = OrderedDict()
    for transport in rr_transport_ids_before | rr_transport_ids:
        for line in transport.transport_line_ids:
            if not line.product_tmpl_id or not line.product_id:
                continue
            product_tmpl = product_tmpls.get(line.product_tmpl_id.id, False)
            product = {
                'id': line.product_id.id,
                'name': line.product_id.display_name,
                'variant_name': line.product_id.product_template_variant_value_ids.name,
                'price_multiplier': rtm._variant_price_multiplier(line.product_id),
            }
            if not product_tmpl:
                product_tmpl = {
                    'product_tmpl_id': line.product_tmpl_id.id,
                    'product_tmpl_name': line.product_tmpl_id.display_name,
                    'products': {line.product_id.id: product},
                }
            else:
                product_tmpl['products'][line.product_id.id] = product
            product_tmpls[line.product_tmpl_id.id] = product_tmpl

    for _p_tmpl_id, p_tmpl in product_tmpls.items():
        p_tmpl['products'] = dict(rtm._sort_products_odict(OrderedDict(p_tmpl['products'])))

    groups_meta = []
    for p_tmpl_id, p_tmpl in product_tmpls.items():
        prod_order = list(p_tmpl['products'].keys())
        prods = p_tmpl['products']
        needs_md = rtm._group_needs_md_column(env, prods, prod_order)
        groups_meta.append({
            'tmpl_id': p_tmpl_id,
            'p_tmpl': p_tmpl,
            'prod_order': prod_order,
            'needs_md': needs_md,
        })

    product_ids_2_col = {}
    md_col_by_tmpl = {}
    qty_col_by_tmpl_id = {}
    col_idx = 4
    for g in groups_meta:
        p_tmpl_id = g['tmpl_id']
        if g['needs_md']:
            for prod_id in g['prod_order']:
                product_ids_2_col[f"{p_tmpl_id}_{prod_id}"] = col_idx
                col_idx += 1
            md_col_by_tmpl[p_tmpl_id] = col_idx
            qty_col_by_tmpl_id[p_tmpl_id] = col_idx
            col_idx += 1
        else:
            prod_id = g['prod_order'][0]
            product_ids_2_col[f"{p_tmpl_id}_{prod_id}"] = col_idx
            qty_col_by_tmpl_id[p_tmpl_id] = col_idx
            col_idx += 1
    last_product_col = col_idx - 1
    start_product_col = 4
    if last_product_col < start_product_col:
        raise UserError(_(
            "Không có vận chuyển đã hoàn thành (done) để tính tiền cho kỳ này. "
            "Vui lòng hoàn thành phiếu vận chuyển rồi xuất lại."
        ))

    total_cells = _cells_from_transports(rr_transport_ids_before | rr_transport_ids)

    if wb is None:
        data, _source = env["rental.template"].sudo().get_template_bytes(
            contract.company_id,
            "transport_matrix_xlsx",
        )
        wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    replacements = {
        '{{b_company}}': contract.b_party.parent_id.name or '',
        '{{b_address}}': contract.b_address or '',
        '{{b_representative}}': contract.b_name or '',
        '{{b_function}}': contract.b_function or '',
        '{{a_representative}}': contract.a_name or '',
        '{{a_company}}': contract.a_party.parent_id.name or '',
        '{{a_function}}': contract.a_function or '',
        '{{construction_work_project}}': contract.construction_work_project_id.name or '',
        '{{construction_work_name}}': contract.construction_work_id.name or '',
        '{{construction_work_address}}': contract.construction_work_address or '',
    }
    replacements.update(transport_matrix_date_replacements(start_date, end_date))
    replace_placeholders_in_sheet(ws, replacements)

    title_row, header_second_row, header_third_row, start_row = find_transport_matrix_layout(ws)
    start_product_col = 4
    cur_product_col = start_product_col
    md_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    length_by_prod_id = {}

    def _xlsx_md_sum(cell_qty_map, tmpl_id, p_tmpl):
        total = 0.0
        for prod_id in p_tmpl['products']:
            if prod_id not in length_by_prod_id:
                prod_rec = env['product.product'].browse(prod_id)
                length_by_prod_id[prod_id] = rtm._linear_meter_factor_for_product(prod_rec)
            length = length_by_prod_id[prod_id]
            if length is None:
                continue
            q = cell_qty_map.get((tmpl_id, prod_id), 0) or 0
            total += q * length
        return total

    def _xlsx_md_cell_value(total):
        if total is None or abs(total) < 1e-12:
            return None
        if abs(total - round(total)) < 1e-9:
            return int(round(total))
        return round(float(total), 2)

    def _write_md_cells(r, cell_qty_map):
        for g in groups_meta:
            if not g['needs_md']:
                continue
            p_tmpl_id = g['tmpl_id']
            mcol = md_col_by_tmpl.get(p_tmpl_id)
            if mcol:
                mdv = _xlsx_md_sum(cell_qty_map, p_tmpl_id, g['p_tmpl'])
                ws.cell(r, mcol).value = _xlsx_md_cell_value(mdv)
                ws.cell(r, mcol).fill = md_fill

    for rng in list(ws.merged_cells.ranges):
        ref = str(rng)
        if (
            rng.min_row == title_row
            and rng.min_col >= start_product_col
            and rng.max_col >= start_product_col
        ):
            ws.unmerge_cells(ref)
    ws.merge_cells(
        start_row=title_row,
        start_column=start_product_col,
        end_row=title_row,
        end_column=last_product_col,
    )

    for rng in list(ws.merged_cells.ranges):
        if rng.min_row in (header_second_row, header_third_row) and rng.min_col >= start_product_col:
            ws.unmerge_cells(str(rng))

    for g in groups_meta:
        p_tmpl_id = g['tmpl_id']
        p_tmpl = g['p_tmpl']
        if g['needs_md']:
            span = len(g['prod_order']) + 1
            ws.merge_cells(
                start_row=header_second_row,
                start_column=cur_product_col,
                end_row=header_second_row,
                end_column=cur_product_col + span - 1,
            )
            ws.cell(header_second_row, cur_product_col).value = p_tmpl['product_tmpl_name']
            i = 0
            for prod_id in g['prod_order']:
                prod = p_tmpl['products'][prod_id]
                c = ws.cell(header_third_row, cur_product_col + i)
                c.value = prod['variant_name'] or prod['name']
                i += 1
            md_cell = ws.cell(header_third_row, cur_product_col + i)
            md_cell.value = "Tổng MD"
            md_cell.fill = md_fill
            cur_product_col += span
        else:
            prod_id = g['prod_order'][0]
            prod = p_tmpl['products'][prod_id]
            ws.merge_cells(
                start_row=header_second_row,
                start_column=cur_product_col,
                end_row=header_third_row,
                end_column=cur_product_col,
            )
            ws.cell(header_second_row, cur_product_col).value = prod['name']
            cur_product_col += 1

    needed_rows = (1 if has_opening else 0) + len(rr_transport_ids) + 1
    adjust_transport_matrix_data_rows(ws, start_row, needed_rows)

    # Template / insert_rows_below may clone horizontal merges across product columns
    # (e.g. I:N). Clear them so each qty / Tổng MD cell stays independent.
    clear_to_col = max(last_product_col, ws.max_column or last_product_col)
    for r in range(start_row, start_row + needed_rows):
        unmerge_row_span(ws, r, start_product_col, clear_to_col)

    row_idx = 0
    opening_row = None
    data_rows_by_date = defaultdict(list)
    data_rows_by_date_and_direction = defaultdict(lambda: defaultdict(list))
    # (excel_row, tmpl_id) → qty shown in HSTT-linked column (piece or Tổng MD)
    qty_by_row_and_tmpl_id = {}

    def _tmpl_display_qty(tmpl_id, cell_qty_map):
        g = next((x for x in groups_meta if x["tmpl_id"] == tmpl_id), None)
        if not g:
            return 0.0
        if g["needs_md"]:
            return _xlsx_md_sum(cell_qty_map, tmpl_id, g["p_tmpl"]) or 0.0
        prod_id = g["prod_order"][0]
        return cell_qty_map.get((tmpl_id, prod_id), 0) or 0.0

    def _record_row_tmpl_qtys(excel_row, cell_qty_map):
        for tmpl_id in qty_col_by_tmpl_id:
            qty_by_row_and_tmpl_id[(excel_row, tmpl_id)] = _tmpl_display_qty(
                tmpl_id, cell_qty_map
            )

    if has_opening:
        r = start_row + row_idx
        opening_row = r
        ws.cell(r, 1).value = row_idx + 1
        cell_tondau = ws.cell(r, 2)
        cell_tondau.value = 'Tồn đầu kỳ'
        cell_tondau.font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        for (tmpl_id, prod_id), qty in opening_cells.items():
            col = product_ids_2_col.get(f'{tmpl_id}_{prod_id}')
            if col:
                ws.cell(r, col).value = qty
        _write_md_cells(r, opening_cells)
        _record_row_tmpl_qtys(r, opening_cells)
        row_idx += 1

    for transport in rr_transport_ids:
        r = start_row + row_idx
        ws.cell(r, 1).value = row_idx + 1
        ws.cell(r, 2).value = transport.start_rental_or_return_date.strftime('%-d/%-m/%Y')
        ws.cell(r, 3).value = transport.plate or ''
        cell_row = _cells_from_transports(transport)
        for (tmpl_id, prod_id), qty in cell_row.items():
            col = product_ids_2_col.get(f'{tmpl_id}_{prod_id}')
            if col:
                ws.cell(r, col).value = qty
        _write_md_cells(r, cell_row)
        _record_row_tmpl_qtys(r, cell_row)
        if transport.start_rental_or_return_date:
            transport_date = transport.start_rental_or_return_date
            direction = (
                "return"
                if transport.type in ("return", "compensation")
                else "delivery"
            )
            data_rows_by_date[transport_date].append(r)
            data_rows_by_date_and_direction[transport_date][direction].append(r)
        row_idx += 1

    r_total = start_row + row_idx
    ws.cell(r_total, 1).value = ''
    ws.cell(r_total, 2).value = f"Tổng KL Cuối T{end_date.strftime('%m.%Y')}"
    ws.merge_cells(start_row=r_total, start_column=2, end_row=r_total, end_column=3)
    for (tmpl_id, prod_id), qty in total_cells.items():
        col = product_ids_2_col.get(f'{tmpl_id}_{prod_id}')
        if col:
            ws.cell(r_total, col).value = qty
    _write_md_cells(r_total, total_cells)
    row_idx += 1

    data_end_row = start_row + row_idx - 1
    apply_product_column_styles(
        ws,
        start_col=start_product_col,
        last_col=last_product_col,
        header_rows=[header_second_row, header_third_row],
        ref_col=start_product_col,
        data_start_row=start_row,
        data_end_row=data_end_row,
        total_row=r_total,
        header_wrap=True,
    )

    def _md_bold_font(cell):
        base = cell.font
        return Font(name=base.name, size=base.size, bold=True, italic=base.italic, color=base.color)

    for mcol in md_col_by_tmpl.values():
        header_cell = ws.cell(header_third_row, mcol)
        header_cell.value = "Tổng MD"
        header_cell.fill = md_fill
        header_cell.font = _md_bold_font(header_cell)
        for r in range(start_row, r_total + 1):
            c = ws.cell(r, mcol)
            c.fill = md_fill
            if c.value not in (None, ""):
                c.font = _md_bold_font(c)

    ws.cell(r_total, 1).font = _md_bold_font(ws.cell(r_total, 1))
    ws.cell(r_total, 2).font = _md_bold_font(ws.cell(r_total, 2))
    for col in range(start_product_col, last_product_col + 1):
        ws.cell(r_total, col).font = _md_bold_font(ws.cell(r_total, col))

    lock_product_column_widths(ws, start_product_col, last_product_col, ref_col=start_product_col)
    apply_product_data_cell_alignment(
        ws,
        start_product_col,
        last_product_col,
        start_row,
        data_end_row,
        total_row=r_total,
    )

    opening_qty_by_tmpl_id = {}
    for g in groups_meta:
        tmpl_id = g["tmpl_id"]
        if g["needs_md"]:
            opening_qty_by_tmpl_id[tmpl_id] = _xlsx_md_sum(
                opening_cells, tmpl_id, g["p_tmpl"]
            )
        else:
            prod_id = g["prod_order"][0]
            opening_qty_by_tmpl_id[tmpl_id] = opening_cells.get(
                (tmpl_id, prod_id), 0
            ) or 0

    meta = {
        "sheet_title": ws.title,
        "opening_row": opening_row,
        "opening_qty_by_tmpl_id": opening_qty_by_tmpl_id,
        "qty_by_row_and_tmpl_id": qty_by_row_and_tmpl_id,
        "total_row": r_total,
        "qty_col_by_tmpl_id": qty_col_by_tmpl_id,
        "data_rows_by_date": dict(data_rows_by_date),
        "data_rows_by_date_and_direction": {
            transport_date: dict(rows_by_direction)
            for transport_date, rows_by_direction
            in data_rows_by_date_and_direction.items()
        },
    }
    return wb, meta


def build_transport_matrix_xlsx_bytes(env, contract, start_date, end_date):
    """Return XLSX bytes for the rental volume confirmation table."""
    wb, _meta = build_transport_matrix_into_workbook(env, contract, start_date, end_date)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
