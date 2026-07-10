# -*- coding: utf-8 -*-
"""Shared helpers for rental Excel templates (placeholders, header styles)."""
from copy import copy

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

_TRANSPORT_MATRIX_FOOTER_MARKERS = ("Hà Nội", "Đơn vị cho thuê")


def replace_placeholders_in_sheet(ws, replacements):
    """Replace {{key}} substrings in all string cells (including merged top-left)."""
    for row in ws.iter_rows():
        for cell in row:
            val = cell.value
            if val is None or not isinstance(val, str):
                continue
            new_val = val
            for key, repl in replacements.items():
                if key in new_val:
                    new_val = new_val.replace(key, str(repl))
            if new_val != val:
                cell.value = new_val


def set_cell_wrap_text(cell, wrap=True, vertical="center", horizontal="center"):
    if cell.alignment:
        al = copy(cell.alignment)
        al.wrapText = wrap
        if vertical:
            al.vertical = vertical
        if horizontal:
            al.horizontal = horizontal
        cell.alignment = al
    else:
        cell.alignment = Alignment(
            wrap_text=wrap,
            vertical=vertical,
            horizontal=horizontal,
        )


def clone_cell_style(ws, src_row, src_col, dst_row, dst_col, wrap_text=False):
    src = ws.cell(src_row, src_col)
    dst = ws.cell(dst_row, dst_col)
    if src.has_style:
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.number_format = copy(src.number_format)
        dst.protection = copy(src.protection)
        dst.alignment = copy(src.alignment)
    if wrap_text:
        set_cell_wrap_text(dst, wrap=True)


def apply_product_column_styles(
    ws,
    start_col,
    last_col,
    header_rows,
    ref_col,
    data_start_row=None,
    data_end_row=None,
    total_row=None,
    header_wrap=True,
):
    """
    Copy styles from ref_col onto every product column (D..last), including
    columns inside the template width that were overwritten when building headers.
    """
    if last_col < start_col:
        return
    for col in range(start_col, last_col + 1):
        for row in header_rows:
            clone_cell_style(
                ws, row, ref_col, row, col, wrap_text=header_wrap
            )
        if data_start_row is not None and data_end_row is not None:
            for row in range(data_start_row, data_end_row + 1):
                clone_cell_style(ws, data_start_row, ref_col, row, col)
        if total_row:
            clone_cell_style(ws, total_row, ref_col, total_row, col)


def find_transport_matrix_layout(ws, default=(14, 15, 16, 17)):
    """Return (title_row, group_header_row, variant_header_row, data_start_row)."""
    for row in ws.iter_rows(min_row=1, max_row=40):
        for cell in row:
            val = cell.value
            if isinstance(val, str) and "Chủng loại" in val:
                title_row = cell.row
                return title_row, title_row + 1, title_row + 2, title_row + 3
    return default


def find_transport_matrix_footer_row(ws, min_row=15, default=None):
    """Return the first row index of the signature/footer block."""
    footer_row = None
    for row in ws.iter_rows(min_row=min_row):
        for cell in row:
            val = cell.value
            if not isinstance(val, str):
                continue
            if any(marker in val for marker in _TRANSPORT_MATRIX_FOOTER_MARKERS):
                if footer_row is None or cell.row < footer_row:
                    footer_row = cell.row
    return footer_row if footer_row is not None else default


def adjust_transport_matrix_data_rows(ws, start_row, needed_rows, footer_row=None):
    """
    Grow or shrink the data area so it fits exactly ``needed_rows`` rows,
    keeping the footer block immediately below the data table.
    Returns the footer row index after adjustment.
    """
    footer_row = footer_row or find_transport_matrix_footer_row(ws, min_row=start_row)
    if not footer_row or footer_row <= start_row:
        return footer_row

    available = footer_row - start_row
    if needed_rows > available:
        from .export_excel_template import insert_rows_below

        extra = needed_rows - available
        insert_rows_below(ws, footer_row - 1, extra)
        footer_row += extra
    elif needed_rows < available:
        remove = available - needed_rows
        ws.delete_rows(start_row + needed_rows, remove)
        footer_row -= remove
    return footer_row


def lock_product_column_widths(ws, start_col, end_col, ref_col=None):
    """Apply a fixed column width from the template reference column."""
    ref_col = ref_col or start_col
    ref_letter = get_column_letter(ref_col)
    ref_dim = ws.column_dimensions.get(ref_letter)
    width = ref_dim.width if ref_dim and ref_dim.width else 7.0
    for col in range(start_col, end_col + 1):
        letter = get_column_letter(col)
        dim = ws.column_dimensions[letter]
        dim.width = width
        dim.bestFit = False


def apply_product_data_cell_alignment(
    ws,
    start_col,
    end_col,
    data_start_row,
    data_end_row,
    total_row=None,
):
    """Shrink-to-fit numeric cells so Excel Windows does not auto-expand columns."""
    rows = list(range(data_start_row, data_end_row + 1))
    if total_row and total_row not in rows:
        rows.append(total_row)
    for row in rows:
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row, col)
            if cell.alignment:
                al = copy(cell.alignment)
                al.shrinkToFit = True
                al.wrapText = False
                al.horizontal = "center"
                al.vertical = "center"
                cell.alignment = al
            else:
                cell.alignment = Alignment(
                    shrink_to_fit=True,
                    wrap_text=False,
                    horizontal="center",
                    vertical="center",
                )


def transport_matrix_date_replacements(start_date, end_date):
    """Placeholder values for the period line in transport matrix templates."""
    date_fmt = "%d/%m/%Y"
    from_text = start_date.strftime(date_fmt) if start_date else ""
    to_text = end_date.strftime(date_fmt) if end_date else ""
    return {
        "{{from_date}}": from_text,
        "{{to_date}}": to_text,
        "{{start_date}}": from_text,
        "{{end_date}}": to_text,
    }


def unmerge_row_span(ws, row, min_col, max_col):
    """Remove horizontal merges on a row between min_col and max_col."""
    to_remove = []
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row == rng.max_row == row and rng.min_col >= min_col and rng.max_col <= max_col:
            to_remove.append(str(rng))
        elif rng.min_row == rng.max_row == row and not (rng.max_col < min_col or rng.min_col > max_col):
            to_remove.append(str(rng))
    for ref in to_remove:
        ws.unmerge_cells(ref)
